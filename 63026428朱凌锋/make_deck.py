# -*- coding: utf-8 -*-
"""生成 5 页极简版 PPT：2024 新课标 II 卷第 19 题（一题一解法）"""
import re
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------- 配色 ----------------
def C(h):
    return RGBColor((h >> 16) & 255, (h >> 8) & 255, h & 255)

INK    = C(0x121C2A)
BODY   = C(0x3A4656)
MUTED  = C(0x7E8896)
GOLD   = C(0xBA8A3A)
GOLD_L = C(0xFAF3E5)
PAPER  = C(0xFAF8F3)
WHITE  = C(0xFFFFFF)
LINE   = C(0xE6DFD2)
NAVY   = C(0x0F1E31)
GHOST  = C(0x142338)

CN = 'Microsoft YaHei'
MATH = 'Cambria'

# ---------------- 富文本：^{} _{} ~斜体开关 ----------------
def parse(s, base_italic=False):
    out, buf = [], ''
    sup = sub = 0
    it = base_italic

    def flush():
        nonlocal buf
        if buf:
            out.append((buf, sup > 0, sub > 0, it))
            buf = ''

    i = 0
    while i < len(s):
        if s.startswith('^{', i):
            flush(); sup += 1; i += 2
        elif s.startswith('_{', i):
            flush(); sub += 1; i += 2
        elif s[i] == '}':
            flush()
            if sup > 0:
                sup -= 1
            elif sub > 0:
                sub -= 1
            i += 1
        elif s[i] == '~':
            flush(); it = not it; i += 1
        else:
            buf += s[i]; i += 1
    flush()
    return out


def set_ea(run, typeface=CN):
    rPr = run.font._rPr
    latin = rPr.find(qn('a:latin'))
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = etree.SubElement(rPr, qn('a:ea'))
        if latin is not None:
            latin.addnext(ea)
    ea.set('typeface', typeface)


def set_baseline(run, v):
    run.font._rPr.set('baseline', str(v))


def put_text(tf, paras, anchor=MSO_ANCHOR.TOP):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = p.get('align', PP_ALIGN.LEFT)
        para.space_before = Pt(p.get('space_before', 0))
        para.space_after = Pt(p.get('space_after', 0))
        para.line_spacing = p.get('line', 1.0)
        for txt, sup, sub, it in parse(p['t'], p.get('italic', False)):
            r = para.add_run()
            r.text = txt
            f = r.font
            f.size = Pt(p.get('size', 14))
            f.bold = p.get('bold', False)
            f.italic = it
            f.color.rgb = p.get('color', BODY)
            f.name = p.get('font', CN)
            set_ea(r, CN)
            if sup:
                set_baseline(r, 30000)
            elif sub:
                set_baseline(r, -25000)


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    put_text(box.text_frame, paras, anchor)
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=0.06,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=0.75):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    s.text_frame.word_wrap = True
    return s


def chip(slide, x, y, d, num, fill=WHITE, fg=INK, edge=GOLD, size=12):
    s = rect(slide, x, y, d, d, fill=fill, line=edge, shape=MSO_SHAPE.OVAL)
    tf = s.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = num
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = fg
    r.font.name = CN
    set_ea(r)
    return s


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def frame(slide, section, title, page):
    bg(slide, PAPER)
    rect(slide, 0.89, 0.55, 0.11, 0.11, GOLD, None, shape=MSO_SHAPE.RECTANGLE)
    text(slide, 1.10, 0.48, 5.0, 0.28,
         [{'t': section, 'size': 9.5, 'color': GOLD, 'bold': True}])
    text(slide, 10.3, 0.48, 2.12, 0.28,
         [{'t': page, 'size': 9.5, 'color': MUTED, 'align': PP_ALIGN.RIGHT}])
    text(slide, 0.86, 0.80, 11.3, 0.72,
         [{'t': title, 'size': 27, 'color': INK, 'bold': True}])
    rect(slide, 0.89, 1.52, 0.62, 0.045, GOLD, None, shape=MSO_SHAPE.RECTANGLE)
    rect(slide, 0.89, 6.93, 11.55, 0.008, LINE, None, shape=MSO_SHAPE.RECTANGLE)
    text(slide, 0.89, 7.05, 8.0, 0.24,
         [{'t': '2024 新课标 Ⅱ 卷 · 第 19 题', 'size': 8.5, 'color': MUTED}])


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

# ============================ 1 封面 ============================
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
text(s, 8.35, 2.55, 4.3, 3.4,
     [{'t': '19', 'size': 210, 'color': GHOST, 'bold': True, 'align': PP_ALIGN.RIGHT}])
rect(s, 0.89, 2.30, 0.045, 2.60, GOLD, None, shape=MSO_SHAPE.RECTANGLE)
text(s, 1.22, 2.26, 7.6, 2.9, [
    {'t': '2024 年 · 新课标 Ⅱ 卷', 'size': 11.5, 'color': GOLD, 'bold': True},
    {'t': '双曲线与数列', 'size': 40, 'color': WHITE, 'bold': True, 'space_before': 14},
    {'t': '压轴第 19 题 · 一题三问，一条主线', 'size': 17,
     'color': C(0xC8D2E0), 'space_before': 12},
])
rect(s, 1.22, 5.62, 2.6, 0.008, C(0x3A4E68), None, shape=MSO_SHAPE.RECTANGLE)
text(s, 1.22, 5.78, 7.0, 0.28,
     [{'t': '17 分　·　解析几何 + 数列', 'size': 10.5, 'color': C(0x8A97AA)}])

# ============================ 2 原题 ============================
s = prs.slides.add_slide(BLANK)
frame(s, '原题', '题目', '02 / 05')

rect(s, 0.89, 1.78, 11.55, 1.30, WHITE, LINE, 0.05)
text(s, 1.14, 2.10, 11.05, 1.00, [
    {'t': '~C~: ~x~^{2} − ~y~^{2} = ~m~ (~m~ > 0)，点 ~P~_{1}(5, 4) 在 ~C~ 上，0 < ~k~ < 1．',
     'size': 15, 'color': INK},
    {'t': '过 ~P~_{~n~−1} 作斜率为 ~k~ 的直线，交左支于 ~Q~_{~n~−1}；~P~_{~n~} 为 ~Q~_{~n~−1} 关于 ~y~ 轴的对称点．',
     'size': 15, 'color': INK, 'space_before': 10},
])

rect(s, 0.89, 3.30, 11.55, 2.82, GOLD_L, LINE, 0.05)
qs = [
    '若 ~k~ = 1/2，求 ~x~_{2}，~y~_{2}．',
    '证明：数列 ~x~_{~n~} − ~y~_{~n~} 是公比为 (1 + ~k~)/(1 − ~k~) 的等比数列．',
    '设 ~S~_{~n~} 为 △~P~_{~n~}~P~_{~n~+1}~P~_{~n~+2} 的面积，证明：~S~_{~n~} = ~S~_{~n~+1}．',
]
for i, q in enumerate(qs):
    y = 3.76 + i * 0.78
    chip(s, 1.14, y, 0.34, str(i + 1))
    text(s, 1.66, y - 0.02, 10.4, 0.34,
         [{'t': q, 'size': 14.5, 'color': INK}])

text(s, 0.89, 6.44, 11.55, 0.26,
     [{'t': '提示：先证 ~P~_{~n~} 都在右支上，后面才敢作除法．', 'size': 10, 'color': MUTED}])

# ============================ 3 第 (1)(2) 问 ============================
s = prs.slides.add_slide(BLANK)
frame(s, '第 (1) (2) 问', '先算出来，再证明', '03 / 05')

# 左卡
rect(s, 0.89, 1.78, 5.55, 4.62, WHITE, LINE, 0.05)
text(s, 1.14, 1.96, 5.0, 0.28, [{'t': '(1)　k = 1/2', 'size': 11, 'color': GOLD, 'bold': True}])
text(s, 1.14, 2.32, 5.06, 2.6, [
    {'t': '~y~ = (~x~ + 3)/2　代入　~x~^{2} − ~y~^{2} = 9', 'size': 14, 'color': INK},
    {'t': '~x~^{2} − 2~x~ − 15 = 0', 'size': 15, 'color': INK, 'space_before': 10},
    {'t': '~x~ = −3　⇒　~Q~_{1}(−3, 0)　⇒　~P~_{2}(3, 0)', 'size': 14,
     'color': INK, 'space_before': 10},
])
rect(s, 1.14, 4.62, 5.06, 1.28, GOLD_L, None, 0.06)
text(s, 1.38, 4.78, 4.6, 0.24, [{'t': '答案', 'size': 9.5, 'color': GOLD, 'bold': True}])
text(s, 1.38, 5.02, 4.6, 0.62,
     [{'t': '~x~_{2} = 3，~y~_{2} = 0', 'size': 24, 'color': INK}],
     anchor=MSO_ANCHOR.TOP)

# 右卡
rect(s, 6.89, 1.78, 5.55, 4.62, WHITE, LINE, 0.05)
text(s, 7.14, 1.96, 5.0, 0.28, [{'t': '(2)　点差法', 'size': 11, 'color': GOLD, 'bold': True}])
text(s, 7.14, 2.34, 5.06, 1.0, [
    {'t': '把“在直线上”和“在双曲线上”各写一遍，再相减：', 'size': 12, 'color': MUTED},
])
rect(s, 7.14, 2.92, 0.035, 1.32, GOLD, None, shape=MSO_SHAPE.RECTANGLE)
text(s, 7.38, 2.94, 4.82, 1.30, [
    {'t': '(1 − ~k~)(~x~_{~n~} − ~y~_{~n~})', 'size': 15, 'color': INK},
    {'t': '= (1 + ~k~)(~x~_{~n~−1} − ~y~_{~n~−1})', 'size': 15, 'color': INK,
     'space_before': 4},
])
text(s, 7.14, 4.42, 5.06, 0.28, [{'t': '于是', 'size': 11, 'color': MUTED}])
text(s, 7.14, 4.72, 5.06, 0.60,
     [{'t': '公比　~t~ = (1 + ~k~)/(1 − ~k~)', 'size': 21, 'color': INK}])
text(s, 7.14, 5.44, 5.06, 0.28,
     [{'t': '首项 ~x~_{1} − ~y~_{1} = 1 ≠ 0，故为等比数列．', 'size': 11, 'color': MUTED}])

# ============================ 4 第 (3) 问 ============================
s = prs.slides.add_slide(BLANK)
frame(s, '第 (3) 问', '面积为什么不变', '04 / 05')

cards = [
    ('①', '等比结构', '~p~_{~n~} = ~x~_{~n~} − ~y~_{~n~}', '~p~_{~n~}·~p~_{~n~+3} = ~p~_{~n~+1}·~p~_{~n~+2}'),
    ('②', '两弦平行', '~p~~q~ 相同 ⇒ 方向相同', '~P~_{~n~}~P~_{~n~+3} ∥ ~P~_{~n~+1}~P~_{~n~+2}'),
    ('③', '同底等高', '共用底边 ~P~_{~n~+1}~P~_{~n~+2}', '~S~_{~n~} = ~S~_{~n~+1}'),
]
x0, w, gap = 0.89, 3.55, 0.45
for i, (num, head, l1, l2) in enumerate(cards):
    x = x0 + i * (w + gap)
    rect(s, x, 1.86, w, 2.86, WHITE, LINE, 0.05)
    chip(s, x + 0.34, 2.16, 0.40, num, WHITE, GOLD, GOLD, 13)
    text(s, x + 0.34, 2.80, w - 0.68, 0.34,
         [{'t': head, 'size': 15, 'color': INK, 'bold': True}])
    rect(s, x + 0.34, 3.26, 0.40, 0.035, GOLD, None, shape=MSO_SHAPE.RECTANGLE)
    text(s, x + 0.34, 3.48, w - 0.68, 1.30, [
        {'t': l1, 'size': 10.5, 'color': MUTED},
        {'t': l2, 'size': 13, 'color': INK, 'space_before': 10},
    ])

rect(s, 0.89, 5.06, 11.55, 0.92, GOLD_L, LINE, 0.06)
text(s, 1.14, 5.22, 11.05, 0.62, [
    {'t': '两个三角形共用底边，第三个顶点在同一条平行线上 —— 底相同、高相同，面积自然相同．',
     'size': 12.5, 'color': INK},
])

text(s, 0.89, 6.30, 11.55, 0.30,
     [{'t': '（顺便：~p~_{~n~}~p~_{~n~+1} ≠ ~p~_{~n~}~p~_{~n~+2}，三点不共线，三角形不会退化．）',
       'size': 9.5, 'color': MUTED}])

# ============================ 5 结论 ============================
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
rect(s, 0.89, 2.55, 0.045, 1.60, GOLD, None, shape=MSO_SHAPE.RECTANGLE)
text(s, 1.22, 2.48, 8.0, 0.30, [{'t': '结论', 'size': 10.5, 'color': GOLD, 'bold': True}])
text(s, 1.22, 2.84, 9.0, 0.90,
     [{'t': '~S~_{~n~} = 36~k~^{3} / (1 − ~k~^{2})^{2}', 'size': 36, 'color': WHITE}])
text(s, 1.22, 3.96, 9.6, 0.34,
     [{'t': '与 ~n~ 无关：每个 △~P~_{~n~}~P~_{~n~+1}~P~_{~n~+2} 的面积都一样．',
       'size': 14, 'color': C(0xC2CDDC)}])
rect(s, 1.22, 5.10, 2.6, 0.008, C(0x3A4E68), None, shape=MSO_SHAPE.RECTANGLE)
text(s, 1.22, 5.26, 8.0, 0.28,
     [{'t': 'k = 1/2 时，~S~_{~n~} ≡ 8．', 'size': 10.5, 'color': C(0x8A97AA)}])

out = 'ai_agent_demo.pptx'
prs.save(out)
print('saved', out, len(prs.slides.__iter__.__self__._sldIdLst), 'slides')
