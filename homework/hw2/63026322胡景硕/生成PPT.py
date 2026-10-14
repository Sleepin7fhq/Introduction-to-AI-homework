#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成「2026 新高考Ⅰ卷 数学第18题」解析 PPT

依赖：python-pptx
运行：python 生成PPT.py
产出：2026新高考I卷数学第18题解析.pptx

排版约定
--------
1. 全部使用纯文本 + Unicode 数学符号，不使用 LaTeX。
2. **只用字体内确实存在的符号**。本机 微软雅黑 缺少
   ⟹ ⇒ ✓ ⁻ ⁶ ₀ ₉ ₁ ₂ 等，写进去会显示成豆腐块（□）。故：
   - 推导箭头一律用 →（两种字体都支持）
   - 对勾用 √（中文数学惯例，且字体支持）
   - 下标 ₁ ₂ 只出现在 Consolas 等宽 run 里（Consolas 支持）
   - 科学计数法写成 4.4e-16，避免上标 ⁻ ⁶
3. 公式用等宽字体 Consolas，说明文字用微软雅黑。
"""
import math
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ----------------------------------------------------------------- 配色与常量
NAVY = RGBColor(0x16, 0x2C, 0x4F)
BLUE = RGBColor(0x1F, 0x6F, 0xB2)
ACCENT = RGBColor(0xD9, 0x6A, 0x1B)
GRAY = RGBColor(0x44, 0x4A, 0x52)
MUTED = RGBColor(0x7A, 0x88, 0x9B)
LIGHT = RGBColor(0xEE, 0xF3, 0xF9)
CREAM = RGBColor(0xFD, 0xEE, 0xDD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xC3, 0xD2, 0xE3)
AXIS = RGBColor(0xB6, 0xC2, 0xD2)

FONT = "微软雅黑"
MONO = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
HERE = os.path.dirname(os.path.abspath(__file__))


# ----------------------------------------------------------------- 基础工具
def style_run(run, size, bold, color, font):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        rPr.append(rPr.makeelement(qn(tag), {"typeface": font}))


def set_para(p, text, size=17, bold=False, color=GRAY, font=FONT,
             space_after=8, align=PP_ALIGN.LEFT, line_spacing=1.15):
    p.text = text
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    for run in p.runs:
        style_run(run, size, bold, color, font)
    return p


def textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def add_title(slide, title, kicker=None):
    y = 0.42
    if kicker:
        tf = textbox(slide, 0.75, 0.30, 11.8, 0.3)
        set_para(tf.paragraphs[0], kicker, size=12.5, bold=True, color=ACCENT, space_after=0)
        y = 0.64
    tf = textbox(slide, 0.75, y, 11.8, 0.7)
    set_para(tf.paragraphs[0], title, size=27, bold=True, color=NAVY, space_after=0)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.78), Inches(y + 0.60),
                                 Inches(1.5), Pt(3.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    bar.shadow.inherit = False


def add_footer(slide, num):
    tf = textbox(slide, 0.75, 6.98, 10.5, 0.28)
    set_para(tf.paragraphs[0],
             "人工智能导论 · 第二次 Demo 实践　|　63026322 胡景硕　|　2026 新高考 I 卷 数学第 18 题",
             size=9.5, color=RGBColor(0x9C, 0xAA, 0xBC), space_after=0)
    tf = textbox(slide, 12.2, 6.98, 0.65, 0.28)
    set_para(tf.paragraphs[0], str(num), size=9.5, color=RGBColor(0x9C, 0xAA, 0xBC),
             space_after=0, align=PP_ALIGN.RIGHT)


def body_lines(slide, items, left=0.9, top=1.62, width=11.55, height=5.1):
    """items: (text, level[, size, bold, color, font])"""
    tf = textbox(slide, left, top, width, height)
    first = True
    for it in items:
        text, level = it[0], it[1]
        size = it[2] if len(it) > 2 else (17 if level == 0 else 15.5)
        bold = it[3] if len(it) > 3 else False
        color = it[4] if len(it) > 4 else (NAVY if level == 0 else GRAY)
        font = it[5] if len(it) > 5 else FONT
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        set_para(p, ("　　" * level) + text, size=size, bold=bold, color=color,
                 font=font, space_after=9 if level == 0 else 6)
    return tf


def panel(slide, left, top, width, height, fill=LIGHT):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                                 Inches(width), Inches(height))
    box.adjustments[0] = 0.06
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = LINE
    box.line.width = Pt(1)
    box.shadow.inherit = False
    return box


def fill_panel(box, items, pad=0.26):
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(0.14)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    first = True
    for it in items:
        text, size = it[0], it[1]
        bold = it[2] if len(it) > 2 else False
        color = it[3] if len(it) > 3 else GRAY
        font = it[4] if len(it) > 4 else FONT
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        set_para(p, text, size=size, bold=bold, color=color, font=font, space_after=5)
    return tf


def add_table(slide, rows, left, top, width, height, col_widths=None, font_size=13):
    nrow, ncol = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nrow, ncol, Inches(left), Inches(top),
                                Inches(width), Inches(height))
    table = gt.table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(val)
            cell.margin_left = Inches(0.12)
            cell.margin_right = Inches(0.10)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (WHITE if r % 2 else LIGHT)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                for run in p.runs:
                    style_run(run, font_size, r == 0, WHITE if r == 0 else GRAY, FONT)
    return table


# ----------------------------------------------------------------- 几何作图
def draw_geometry(slide, cx=3.05, cy=4.00, unit=1.12):
    """按 k = √5/2 的真实坐标作图"""
    k = math.sqrt(5) / 2

    def M(x, y):
        return Inches(cx + x * unit), Inches(cy - y * unit)

    a, b = 2.0, math.sqrt(3)

    ell = slide.shapes.add_shape(MSO_SHAPE.OVAL, *M(-a, b),
                                 Inches(2 * a * unit), Inches(2 * b * unit))
    ell.fill.background()
    ell.line.color.rgb = BLUE
    ell.line.width = Pt(2.0)
    ell.shadow.inherit = False

    for p0, p1 in (((-2.45, 0), (2.45, 0)), ((0, -2.05), (0, 2.05))):
        cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *M(*p0), *M(*p1))
        cn.line.color.rgb = AXIS
        cn.line.width = Pt(1)

    P = (0.5, k * 1.5)
    Q = (-1.75, k * (-0.75))
    R = (-P[0], -P[1])
    fb = slide.shapes.build_freeform(*M(*P), scale=1.0)
    fb.add_line_segments([M(*Q), M(*R)], close=True)
    tri = fb.convert_to_shape()
    tri.fill.solid()
    tri.fill.fore_color.rgb = CREAM
    tri.line.color.rgb = ACCENT
    tri.line.width = Pt(1.25)
    tri.shadow.inherit = False

    x0, x1 = -2.15, 0.85
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                    *M(x0, k * (x0 + 1)), *M(x1, k * (x1 + 1)))
    cn.line.color.rgb = ACCENT
    cn.line.width = Pt(1.75)

    for (px, py), label, dx, dy, col in (
        (P, "P", 0.11, -0.27, BLUE),
        (Q, "Q", -0.32, -0.17, BLUE),
        (R, "R", -0.34, 0.06, BLUE),
        ((0, 0), "O", -0.28, 0.03, GRAY),
        ((-1, 0), "F", -0.07, 0.10, ACCENT),
    ):
        d = slide.shapes.add_shape(MSO_SHAPE.OVAL, *M(px - 0.08, py + 0.08),
                                   Inches(0.16 * unit), Inches(0.16 * unit))
        d.fill.solid()
        d.fill.fore_color.rgb = col
        d.line.fill.background()
        d.shadow.inherit = False
        tf = textbox(slide, cx + px * unit + dx, cy - py * unit + dy, 0.9, 0.28)
        set_para(tf.paragraphs[0], label, size=14, bold=True, color=col, space_after=0)

    tf = textbox(slide, 0.75, cy + (b + 0.30) * unit, 4.7, 0.75)
    set_para(tf.paragraphs[0], "以 (i) 的解 k = √5/2 作图", size=10.5, color=MUTED, space_after=2)
    p = tf.add_paragraph()
    set_para(p, "Q(-1.75, -0.839)  P(0.5, 1.677)  R(-0.5, -1.677)",
             size=10.5, color=MUTED, font=MONO, space_after=0)


# ================================================================== 幻灯片
def new_slide():
    return prs.slides.add_slide(BLANK)


# ---------------------------------------------------------------- 1 封面
s = new_slide()
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(2.55))
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
bg.shadow.inherit = False
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.55), Inches(2.2), Pt(5))
bar.fill.solid()
bar.fill.fore_color.rgb = ACCENT
bar.line.fill.background()
bar.shadow.inherit = False

tf = textbox(s, 0.9, 0.70, 11.5, 1.7)
set_para(tf.paragraphs[0], "2026 年新高考 I 卷 数学 第 18 题", size=40, bold=True,
         color=WHITE, space_after=8)
p = tf.add_paragraph()
set_para(p, "椭圆中的面积比与夹角最值", size=22, color=RGBColor(0x9E, 0xC3, 0xE8), space_after=0)

tf = textbox(s, 0.9, 3.30, 11.5, 2.1)
set_para(tf.paragraphs[0], "从原题核实、严格推导，到数值验证的完整过程", size=19,
         color=GRAY, space_after=24)
for line in ("工具：DeepSeek Harness（DSH）",
             "提交人：63026322 胡景硕",
             "课程：人工智能导论 · 第二次 Demo 实践"):
    p = tf.add_paragraph()
    set_para(p, line, size=15, color=MUTED, space_after=8)

fill_panel(panel(s, 0.9, 5.50, 11.55, 1.10), [
    ("结论速览", 12, True, ACCENT),
    ("（1）x²/4 + y²/3 = 1　　（2)(i)  l: y = (√5/2)(x+1)　　（2)(ii)  min tan∠PQR = 4√3",
     15.5, True, NAVY),
])

# ---------------------------------------------------------------- 2 原题
s = new_slide()
add_title(s, "原题", kicker="题目")
fill_panel(panel(s, 0.9, 1.62, 11.55, 2.15, WHITE), [
    ("已知椭圆 C：x²/a² + y²/b² = 1（a > b > 0）的左焦点为 F(-1, 0)，离心率为 1/2。",
     16.5, False, GRAY),
    ("（1）求 C 的方程；", 16.5, True, NAVY),
    ("（2）设 O 为坐标原点，过 F 且斜率大于 0 的动直线 l 与 C 交于 P、Q 两点，其中 Q 在第三象限，",
     16.5, False, GRAY),
    ("　　　直线 PO 与 C 的另一个交点为 R。", 16.5, False, GRAY),
])
body_lines(s, [
    ("（i）若 △PQR 的面积是 △PFO 的面积的 3 倍，求 l 的方程；", 0, 17, True, NAVY),
    ("（ii）求 tan∠PQR 的最小值。", 0, 17, True, NAVY),
], top=4.00)
fill_panel(panel(s, 0.9, 4.95, 11.55, 1.75), [
    ("题干要点", 12, True, ACCENT),
    ("· 椭圆参数未知，但由「左焦点 F(-1, 0)」与「离心率 1/2」已完全确定", 14.5, False, GRAY),
    ("· 动直线 l 过焦点 F，斜率 k > 0，与椭圆有两个交点 P、Q", 14.5, False, GRAY),
    ("· 关键构造：直线 PO 与椭圆的另一交点 R —— 即 P 关于中心的对称点", 14.5, False, GRAY),
])

# ---------------------------------------------------------------- 3 来源
s = new_slide()
add_title(s, "题目来源与核对", kicker="资料核实")
add_table(s, [
    ["来源", "用途"],
    ["良师宝 · 题库", "本题面以此为准（公式完整可读）"],
    ["21 世纪教育网", "佐证题号 18、分值 17 分与三问结构"],
    ["21 世纪教育网 · 组卷", "佐证该题属「2026 年新高考 I 卷数学试卷」"],
    ["知乎 · 圆锥曲线真题 G001", "佐证题型为圆锥曲线"],
    ["中国知网《中学数学》2026(07)", "佐证该题确为 2026 新高考 I 卷第 18 题"],
], left=0.9, top=1.66, width=11.55, height=2.45, col_widths=[3.5, 8.05], font_size=13)
fill_panel(panel(s, 0.9, 4.32, 11.55, 2.35), [
    ("核对中的两个关键点", 12, True, ACCENT),
    ("① 公式来源：部分网站在未登录时会剥离公式，公开页面里公式位置为空，", 14, False, GRAY),
    ("　　故题面采自公式完整的良师宝，并与 21 世纪教育网的文字骨架逐字比对。", 14, False, GRAY),
    ("② 字母命名考据：不同网站用 M,N,P 或 P,Q,R，数学内容相同。", 14, False, GRAY),
    ("　　本文采用 P,Q,R —— 因为只有该解读下面积条件有解", 14, False, GRAY),
    ("　　（另一解读会导出 √(1+k²) = 3/14 < 1 的矛盾）。", 14, False, GRAY),
])

# ---------------------------------------------------------------- 4 第(1)问
s = new_slide()
add_title(s, "第（1）问　求 C 的方程", kicker="解答")
body_lines(s, [("由左焦点 F(-1, 0) 得半焦距　c = |OF| = 1", 0, 17),
               ("由离心率　e = c / a = 1/2，得", 0, 17)], top=1.66)
fill_panel(panel(s, 2.6, 2.72, 8.1, 0.85, WHITE),
           [("a = c / e = 1 / (1/2) = 2", 20, True, NAVY, MONO)])
body_lines(s, [("于是", 0, 17)], top=3.72)
fill_panel(panel(s, 2.6, 4.16, 8.1, 0.85, WHITE),
           [("b² = a² - c² = 4 - 1 = 3", 20, True, NAVY, MONO)])
fill_panel(panel(s, 0.9, 5.35, 11.55, 1.15), [
    ("所以", 12, True, ACCENT),
    ("C: x²/4 + y²/3 = 1", 24, True, ACCENT, MONO),
])

# ---------------------------------------------------------------- 5 预备
s = new_slide()
add_title(s, "第（2）问 · 预备：直线与椭圆联立", kicker="解答")
body_lines(s, [("设 l：y = k(x + 1)（k > 0），即 kx - y + k = 0", 0, 17),
               ("代入椭圆方程 3x² + 4y² = 12：", 0, 17)], top=1.66)
fill_panel(panel(s, 1.4, 2.72, 10.5, 1.35, WHITE), [
    ("3x² + 4k²(x + 1)² = 12", 17, True, NAVY, MONO),
    ("  →  (3 + 4k²)x² + 8k²x + (4k² - 12) = 0", 17, True, NAVY, MONO),
])
body_lines(s, [("判别式与韦达定理：", 0, 17)], top=4.30)
fill_panel(panel(s, 1.4, 4.78, 10.5, 1.72), [
    ("Δ = (8k²)² - 4(3 + 4k²)(4k² - 12) = 144(k² + 1) > 0", 15.5, True, GRAY, MONO),
    ("故 l 与 C 恒有两交点。设 P(x₁, y₁)、Q(x₂, y₂)，由韦达定理：", 14.5, False, GRAY, MONO),
    ("x₁ + x₂ = -8k² / (3 + 4k²)      x₁·x₂ = (4k² - 12) / (3 + 4k²)", 15.5, True, GRAY, MONO),
])

# ---------------------------------------------------------------- 6 几何事实
s = new_slide()
add_title(s, "两个关键几何事实", kicker="解答 · 核心")
fill_panel(panel(s, 0.9, 1.70, 5.55, 4.55, WHITE), [
    ("事实 ①　R 是 P 关于原点的对称点", 14.5, True, ACCENT),
    ("椭圆中心为原点，直线 PO 过中心，", 14, False, GRAY),
    ("故它与椭圆的另一交点 R 必满足", 14, False, GRAY),
    ("R = -P = (-x₁, -y₁)", 19, True, NAVY, MONO),
    ("", 10, False, GRAY),
    ("事实 ②　d(R, l) = 2·d(O, l)", 14.5, True, ACCENT),
    ("由 F ∈ l，得", 14, False, GRAY),
    ("d(O, l) = k / √(k² + 1)", 19, True, NAVY, MONO),
])
fill_panel(panel(s, 6.75, 1.70, 5.70, 4.55), [
    ("事实 ② 的推导", 14.5, True, ACCENT),
    ("代入 R = -P 且 y₁ = k(x₁ + 1)：", 14, False, GRAY, MONO),
    ("d(R, l) = |k(-x₁) - (-y₁) + k| / √(k²+1)", 14, True, GRAY, MONO),
    ("         = |-k·x₁ + k(x₁+1) + k| / √(k²+1)", 14, True, GRAY, MONO),
    ("         = 2k / √(k² + 1)", 14, True, GRAY, MONO),
    ("         = 2·d(O, l)", 14, True, ACCENT, MONO),
    ("", 8, False, GRAY),
    ("这两个事实把面积问题转化成「同底比高」，", 14, False, GRAY),
    ("是整道题的枢纽。", 14, True, NAVY),
])

# ---------------------------------------------------------------- 7 (i) 面积比
s = new_slide()
add_title(s, "第（2)(i) 问 · 面积比转化", kicker="解答")
draw_geometry(s)
body_lines(s, [("△PQR 与 △PFO 都以 l 上的线段作底边：", 0, 15)], left=5.95, top=1.68, width=6.5)
fill_panel(panel(s, 5.95, 2.14, 6.5, 1.25, WHITE), [
    ("S(△PQR) = (1/2)·|PQ|·d(R, l)", 13.5, True, GRAY, FONT),
    ("S(△PFO) = (1/2)·|PF|·d(O, l)", 13.5, True, GRAY, FONT),
])
body_lines(s, [("两式相除，代入 d(R, l) = 2·d(O, l) 及面积比为 3：", 0, 15)],
           left=5.95, top=3.52, width=6.5)
fill_panel(panel(s, 5.95, 4.18, 6.5, 1.45), [
    ("3 = (|PQ| / |PF|) · 2", 16, True, GRAY, MONO),
    ("  →  |PQ| / |PF| = 3/2      …… (*)", 16, True, ACCENT, MONO),
])
fill_panel(panel(s, 5.95, 5.78, 6.5, 0.95, WHITE), [
    ("共底在两直线 l 上时，面积比 = 底边比 × 高的比", 13, False, MUTED),
])

# ---------------------------------------------------------------- 8 (i) 求解
s = new_slide()
add_title(s, "第（2)(i) 问 · 求解斜率 k", kicker="解答")
body_lines(s, [("直线 l 上的长度 = 横坐标之差 × 弧长因子 √(1 + k²)：", 0, 15.5)], top=1.66)
fill_panel(panel(s, 1.4, 2.10, 10.5, 1.00, WHITE), [
    ("|PQ| = √(1+k²)·(x₁ - x₂)      |PF| = √(1+k²)·(x₁ + 1)", 15.5, True, GRAY, MONO),
])
body_lines(s, [("由求根公式：", 0, 15.5)], top=3.28)
fill_panel(panel(s, 1.4, 3.72, 10.5, 1.45, WHITE), [
    ("x₁ - x₂ = √Δ / (3+4k²) = 12√(k²+1) / (3+4k²)", 14.5, True, GRAY, MONO),
    ("x₁ + 1 = (3 + 6√(k²+1)) / (3+4k²)", 14.5, True, GRAY, MONO),
])
body_lines(s, [("代入 (*)：", 0, 15.5)], top=5.30)
fill_panel(panel(s, 1.4, 5.74, 10.5, 1.02), [
    ("12√(k²+1) / (3 + 6√(k²+1)) = 3/2   →   √(k²+1) = 3/2   →   k = √5/2",
     15.5, True, ACCENT, MONO),
])

# ---------------------------------------------------------------- 9 (i) 结论
s = new_slide()
add_title(s, "第（2)(i) 问 · 结论", kicker="解答")
fill_panel(panel(s, 0.9, 1.78, 11.55, 1.50), [
    ("k² + 1 = 9/4   →   k² = 5/4，因 k > 0 得", 14.5, False, GRAY),
    ("k = √5 / 2", 26, True, ACCENT, MONO),
])
fill_panel(panel(s, 0.9, 3.52, 11.55, 1.30, WHITE), [
    ("所以直线 l 的方程为", 13, True, GRAY),
    ("y = (√5/2)(x + 1)     即    √5·x - 2y + √5 = 0", 21, True, NAVY, MONO),
])
fill_panel(panel(s, 0.9, 5.05, 11.55, 1.45, WHITE), [
    ("几何核对（代入 k = √5/2）", 12, True, ACCENT),
    ("Q(-1.75, -0.839) 在第三象限 √", 14, False, GRAY),
    ("|PQ| / |PF| = 3/2 √　　d(R,l) / d(O,l) = 2 √", 14, False, GRAY),
])

# ---------------------------------------------------------------- 10 (ii) 斜率
s = new_slide()
add_title(s, "第（2)(ii) 问 · 求 QR 的斜率", kicker="解答")
body_lines(s, [("∠PQR 的顶点为 Q，两射线为 QP 与 QR。", 0, 16),
               ("其中 QP 就是直线 l，斜率 k(QP) = k。", 0, 16)], top=1.66)
body_lines(s, [("求 k(QR)：", 0, 16)], top=2.82)
fill_panel(panel(s, 1.4, 3.28, 10.5, 1.45, WHITE), [
    ("k(QR) = (y_R - y_Q) / (x_R - x_Q) = (-y₁ - y₂) / (-x₁ - x₂)", 15, True, GRAY, MONO),
    ("      = (y₁ + y₂) / (x₁ + x₂)", 15, True, GRAY, MONO),
])
body_lines(s, [("由 y = k(x + 1) 得 y₁ + y₂ = k(x₁ + x₂) + 2k = 6k / (3 + 4k²)", 0, 15, False, GRAY, MONO)],
           top=4.88)
fill_panel(panel(s, 1.4, 5.38, 10.5, 1.05), [
    ("k(QR) = [6k / (3+4k²)] ÷ [-8k² / (3+4k²)] = -3 / (4k)", 18, True, ACCENT, MONO),
])

# ---------------------------------------------------------------- 11 夹角公式
s = new_slide()
add_title(s, "第（2)(ii) 问 · 夹角公式", kicker="解答")
fill_panel(panel(s, 0.9, 1.70, 11.55, 0.95, WHITE), [
    ("k(QP)·k(QR) = k·(-3/(4k)) = -3/4 ≠ -1   →   两直线不垂直，∠PQR ≠ 90°",
     15.5, True, GRAY, FONT),
])
body_lines(s, [("由两直线夹角的正切公式：", 0, 16)], top=2.80)
fill_panel(panel(s, 1.2, 3.28, 10.9, 1.60), [
    ("tan∠PQR = | (k(QR) - k(QP)) / (1 + k(QP)·k(QR)) |", 16.5, True, GRAY, FONT),
    ("= | (-3/(4k) - k) / (1 - 3/4) | = 4·| 3/(4k) + k |", 16.5, True, GRAY, FONT),
])
body_lines(s, [("由 k > 0 得 3/(4k) + k > 0，去掉绝对值：", 0, 16)], top=5.05)
fill_panel(panel(s, 1.2, 5.50, 10.9, 1.15, WHITE), [
    ("tan∠PQR = 4·(3/(4k) + k) = 3/k + 4k", 22, True, NAVY, FONT),
])

# ---------------------------------------------------------------- 12 基本不等式
s = new_slide()
add_title(s, "第（2)(ii) 问 · 基本不等式求最值", kicker="解答")
fill_panel(panel(s, 0.9, 1.75, 11.55, 1.70, WHITE), [
    ("因 k > 0，由基本不等式（AM-GM）：", 15, False, GRAY),
    ("tan∠PQR = 3/k + 4k ≥ 2·√( (3/k)·4k ) = 2√12 = 4√3", 21, True, NAVY, FONT),
])
fill_panel(panel(s, 0.9, 3.70, 11.55, 1.30), [
    ("等号成立当且仅当", 14, False, GRAY),
    ("3/k = 4k   →   k² = 3/4   →   k = √3/2   (> 0，符合题设)", 18, True, ACCENT, MONO),
])
fill_panel(panel(s, 0.9, 5.20, 11.55, 1.30), [
    ("所以", 12, True, ACCENT),
    ("tan∠PQR 的最小值为 4√3　（此时 l 的斜率 k = √3/2）", 21, True, ACCENT),
])

# ---------------------------------------------------------------- 13 数值验证
s = new_slide()
add_title(s, "数值验证", kicker="独立复核")
body_lines(s, [("用 Python 3 标准库独立复核（二分法求根 + 20 万点扫描 + 三分法求极值）：", 0, 15)],
           top=1.66)
add_table(s, [
    ["核验项", "数值结果", "解析值", "偏差（科学计数法）"],
    ["(2)(i) 面积比为 3 时的 k", "1.118033988750", "√5/2", "4.4e-16"],
    ["解析式 tan∠PQR = 4k + 3/k", "8 个取样点比对", "—", "最大 3.7e-14"],
    ["(2)(ii) tan∠PQR 的最小值", "6.928203230275", "4√3", "9.8e-15"],
], left=0.9, top=2.32, width=11.55, height=2.05,
    col_widths=[4.3, 3.3, 1.9, 2.05], font_size=13)
fill_panel(panel(s, 0.9, 4.65, 11.55, 1.90), [
    ("复核要点", 12, True, ACCENT),
    ("· |PQ|/|PF| = 3/2 与 d(R,l)/d(O,l) = 2 均精确成立", 14, False, GRAY),
    ("· 解析式在 8 个取样点上与数值结果的最大偏差仅 3.7e-14", 14, False, GRAY),
    ("· 数值实验只作佐证，严格结论仍以上述解析证明为准", 14, True, NAVY),
])

# ---------------------------------------------------------------- 14 汇总
s = new_slide()
add_title(s, "答案汇总与考点", kicker="小结")
add_table(s, [
    ["小问", "结论"],
    ["（1）", "C：x²/4 + y²/3 = 1"],
    ["（2)(i)", "l：y = (√5/2)(x + 1)，即 √5·x - 2y + √5 = 0"],
    ["（2)(ii)", "tan∠PQR 的最小值为 4√3（k = √3/2 时取得）"],
], left=0.9, top=1.68, width=11.55, height=2.00, col_widths=[2.2, 9.35], font_size=15)
fill_panel(panel(s, 0.9, 3.95, 11.55, 2.55), [
    ("考点与方法", 12, True, ACCENT),
    ("① 椭圆标准方程与几何性质：由焦点、离心率反解 a、b", 14, False, GRAY),
    ("② 直线与椭圆联立：判别式判断交点个数，韦达定理沟通坐标", 14, False, GRAY),
    ("③ 面积转化的核心技巧：两三角形共底在 l 上时，面积比 = 底边比 × 高的比", 14, False, GRAY),
    ("④ 中心对称的妙用：R 是 P 的对径点，使 d(R,l) = 2·d(O,l)，一步消去复杂坐标", 14, False, GRAY),
    ("⑤ 夹角公式 + 基本不等式：化为 3/k + 4k 的 AM-GM，等号条件顺带给出取到最值的 k",
     14, False, GRAY),
])

# ---------------------------------------------------------------- 页脚
for i, sl in enumerate(prs.slides, 1):
    if i > 1:
        add_footer(sl, i)

out = os.path.join(HERE, "2026新高考I卷数学第18题解析.pptx")
prs.save(out)
print("已生成：", out)
print("幻灯片总数：", len(prs.slides._sldIdLst))
