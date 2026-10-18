"""
生成《第19题 集合D与函数增长》12页数学课件。

依赖：
    pip install python-pptx

运行：
    python build_deck.py

输出：
    与本脚本同目录的《第19题_集合D与函数增长_解析.pptx》
"""
# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

# ---------------- palette ----------------
INK      = RGBColor(0x10, 0x22, 0x33)
INK_2    = RGBColor(0x1C, 0x34, 0x46)
BG       = RGBColor(0xF6, 0xF2, 0xE9)
PAPER    = RGBColor(0xFF, 0xFD, 0xF8)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
GOLD     = RGBColor(0xB8, 0x89, 0x4B)
GOLD_L   = RGBColor(0xEA, 0xD9, 0xBF)
TEAL     = RGBColor(0x1E, 0x6B, 0x6E)
TEAL_D   = RGBColor(0x15, 0x50, 0x53)
TEAL_L   = RGBColor(0xDD, 0xEB, 0xE7)
CORAL    = RGBColor(0xB4, 0x4C, 0x3B)
CORAL_L  = RGBColor(0xF6, 0xE1, 0xDB)
MUTED    = RGBColor(0x62, 0x71, 0x7C)
LINE     = RGBColor(0xDD, 0xD2, 0xC0)
CREAM    = RGBColor(0xF1, 0xE7, 0xD4)

CN = "Microsoft YaHei"
MATH = "Cambria Math"
SW, SH = 13.333, 7.5
M = 0.72
TOTAL = 12

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
prs.core_properties.title = "第19题：集合D与函数增长"
prs.core_properties.subject = "高三数学新定义压轴题解析"
prs.core_properties.author = "AI 教学助理"
prs.core_properties.keywords = "集合D, 函数单调性, 反证法, 高考数学"


def set_run_font(run, size, color, bold=False, italic=False, latin=MATH, ea=CN,
                 baseline=None, spc=None):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = latin
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea_el = rPr.find(qn('a:ea'))
    if ea_el is None:
        ea_el = etree.SubElement(rPr, qn('a:ea'))
    ea_el.set('typeface', ea)
    cs_el = rPr.find(qn('a:cs'))
    if cs_el is None:
        cs_el = etree.SubElement(rPr, qn('a:cs'))
    cs_el.set('typeface', latin)
    rPr.set('lang', 'zh-CN')
    rPr.set('altLang', 'en-US')
    if baseline is not None:
        rPr.set('baseline', str(baseline))
    if spc is not None:
        rPr.set('spc', str(spc))


def _tokens(s):
    out, buf, i = [], "", 0
    while i < len(s):
        if s.startswith('^{', i):
            if buf:
                out.append((buf, 'n')); buf = ""
            j = s.find('}', i + 2)
            if j == -1:
                buf += s[i:]; break
            out.append((s[i + 2:j], 'sup')); i = j + 1
        elif s.startswith('_{', i):
            if buf:
                out.append((buf, 'n')); buf = ""
            j = s.find('}', i + 2)
            if j == -1:
                buf += s[i:]; break
            out.append((s[i + 2:j], 'sub')); i = j + 1
        elif s.startswith('**', i):
            if buf:
                out.append((buf, 'n')); buf = ""
            j = s.find('**', i + 2)
            if j == -1:
                buf += s[i:]; break
            out.append((s[i + 2:j], 'b')); i = j + 2
        else:
            buf += s[i]; i += 1
    if buf:
        out.append((buf, 'n'))
    return out


def add_runs(p, text, size, color, bold=False, italic=False,
             latin=MATH, ea=CN, spc=None):
    for txt, kind in _tokens(text):
        r = p.add_run()
        r.text = txt
        if kind == 'sup':
            set_run_font(r, size * 0.68, color, bold=bold, italic=italic,
                         latin=latin, ea=ea, baseline=48000, spc=spc)
        elif kind == 'sub':
            set_run_font(r, size * 0.66, color, bold=bold, italic=italic,
                         latin=latin, ea=ea, baseline=-22000, spc=spc)
        elif kind == 'b':
            set_run_font(r, size, color, bold=True, italic=italic,
                         latin=latin, ea=ea, spc=spc)
        else:
            set_run_font(r, size, color, bold=bold, italic=italic,
                         latin=latin, ea=ea, spc=spc)


def fill_text(tf, lines, size=15, color=INK, bold=False, italic=False,
              align=PP_ALIGN.LEFT, line_spacing=1.12, latin=MATH, ea=CN,
              margins=(0, 0, 0, 0), anchor=MSO_ANCHOR.TOP, spc=None):
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom = [Inches(x) for x in margins]
    if isinstance(lines, (str, dict)):
        lines = [lines]
    for idx, item in enumerate(lines):
        d = item if isinstance(item, dict) else {"text": item}
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = d.get("align", align)
        p.line_spacing = d.get("line_spacing", line_spacing)
        p.space_before = Pt(d.get("before", 0))
        p.space_after = Pt(d.get("after", 0))
        add_runs(p, d.get("text", ""), d.get("size", size), d.get("color", color),
                 d.get("bold", bold), d.get("italic", italic),
                 d.get("latin", latin), d.get("ea", ea), d.get("spc", spc))


def textbox(slide, x, y, w, h, lines, size=15, color=INK, bold=False,
            align=PP_ALIGN.LEFT, line_spacing=1.12, latin=MATH, ea=CN,
            margins=(0, 0, 0, 0), anchor=MSO_ANCHOR.TOP, spc=None, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_text(box.text_frame, lines, size=size, color=color, bold=bold, italic=italic,
              align=align, line_spacing=line_spacing, latin=latin, ea=ea,
              margins=margins, anchor=anchor, spc=spc)
    return box


def rect(slide, x, y, w, h, fill, line=None, radius=None, line_w=0.9):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shp_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def card(slide, x, y, w, h, fill=PAPER, line=LINE, radius=0.045):
    return rect(slide, x, y, w, h, fill, line, radius=radius)


def line(slide, x1, y1, x2, y2, color=LINE, width=1.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def oval(slide, x, y, w, h, fill, line_color=None, line_w=1.0):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def pill(slide, x, y, w, h, text, fill, color, size=10, border=None):
    shp = rect(slide, x, y, w, h, fill, border, radius=0.5, line_w=0.8)
    fill_text(shp.text_frame, [text], size=size, color=color, bold=True,
              align=PP_ALIGN.CENTER, latin=CN, ea=CN,
              margins=(0.02, 0.02, 0.01, 0.01), anchor=MSO_ANCHOR.MIDDLE)
    return shp


def set_bg(slide, color):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def header(slide, part, title, accent=GOLD):
    rect(slide, 0.47, 0.61, 0.065, 0.83, accent, None)
    textbox(slide, 0.73, 0.43, 9.8, 0.30, [part], size=10, color=GOLD,
            bold=True, latin=CN, ea=CN, spc=80)
    textbox(slide, 0.73, 0.78, 11.35, 0.55, [title], size=24.5, color=INK,
            bold=True, latin=CN, ea=CN)
    textbox(slide, 10.35, 0.46, 2.25, 0.30, ["第 19 题 · 17 分"], size=9.5, color=MUTED,
            align=PP_ALIGN.RIGHT, latin=CN, ea=CN)
    line(slide, M, 1.48, SW - M, 1.48, LINE, 0.8)


def footer(slide, num):
    line(slide, M, 7.06, SW - M, 7.06, LINE, 0.6)
    textbox(slide, M, 7.13, 7.5, 0.22, ["集合 D 与函数增长 · 第 19 题解析"], size=8, color=MUTED,
            latin=CN, ea=CN)
    textbox(slide, 10.4, 7.13, 2.2, 0.22, [f"{num:02d} / {TOTAL:02d}"], size=8, color=MUTED,
            align=PP_ALIGN.RIGHT, latin=CN, ea=CN)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def step_badge(slide, x, y, num, color=TEAL, d=0.34):
    oval(slide, x, y, d, d, color, None)
    textbox(slide, x, y + 0.01, d, d - 0.02, [str(num)], size=11.5, color=WHITE,
            bold=True, align=PP_ALIGN.CENTER, latin=CN, ea=CN,
            anchor=MSO_ANCHOR.MIDDLE)


# ================================================================
# 1 cover
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, INK)
rect(s, 8.28, 0, 5.053, 7.5, INK_2, None)
line(s, 0.92, 1.18, 0.92, 6.12, GOLD, 1.0)
textbox(s, 1.18, 1.14, 6.6, 0.32, ["高考数学 · 新定义压轴题"], size=12, color=GOLD,
        bold=True, latin=CN, ea=CN, spc=100)
textbox(s, 1.16, 1.72, 6.9, 1.95, ["从集合 D 读懂", "函数的“增长方向”"], size=40,
        color=CREAM, bold=True, line_spacing=1.10, latin=CN, ea=CN)
line(s, 1.18, 3.92, 4.28, 3.92, GOLD, 1.2)
textbox(s, 1.18, 4.18, 6.55, 0.9, ["第 19 题（17 分）完整解析", "分类讨论 · 反证法 · 集合包含"],
        size=16, color=RGBColor(0xD9, 0xE2, 0xE6), line_spacing=1.55, latin=CN, ea=CN)
textbox(s, 1.18, 6.18, 5.2, 0.32, ["高三数学 · 解题过程与思路复盘"], size=10,
        color=RGBColor(0x9F, 0xB0, 0xBA), latin=CN, ea=CN)
oval(s, 8.53, 1.15, 4.20, 4.20, None, RGBColor(0x6E, 0x58, 0x3B), 1.1)
oval(s, 9.48, 2.10, 2.30, 2.30, None, RGBColor(0x8B, 0x6B, 0x43), 0.8)
for (dx, dy) in [(10.59, 1.02), (12.58, 3.20), (10.59, 5.28), (8.47, 3.20)]:
    oval(s, dx, dy, 0.10, 0.10, GOLD, None)
textbox(s, 8.53, 2.55, 4.20, 0.9, ["D(x_{0})"], size=43, color=GOLD,
        align=PP_ALIGN.CENTER, latin=MATH, ea=MATH)
textbox(s, 8.40, 3.58, 4.46, 0.7, ["{ d ∈ ℝ  |  f(x_{0}+d) > f(x_{0}) }"], size=12.5,
        color=RGBColor(0xE6, 0xD9, 0xC4), align=PP_ALIGN.CENTER, latin=MATH, ea=MATH)
textbox(s, 8.53, 5.62, 4.20, 0.32, ["一个集合，三种函数性质"], size=10.5,
        color=RGBColor(0xB9, 0xA9, 0x91), align=PP_ALIGN.CENTER, latin=CN, ea=CN)
notes(s, "开场：D(x0)只有一句话——记录所有能让函数值变大的平移量。三问分别在算集合、比集合、用集合反推函数性质。")

# ================================================================
# 2 problem
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s, BG)
header(s, "READ · 题目", "题目条件与三个小问")

# definition banner
b = rect(s, 0.72, 1.66, 11.88, 1.36, INK, None, radius=0.05)
textbox(s, 0.98, 1.82, 11.35, 0.38,
        ["当 x<0 时，f(x)=2^{x}；  D(x_{0})={ d∈ℝ | f(x_{0}+d)>f(x_{0}) }。"],
        size=17.5, color=CREAM, bold=True, latin=MATH, ea=CN)
textbox(s, 0.98, 2.35, 11.35, 0.34,
        ["f 的定义域为 ℝ，x_{0} 可取任意实数。"],
        size=11.5, color=RGBColor(0xBF,0xCD,0xD4), latin=MATH, ea=CN)

# (1)
card(s, 0.72, 3.16, 11.88, 0.86)
pill(s, 0.94, 3.31, 0.66, 0.28, "（1）", TEAL_L, TEAL_D, 9.5)
textbox(s, 1.78, 3.33, 10.45, 0.36,
        ["若当 x≥0 时，f(x)=1−x，求 D(−1)。"],
        size=14.3, color=INK, latin=MATH, ea=CN)

# (2)
card(s, 0.72, 4.12, 11.88, 0.98)
pill(s, 0.94, 4.28, 0.66, 0.28, "（2）", GOLD_L, INK, 9.5)
textbox(s, 1.78, 4.20, 10.48, 0.74,
        ["若 f(x) 是奇函数，f(x_{1})≤f(x_{2})，且 x_{1}x_{2}≠0，",
         "证明：D(x_{2})⊆D(x_{1})。"],
        size=14.0, color=INK, line_spacing=1.22, latin=MATH, ea=CN)

# (3)
card(s, 0.72, 5.20, 11.88, 1.70)
pill(s, 0.94, 5.36, 0.66, 0.28, "（3）", CORAL_L, CORAL, 9.5)
textbox(s, 1.78, 5.28, 10.48, 1.50,
        [{"text": "设 f 满足：① 若 f(x_{1})≤f(x_{2})，则 D(x_{2})⊆D(x_{1})；", "size": 13.6, "after": 1},
         {"text": "② 当 0<x<1 时，f(x)<f(0)。", "size": 13.6, "after": 2},
         {"text": "证明：(i) f(0)≥1；(ii) f(x) 在区间 (0,+∞) 单调递增。", "size": 13.6}],
        line_spacing=1.20, latin=MATH, ea=CN)
footer(s, 2)
notes(s, "先把题目完整读一遍。条件①的方向是本题的骨架：函数值小的一方，D 反而更大。")

# ================================================================
# 3 definition
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s, BG)
header(s, "READ · 翻译", "把新定义翻译成一句人话")

card(s, 0.82, 1.72, 11.68, 1.62, TEAL_L, TEAL_L)
textbox(s, 0.82, 1.94, 11.68, 0.66, ["d∈D(x_{0})  ⇔  f(x_{0}+d)>f(x_{0})"],
        size=27, color=TEAL_D, bold=True, align=PP_ALIGN.CENTER, latin=MATH, ea=MATH)
textbox(s, 0.82, 2.74, 11.68, 0.36, ["从 x_{0} 平移到 x_{0}+d 后，函数值比原来大。"],
        size=14, color=INK_2, align=PP_ALIGN.CENTER, latin=CN, ea=CN)

# two inference cards
card(s, 0.82, 3.62, 5.72, 2.46)
textbox(s, 1.06, 3.82, 5.18, 0.32, ["推论 ①（最常用）"], size=11, color=GOLD, bold=True, latin=CN, ea=CN)
textbox(s, 1.06, 4.22, 5.18, 0.42, ["f(a)≤f(b) ⇒ D(b)⊆D(a)"], size=18, color=INK, bold=True, latin=MATH, ea=MATH)
textbox(s, 1.06, 4.80, 5.12, 0.94,
        ["函数值小的一方，增长方向集合反而更大。", "条件①是“反向包含”，方向不能写反。"],
        size=12.6, color=MUTED, line_spacing=1.28, latin=CN, ea=CN)

card(s, 6.88, 3.62, 5.72, 2.46)
textbox(s, 7.12, 3.82, 5.18, 0.32, ["推论 ②（同值可换）"], size=11, color=TEAL, bold=True, latin=CN, ea=CN)
textbox(s, 7.12, 4.22, 5.18, 0.42, ["f(a)=f(b) ⇒ D(a)=D(b)"], size=18, color=INK, bold=True, latin=MATH, ea=MATH)
textbox(s, 7.12, 4.80, 5.12, 0.94,
        ["两个方向都包含，所以两个集合相等。", "遇到“同值”时，可以直接替换使用。"],
        size=12.6, color=MUTED, line_spacing=1.28, latin=CN, ea=CN)

textbox(s, 0.88, 6.38, 11.6, 0.34,
        ["后面所有证明，都在反复使用这两条推论。"], size=12, color=GOLD, italic=True,
        align=PP_ALIGN.CENTER, latin=CN, ea=CN)
footer(s, 3)
notes(s, "这一页只做一件事：让学生把 d∈D(x0) 和 f(x0+d)>f(x0) 自由互换，并记住条件①的包含方向。")

# ================================================================
# 4 part 1
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s, BG)
header(s, "PART 01 · 直接计算", "第（1）问：把集合语言化成不等式")
textbox(s, 0.78, 1.70, 11.7, 0.42,
        ["f(−1)=1/2。令 t=−1+d，则 d∈D(−1) ⇔ f(t)>1/2。"],
        size=15.2, color=INK, latin=MATH, ea=CN)

card(s, 0.72, 2.40, 5.72, 2.06)
pill(s, 0.96, 2.60, 1.12, 0.30, "t < 0", GOLD_L, INK, 10.5)
textbox(s, 0.96, 3.04, 5.15, 1.12,
        [{"text": "f(t)=2^{t}>2^{−1}", "size": 16.5, "after": 2},
         {"text": "⇒ t>−1", "size": 16.5, "after": 4},
         {"text": "所以 0<d<1。", "size": 15.0, "color": TEAL, "bold": True}],
        line_spacing=1.16, latin=MATH, ea=MATH)

card(s, 6.90, 2.40, 5.72, 2.06)
pill(s, 7.14, 2.60, 1.30, 0.30, "t ≥ 0", TEAL_L, TEAL_D, 10.5)
textbox(s, 7.14, 3.04, 5.15, 1.12,
        [{"text": "f(t)=1−t>1/2", "size": 16.5, "after": 2},
         {"text": "⇒ 0≤t<1/2", "size": 16.5, "after": 4},
         {"text": "所以 1≤d<3/2。", "size": 15.0, "color": TEAL, "bold": True}],
        line_spacing=1.16, latin=MATH, ea=MATH)

# number line
yline = 5.15
line(s, 1.20, yline, 11.02, yline, RGBColor(0xC8, 0xBD, 0xAA), 2.0)
scale = (11.02 - 1.20) / 2.0
line(s, 1.20, yline, 1.20 + scale, yline, TEAL, 7.0)
line(s, 1.20 + scale, yline, 1.20 + 1.5 * scale, yline, GOLD, 7.0)
oval(s, 1.15, yline - 0.06, 0.12, 0.12, BG, TEAL, 1.4)
oval(s, 1.20 + scale - 0.065, yline - 0.065, 0.13, 0.13, TEAL, None)
oval(s, 1.20 + 1.5 * scale - 0.06, yline - 0.06, 0.12, 0.12, BG, GOLD, 1.4)
for val, lab in [(0, "0"), (1, "1"), (1.5, "3/2"), (2, "2")]:
    xx = 1.20 + val * scale
    textbox(s, xx - 0.30, yline + 0.18, 0.60, 0.26, [lab], size=11.5, color=MUTED,
            align=PP_ALIGN.CENTER, latin=MATH, ea=MATH)
textbox(s, 11.05, yline - 0.30, 0.30, 0.25, ["d"], size=12, color=MUTED, latin=MATH, ea=MATH)
textbox(s, 1.40, yline + 0.48, 7.9, 0.28, ["d=1 属于；d=3/2 不属于。"], size=10.5, color=MUTED,
        latin=CN, ea=CN)

ans = rect(s, 0.72, 5.98, 5.55, 0.70, TEAL, None, radius=0.08)
fill_text(ans.text_frame, ["D(−1) = ( 0 , 3/2 )"], size=21, color=WHITE, bold=True,
          align=PP_ALIGN.CENTER, latin=MATH, ea=MATH, anchor=MSO_ANCHOR.MIDDLE)
textbox(s, 6.55, 6.08, 6.0, 0.56,
        ["两个区间在 d=1 处相接，所以并起来是一个连续区间。"], size=12.2, color=MUTED,
        line_spacing=1.15, latin=CN, ea=CN)
footer(s, 4)
notes(s, "端点最容易丢分：d=1 时 t=0，f(0)=1>1/2，属于；d=3/2 时 t=1/2，函数值恰好等于1/2，不属于。")

# ================================================================
# 5 part 2 formulas
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s, BG)
header(s, "PART 02 · 解析式", "第（2）问：奇函数补全正半轴")

card(s, 0.72, 1.70, 4.30, 4.92)
textbox(s, 0.96, 1.90, 3.82, 0.30, ["由奇函数"], size=12.5, color=TEAL, bold=True, latin=CN, ea=CN)
textbox(s, 0.96, 2.30, 3.82, 0.90,
        [{"text": "f(0)=0", "size": 17, "after": 3},
         {"text": "x>0 时：", "size": 13.5, "after": 1},
         {"text": "f(x)=−f(−x)=−2^{−x}", "size": 13.8}],
        line_spacing=1.22, latin=MATH, ea=CN)
line(s, 0.98, 3.42, 4.76, 3.42, LINE, 0.8)
textbox(s, 0.98, 3.58, 3.80, 0.26, ["完整解析式"], size=10.5, color=MUTED, latin=CN, ea=CN)
textbox(s, 1.02, 3.92, 3.78, 2.10,
        [{"text": "f(x) =", "size": 14.5, "after": 6},
         {"text": "2^{x}，      x<0", "size": 15.5, "after": 3},
         {"text": "0，          x=0", "size": 15.5, "after": 3},
         {"text": "−2^{−x}，  x>0", "size": 15.5}],
        line_spacing=1.16, latin=MATH, ea=MATH)
textbox(s, 0.96, 6.10, 3.82, 0.30, ["负半轴已知，正半轴由奇性得到。"], size=10.5, color=MUTED,
        latin=CN, ea=CN)

card(s, 5.34, 1.70, 7.26, 2.22)
textbox(s, 5.58, 1.90, 6.74, 0.28, ["x < 0 时"], size=12, color=TEAL, bold=True, latin=CN, ea=CN)
textbox(s, 5.58, 2.20, 6.64, 0.44, ["D(x) = ( 0 , −x )"], size=21, color=INK, bold=True,
        latin=MATH, ea=MATH)
textbox(s, 5.58, 2.76, 6.70, 0.96,
        [{"text": "f(x)=2^{x}>0。要 f(x+d)>f(x)，", "size": 12.4, "after": 1},
         {"text": "必须仍落在负半轴，于是 0<d<−x。", "size": 12.4}],
        line_spacing=1.22, latin=MATH, ea=CN)

card(s, 5.34, 4.16, 7.26, 2.46)
textbox(s, 5.58, 4.36, 6.74, 0.28, ["x > 0 时"], size=12, color=GOLD, bold=True, latin=CN, ea=CN)
textbox(s, 5.58, 4.66, 6.64, 0.44, ["D(x) = (−∞, −x] ∪ (0, +∞)"], size=19, color=INK, bold=True,
        latin=MATH, ea=MATH)
textbox(s, 5.58, 5.18, 6.70, 1.20,
        [{"text": "f(x)=−2^{−x}<0。", "size": 12.4, "after": 1},
         {"text": "当 x+d≤0 时，f(x+d)≥0>f(x)；", "size": 12.4, "after": 1},
         {"text": "当 x+d>0 时，需且只需 d>0。", "size": 12.4}],
        line_spacing=1.22, latin=MATH, ea=CN)
footer(s, 5)
notes(s, "正半轴的 D 有一个向左的负区间 (-∞,-x]：只要把点平移到 0 或负半轴，函数值就从负数变成非负数。")

# ================================================================
# 6 part 2 proof
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s, BG)
header(s, "PART 02 · 证明", "第（2）问：分类比较两个集合")
textbox(s, 0.78, 1.68, 11.6, 0.36,
        ["f 在 (−∞,0) 与 (0,+∞) 上分别严格递增，先判断 x_{1}、x_{2} 的位置。"],
        size=12.0, color=MUTED, latin=MATH, ea=CN)

card(s, 0.72, 2.10, 3.86, 3.56)
pill(s, 0.94, 2.28, 1.10, 0.30, "① 同负", TEAL_L, TEAL_D, 10)
textbox(s, 0.94, 2.68, 3.44, 0.30, ["x_{1}≤x_{2}<0"], size=13.5, color=TEAL_D, bold=True,
        latin=MATH, ea=MATH)
textbox(s, 0.94, 3.06, 3.40, 1.05,
        [{"text": "D(x_{2})=(0,−x_{2})", "size": 12.8, "after": 2},
         {"text": "D(x_{1})=(0,−x_{1})", "size": 12.8}],
        line_spacing=1.18, latin=MATH, ea=MATH)
textbox(s, 0.94, 4.22, 3.38, 1.18,
        [{"text": "−x_{2}≤−x_{1}，", "size": 12.3, "after": 2},
         {"text": "所以 D(x_{2})⊆D(x_{1})。", "size": 12.3, "color": TEAL_D, "bold": True}],
        line_spacing=1.24, latin=MATH, ea=CN)

card(s, 4.72, 2.10, 3.86, 3.56)
pill(s, 4.94, 2.28, 1.10, 0.30, "② 同正", CREAM, GOLD, 10)
textbox(s, 4.94, 2.68, 3.44, 0.30, ["0<x_{1}≤x_{2}"], size=13.5, color=GOLD, bold=True,
        latin=MATH, ea=MATH)
textbox(s, 4.94, 3.06, 3.40, 0.74,
        ["两集合都含 (0,+∞)，", "只需比较负半轴部分："] , size=11.8, color=INK_2,
        line_spacing=1.24, latin=MATH, ea=CN)
textbox(s, 4.94, 3.84, 3.40, 0.38,
        ["(−∞,−x_{2}] ⊆ (−∞,−x_{1}]"], size=13.0, color=GOLD, bold=True,
        latin=MATH, ea=MATH)
textbox(s, 4.94, 4.48, 3.38, 0.64,
        ["因为 −x_{2}≤−x_{1}，", "所以包含成立。"], size=11.8, color=INK_2,
        line_spacing=1.22, latin=MATH, ea=CN)

card(s, 8.72, 2.10, 3.86, 3.56)
pill(s, 8.94, 2.28, 1.10, 0.30, "③ 异号", CORAL_L, CORAL, 10)
textbox(s, 8.94, 2.68, 3.44, 0.30, ["x_{2}<0<x_{1}"], size=13.5, color=CORAL, bold=True,
        latin=MATH, ea=MATH)
textbox(s, 8.94, 3.06, 3.42, 1.08,
        [{"text": "f(x_{2})>0>f(x_{1})，", "size": 12.4, "after": 2},
         {"text": "前提自动成立。", "size": 12.4}],
        line_spacing=1.18, latin=MATH, ea=CN)
textbox(s, 8.94, 4.06, 3.40, 1.28,
        [{"text": "D(x_{2})=(0,−x_{2})", "size": 11.5, "after": 1},
         {"text": "⊂(0,+∞)", "size": 11.5, "after": 1},
         {"text": "⊂D(x_{1})", "size": 11.5, "color": CORAL, "bold": True}],
        line_spacing=1.18, latin=MATH, ea=MATH)

concl = rect(s, 0.72, 5.94, 11.88, 0.70, TEAL, None, radius=0.08)
fill_text(concl.text_frame,
          ["综上，D(x_{2})⊆D(x_{1})。若 x_{1}<0<x_{2}，则 f(x_{1})>0>f(x_{2})，前提不成立。"],
          size=13.2, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
          latin=MATH, ea=CN, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 6)
notes(s, "三种位置关系不要漏。异号时只有 x2<0<x1 会满足前提；另一种异号情形函数值顺序反了，条件①根本用不上。")

# ================================================================
# 7 part 3 i
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s, BG)
header(s, "PART 03（i）· 反证", "第（3）（i）问：制造一个“该在却不在”的 d")

steps = [
    ("反设", ["f(0)<1。"]),
    ("取点", ["取 y∈(−1,0)，使 f(y)=2^{y}>f(0)；", "令 d=−y/2。"]),
    ("入 D(y)", ["0<d<1，且 y+d=y/2>y，", "故 f(y+d)>f(y)，", "即 d∈D(y)。"]),
    ("出 D(0)", ["由条件②，f(d)<f(0)，", "所以 d∉D(0)。"]),
    ("矛盾", ["由 f(y)>f(0) 及条件①，", "得 D(y)⊆D(0)，", "于是 d∈D(0)，与 d∉D(0) 矛盾。"]),
]
for i, (lab, body) in enumerate(steps):
    yy = 1.76 + i * 1.00
    step_badge(s, 0.90, yy + 0.02, i + 1, TEAL if i < 4 else CORAL)
    textbox(s, 1.38, yy - 0.01, 1.02, 0.30, [lab], size=11, color=GOLD, bold=True,
            latin=CN, ea=CN)
    textbox(s, 2.38, yy - 0.02, 4.70, 0.86, body, size=12.2, color=INK,
            line_spacing=1.18, latin=MATH, ea=CN)

card(s, 7.52, 1.76, 5.06, 4.42, INK, None)
textbox(s, 7.78, 1.98, 4.52, 0.30, ["矛盾的核心"], size=12, color=GOLD, bold=True, latin=CN, ea=CN)
textbox(s, 7.78, 2.40, 4.50, 0.44, ["d ∈ D(y)  ⊆  D(0)"], size=19, color=CREAM, bold=True,
        align=PP_ALIGN.CENTER, latin=MATH, ea=MATH)
textbox(s, 7.78, 2.86, 4.50, 0.30, ["但由 f(d)<f(0) 得到"], size=11.5,
        color=RGBColor(0xBF,0xCD,0xD4), align=PP_ALIGN.CENTER, latin=CN, ea=CN)
textbox(s, 7.78, 3.34, 4.50, 0.44, ["d ∉ D(0)"], size=19, color=RGBColor(0xF0,0xA8,0x95), bold=True,
        align=PP_ALIGN.CENTER, latin=MATH, ea=MATH)
line(s, 8.18, 4.04, 11.92, 4.04, RGBColor(0x48,0x5E,0x6B), 0.8)
textbox(s, 7.78, 4.24, 4.50, 0.32, ["同一元素 d，既在又不在 D(0)。"], size=11.2,
        color=RGBColor(0xD9,0xE2,0xE6), align=PP_ALIGN.CENTER, latin=CN, ea=CN)
textbox(s, 7.78, 5.12, 4.50, 0.52, ["故  f(0) ≥ 1"], size=23.5, color=WHITE, bold=True,
        align=PP_ALIGN.CENTER, latin=MATH, ea=MATH)
footer(s, 7)
notes(s, "关键不是硬算 f(0)，而是让一个具体的 d 同时落入和不落入 D(0)。d=-y/2 既保证在 (0,1) 内能用条件②，又保证 y+d 仍在负半轴。")

# ================================================================
# 8 lemma
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s, BG)
header(s, "PART 03（ii）· 引理", "先准备一个“值序—平移”引理")

card(s, 0.86, 1.72, 11.58, 2.36, TEAL_L, TEAL_L)
pill(s, 1.12, 1.92, 1.18, 0.31, "核心引理", TEAL, WHITE, 10.5)
textbox(s, 1.12, 2.38, 10.98, 0.54,
        ["若  f(a) ≤ f(b) < f(c)，则  f(a) < f(a−b+c)。"],
        size=21, color=INK, bold=True, align=PP_ALIGN.CENTER, latin=MATH, ea=MATH)
textbox(s, 1.12, 3.08, 10.98, 0.64,
        ["把 c−b 看成“从 b 走到 c”的平移量；", "它属于 D(b)，也就属于 D(a)。"],
        size=12.5, color=TEAL_D, align=PP_ALIGN.CENTER, line_spacing=1.18, latin=MATH, ea=CN)

textbox(s, 0.92, 4.32, 0.95, 0.30, ["证明"], size=12, color=GOLD, bold=True, latin=CN, ea=CN)
proof_lines = [
    {"text": "①  f(c)>f(b)  ⇒  c−b∈D(b)。", "size": 13.8, "after": 3},
    {"text": "②  f(a)≤f(b)  ⇒  D(b)⊆D(a)。", "size": 13.8, "after": 3},
    {"text": "③  所以 c−b∈D(a)，即 f(a+c−b)>f(a)。", "size": 13.8},
]
textbox(s, 1.72, 4.28, 8.72, 1.26, proof_lines, line_spacing=1.26)

card(s, 10.72, 4.26, 1.72, 1.28, CORAL_L, CORAL_L)
textbox(s, 10.82, 4.42, 1.52, 0.26, ["方向"], size=9.5, color=CORAL, bold=True,
        align=PP_ALIGN.CENTER, latin=CN, ea=CN)
textbox(s, 10.82, 4.72, 1.52, 0.72, ["值小的一方", "D 更大"], size=9.8, color=INK_2,
        align=PP_ALIGN.CENTER, line_spacing=1.16, latin=CN, ea=CN)

line(s, 0.92, 5.88, 12.28, 5.88, LINE, 0.8)
textbox(s, 0.92, 6.08, 11.4, 0.50,
        ["用法：要证 f(v)>f(u)，只需让 v−u 成为某个 D(u) 中的元素。"],
        size=12.7, color=MUTED, latin=MATH, ea=CN)
footer(s, 8)
notes(s, "引理把“两个函数值的大小”翻译成“一个平移量属于 D(a)”。最后证明递增时，就是让 v-u 充当这个平移量。")

# ================================================================
# 9 claim 1
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s, BG)
header(s, "PART 03（ii）· 排除正值 ①", "先排除 0<x<1 上的正值")

card(s, 0.82, 1.72, 11.68, 4.80)
textbox(s, 1.08, 1.94, 11.10, 0.34, ["反设存在 x∈(0,1)，使 f(x)>0。"], size=13.2, color=CORAL, bold=True,
        latin=MATH, ea=CN)
claims = [
    ("①", "取 m<0，使 f(m)=2^{m}<f(x)。"),
    ("②", "由 f(m)<f(x)，得 D(x)⊆D(m)。"),
    ("③", "因 0<x<1，f(x)<f(0)，故 −x∈D(x)。"),
    ("④", "于是 −x∈D(m)，即 f(m−x)>f(m)。"),
    ("⑤", "但 m−x<m<0，函数值反而更小："),
]
for i, (num, body) in enumerate(claims):
    yy = 2.42 + i * 0.70
    step_badge(s, 1.10, yy + 0.02, i + 1, TEAL if i < 4 else CORAL, 0.30)
    textbox(s, 1.56, yy, 10.40, 0.60, [body], size=13.3, color=INK, line_spacing=1.18,
            latin=MATH, ea=CN)

eq = rect(s, 3.02, 5.82, 7.20, 0.44, CORAL_L, CORAL_L, radius=0.08)
fill_text(eq.text_frame, ["f(m−x)=2^{m−x}<2^{m}=f(m)，矛盾。"], size=12.8, color=CORAL,
          bold=True, align=PP_ALIGN.CENTER, latin=MATH, ea=MATH, anchor=MSO_ANCHOR.MIDDLE)
textbox(s, 1.56, 6.58, 10.40, 0.30, ["所以 0<x<1 时，f(x)≤0。"], size=12.5, color=TEAL_D, bold=True,
        latin=CN, ea=CN)
footer(s, 9)
notes(s, "先用 -x 把 (0,1) 上的正值打掉：f(x)<f(0) 使 -x∈D(x)，但 m-x<m<0 又让 -x 不可能属于 D(m)。")

# ================================================================
# 10 claim 2
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s, BG)
header(s, "PART 03（ii）· 排除正值 ②", "再排除 x≥1 上的正值")

card(s, 0.82, 1.72, 11.68, 4.56)
textbox(s, 1.08, 1.94, 11.10, 0.34, ["反设存在 x≥1，使 f(x)>0。"], size=13.2, color=CORAL, bold=True,
        latin=MATH, ea=CN)
claims2 = [
    ("①", "取 m<0 使 f(m)<f(x)，令 n=m−x+1/2。"),
    ("②", "因 x≥1，n<m<0，故 f(n)<f(m)，从而 D(m)⊆D(n)。"),
    ("③", "又 x−m∈D(m)（因为 m+(x−m)=x，且 f(x)>f(m)）。"),
    ("④", "所以 x−m∈D(n)，即 f(1/2)=f(n+x−m)>f(n)>0。"),
]
for i, (num, body) in enumerate(claims2):
    yy = 2.42 + i * 0.76
    step_badge(s, 1.10, yy + 0.02, i + 1, TEAL if i < 3 else GOLD, 0.30)
    textbox(s, 1.56, yy - 0.02, 10.45, 0.68, [body], size=13.3, color=INK, line_spacing=1.18,
            latin=MATH, ea=CN)

conf = rect(s, 1.34, 5.46, 10.60, 0.64, CORAL_L, CORAL_L, radius=0.08)
fill_text(conf.text_frame,
          ["但 1/2∈(0,1)，由第一步 f(1/2)≤0，与 f(1/2)>0 矛盾。"],
          size=12.6, color=CORAL, bold=True, align=PP_ALIGN.CENTER,
          latin=MATH, ea=CN, anchor=MSO_ANCHOR.MIDDLE)

concl2 = rect(s, 0.82, 6.40, 11.68, 0.42, None, None)
textbox(s, 0.98, 6.38, 11.36, 0.42,
        ["所以  f(x)≤0  对所有 x>0 成立。"], size=14.5, color=TEAL_D, bold=True,
        align=PP_ALIGN.CENTER, latin=MATH, ea=CN)
footer(s, 10)
notes(s, "第二步把 x≥1 的正值逼到 1/2 处：取 n=m-x+1/2，利用 D(m)⊆D(n) 与 x-m∈D(m)，得到 f(1/2)>0，与第一步矛盾。")

# ================================================================
# 11 final
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s, BG)
header(s, "PART 03（ii）· 收口", "最后一步：用引理证明严格递增")

textbox(s, 0.82, 1.70, 6.90, 0.34, ["任取 0<u<v。"], size=14.5, color=INK, bold=True,
        latin=CN, ea=CN)
proof2 = [
    {"text": "① 由上面的结论，f(u)≤0。", "size": 13.5, "after": 4},
    {"text": "② 又 −v<−u<0，且 f(t)=2^{t} 在负半轴递增，", "size": 13.5, "after": 1},
    {"text": "所以 0<f(−v)<f(−u)。", "size": 13.5, "after": 4},
    {"text": "③ 于是 f(u)≤0<f(−v)<f(−u)。", "size": 13.5, "after": 4},
    {"text": "④ 对 (a,b,c)=(u,−v,−u) 使用引理：", "size": 13.5, "after": 2},
]
textbox(s, 0.82, 2.10, 7.02, 2.58, proof2, line_spacing=1.26)

eqbox = rect(s, 0.86, 4.72, 7.10, 0.76, TEAL, None, radius=0.07)
fill_text(eqbox.text_frame, ["f(u) < f(u+v−u) = f(v)"], size=16.6,
          color=WHITE, bold=True, align=PP_ALIGN.CENTER, latin=MATH, ea=MATH,
          anchor=MSO_ANCHOR.MIDDLE)

textbox(s, 0.86, 5.78, 7.10, 0.42,
        ["即对任意 0<u<v，都有 f(u)<f(v)。"], size=13.8, color=INK, bold=True,
        latin=CN, ea=CN)
textbox(s, 0.86, 6.30, 7.10, 0.42,
        ["所以 f(x) 在 (0,+∞) 上严格单调递增。"], size=15.2, color=TEAL_D, bold=True,
        latin=CN, ea=CN)

card(s, 8.34, 1.70, 4.26, 4.86, PAPER, LINE)
textbox(s, 8.58, 1.92, 3.80, 0.28, ["数轴上的位置"], size=10.5, color=MUTED, bold=True,
        latin=CN, ea=CN)
ny = 3.30
line(s, 8.55, ny, 12.35, ny, RGBColor(0xC8,0xBD,0xAA), 1.8)
pts = [(-1.65, "−v", CORAL), (-0.90, "−u", GOLD), (0, "0", MUTED), (0.90, "u", TEAL), (1.65, "v", TEAL)]
for off, lab, col in pts:
    xx = 10.40 + off
    oval(s, xx - 0.055, ny - 0.055, 0.11, 0.11, col, None)
    textbox(s, xx - 0.26, ny + 0.17, 0.52, 0.26, [lab], size=10.8, color=col,
            bold=True, align=PP_ALIGN.CENTER, latin=MATH, ea=MATH)
# negative segment bracket
line(s, 8.75, ny - 0.48, 9.50, ny - 0.48, GOLD, 0.9)
line(s, 8.75, ny - 0.48, 8.75, ny - 0.20, GOLD, 0.9)
line(s, 9.50, ny - 0.48, 9.50, ny - 0.20, GOLD, 0.9)
textbox(s, 8.65, ny - 0.82, 1.00, 0.26, ["负半轴"], size=9.5, color=GOLD, latin=CN, ea=CN)
textbox(s, 8.58, 4.04, 3.78, 1.18,
        ["负半轴：f(x)>0，且随 x 增大；", "正半轴：f(x)≤0。", "u 在正半轴，−v,−u 在负半轴。"],
        size=10.8, color=INK_2, line_spacing=1.38, latin=CN, ea=CN)
textbox(s, 8.58, 5.44, 3.80, 0.82,
        ["关键：把 v−u 变成", "从 −v 到 −u 的差。"],
        size=10.8, color=MUTED, line_spacing=1.30, latin=MATH, ea=CN)

footer(s, 11)
notes(s, "最后一步非常短：u 在正半轴，f(u)≤0；-v 和 -u 在负半轴，函数值一正一更大。引理正好把这两个负半轴点的差转换成 v-u。")

# ================================================================
# 12 review
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s, BG)
header(s, "REVIEW · 复盘", "这道题的方法骨架")

review = [
    ("01", "翻译", ["d∈D(x_{0}) ⇔", "f(x_{0}+d)>f(x_{0})。", "D 是“增长方向”的集合。"], TEAL, TEAL_L),
    ("02", "转化", ["条件①：函数值小", "⇒ D 反而大。", "同值 ⇒ D 相等。"], GOLD, CREAM),
    ("03", "构造", ["反证 + 取 m<0；", "用引理把目标变成", "v−u∈D(u)。"], CORAL, CORAL_L),
]
for i, (num, title, body, accent, light) in enumerate(review):
    x = 0.72 + i * 4.10
    card(s, x, 1.70, 3.82, 2.62)
    textbox(s, x + 0.24, 1.88, 0.82, 0.52, [num], size=22, color=accent, bold=True, latin=CN, ea=CN)
    textbox(s, x + 1.10, 1.95, 2.40, 0.36, [title], size=14.5, color=INK, bold=True, latin=CN, ea=CN)
    textbox(s, x + 0.24, 2.58, 3.32, 1.48, body, size=11.8, color=MUTED, line_spacing=1.40,
            latin=MATH, ea=CN)

card(s, 0.72, 4.60, 7.52, 1.40, INK, None)
textbox(s, 0.96, 4.76, 7.02, 0.26, ["易错提醒"], size=10.5, color=GOLD, bold=True, latin=CN, ea=CN)
textbox(s, 0.96, 5.06, 7.04, 0.80,
        ["① 条件①方向别写反；", "② 端点 0 与 3/2 的取舍；", "③ 正半轴解析式是 −2^{−x}。"],
        size=10.4, color=RGBColor(0xD9,0xE2,0xE6), line_spacing=1.26, latin=MATH, ea=CN)

card(s, 8.48, 4.60, 4.12, 1.40, TEAL_L, TEAL_L)
textbox(s, 8.72, 4.74, 3.66, 0.26, ["答案速览"], size=10.5, color=TEAL_D, bold=True, latin=CN, ea=CN)
textbox(s, 8.72, 5.02, 3.66, 0.84,
        [{"text": "（1）D(−1)=(0,3/2)", "size": 11.4, "after": 2},
         {"text": "（2）D(x_{2})⊆D(x_{1})", "size": 10.6, "after": 1},
         {"text": "（3）(i)(ii) 均得证", "size": 10.6}],
        line_spacing=1.18, latin=MATH, ea=CN)

line(s, 0.72, 6.26, 12.60, 6.26, LINE, 0.8)
textbox(s, 0.72, 6.44, 11.8, 0.40,
        ["通用节奏：先翻译定义，再把条件变成集合关系，最后用反证或构造收口。"],
        size=11.8, color=MUTED, italic=True, latin=CN, ea=CN)
footer(s, 12)
notes(s, "复盘时强调三件事：定义翻译、条件方向、反证构造。比记住某一步计算更重要的是知道为什么每一步会发生。")

import os
out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "第19题_集合D与函数增长_解析.pptx")
prs.save(out)
print("saved", out, "slides:", len(prs.slides))
