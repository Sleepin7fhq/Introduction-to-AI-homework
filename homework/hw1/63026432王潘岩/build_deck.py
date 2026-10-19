# -*- coding: utf-8 -*-
"""Build a polished Chinese math presentation for the 217 problem."""

from math import gcd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

# ----------------------------------------------------------------------------
# Palette
# ----------------------------------------------------------------------------
NAVY       = RGBColor(0x0B, 0x1F, 0x3A)
NAVY_2     = RGBColor(0x12, 0x33, 0x51)
NAVY_SOFT  = RGBColor(0x1B, 0x47, 0x63)
TEAL       = RGBColor(0x0F, 0x8B, 0x84)
TEAL_LIGHT = RGBColor(0xE8, 0xF6, 0xF4)
TEAL_PALE  = RGBColor(0xCF, 0xEA, 0xE7)
GOLD       = RGBColor(0xD8, 0xA8, 0x4E)
GOLD_LIGHT = RGBColor(0xFD, 0xF4, 0xDF)
RED        = RGBColor(0xD9, 0x5B, 0x4B)
RED_LIGHT  = RGBColor(0xFB, 0xEA, 0xE7)
INK        = RGBColor(0x16, 0x2A, 0x40)
GRAY       = RGBColor(0x5E, 0x71, 0x85)
GRAY_2     = RGBColor(0x8A, 0x9A, 0xAA)
BORDER     = RGBColor(0xDC, 0xE6, 0xEF)
BG         = RGBColor(0xF7, 0xF9, 0xFC)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
PALE_BLUE  = RGBColor(0xEC, 0xF3, 0xF9)

CN_FONT = "Microsoft YaHei"
MATH_FONT = "Cambria Math"

W, H = 13.333, 7.5
TOTAL = 8


# ----------------------------------------------------------------------------
# Geometry / low-level helpers
# ----------------------------------------------------------------------------
def rect(slide, x, y, w, h, fill=WHITE, line=None, line_w=0.75,
         shape=MSO_SHAPE.RECTANGLE, radius=None, shadow=False):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    s.shadow.inherit = False
    if shadow:
        add_shadow(s)
    s.text_frame.word_wrap = True
    return s


def add_shadow(shape, blur=14, dist=3, direction=54000, color="B7C7D8", alpha=22):
    """Add a very subtle outer shadow."""
    spPr = shape._element.spPr
    effectLst = spPr.find(qn("a:effectLst"))
    if effectLst is None:
        effectLst = etree.SubElement(spPr, qn("a:effectLst"))
    shdw = etree.SubElement(effectLst, qn("a:outerShdw"))
    shdw.set("blurRad", str(int(blur * 12700)))
    shdw.set("dist", str(int(dist * 12700)))
    shdw.set("dir", str(direction))
    shdw.set("algn", "ctr")
    shdw.set("rotWithShape", "0")
    clr = etree.SubElement(shdw, qn("a:srgbClr"))
    clr.set("val", color)
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(alpha * 1000)))


def line(slide, x, y, w, h, color=BORDER, width=1.0):
    s = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y + h))
    s.line.color.rgb = color
    s.line.width = Pt(width)
    return s


def text(slide, x, y, w, h, content="", size=14, color=INK, bold=False,
         font=CN_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line_spacing=1.05, margins=(0, 0, 0, 0), italic=False, char_space=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(margins[0])
    tf.margin_right = Inches(margins[1])
    tf.margin_top = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    r = p.add_run()
    r.text = content
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    if char_space is not None:
        r.font._rPr.set("spc", str(int(char_space * 100)))
    return box


def rich(slide, x, y, w, h, parts, size=14, color=INK, font=CN_FONT,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.05,
         margins=(0, 0, 0, 0), space_after=0):
    """parts: list of (text, {overrides}) or plain strings."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(margins[0])
    tf.margin_right = Inches(margins[1])
    tf.margin_top = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    p.space_before = Pt(0)
    p.space_after = Pt(space_after)
    for part in parts:
        if isinstance(part, str):
            txt, over = part, {}
        else:
            txt, over = part
        r = p.add_run()
        r.text = txt
        r.font.name = over.get("font", font)
        r.font.size = Pt(over.get("size", size))
        r.font.bold = over.get("bold", False)
        r.font.italic = over.get("italic", False)
        r.font.color.rgb = over.get("color", color)
    return box


def add_paragraphs(slide, x, y, w, h, paras, size=14, color=INK, font=CN_FONT,
                   align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.08,
                   margins=(0, 0, 0, 0), space_after=5):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(margins[0])
    tf.margin_right = Inches(margins[1])
    tf.margin_top = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_before = Pt(0)
        p.space_after = Pt(space_after)
        if isinstance(para, str):
            runs = [(para, {})]
        else:
            runs = para
        for part in runs:
            if isinstance(part, str):
                txt, over = part, {}
            else:
                txt, over = part
            r = p.add_run()
            r.text = txt
            r.font.name = over.get("font", font)
            r.font.size = Pt(over.get("size", size))
            r.font.bold = over.get("bold", False)
            r.font.italic = over.get("italic", False)
            r.font.color.rgb = over.get("color", color)
    return box


def pill(slide, x, y, w, h, label, fill, fg=WHITE, size=11, bold=True,
         line_color=None, radius=0.5):
    s = rect(slide, x, y, w, h, fill=fill, line=line_color,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    tf = s.text_frame
    tf.word_wrap = False
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    r = p.add_run()
    r.text = label
    r.font.name = CN_FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = fg
    return s


def dot(slide, x, y, d, fill, line_color=None):
    return rect(slide, x, y, d, d, fill=fill, line=line_color,
                shape=MSO_SHAPE.OVAL)


def set_cell_border(cell, color=BORDER, width=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for e in tcPr.findall(qn(tag)):
            tcPr.remove(e)
    # Schema order: lnL, lnR, lnT, lnB. Insert at the front in reverse.
    for offset, tag in enumerate(("a:lnB", "a:lnT", "a:lnR", "a:lnL")):
        ln = etree.Element(qn(tag))
        ln.set("w", str(int(width * 12700)))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        fill = etree.SubElement(ln, qn("a:solidFill"))
        clr = etree.SubElement(fill, qn("a:srgbClr"))
        clr.set("val", str(color))
        prst = etree.SubElement(ln, qn("a:prstDash"))
        prst.set("val", "solid")
        tcPr.insert(0, ln)


def style_cell(cell, fill, value, size=12, color=INK, bold=False,
               align=PP_ALIGN.LEFT, margins=(0.08, 0.08, 0.02, 0.02)):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = Inches(margins[0])
    cell.margin_right = Inches(margins[1])
    cell.margin_top = Inches(margins[2])
    cell.margin_bottom = Inches(margins[3])
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    # Clear and write one run (keep the paragraph properties node).
    for child in list(p._p):
        if child.tag != qn("a:pPr"):
            p._p.remove(child)
    r = p.add_run()
    r.text = value
    r.font.name = CN_FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    set_cell_border(cell, BORDER, 0.6)


# ----------------------------------------------------------------------------
# Slide chrome
# ----------------------------------------------------------------------------
def background(slide, color=BG):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def chrome(slide, section, title, subtitle=None, idx=2, total=TOTAL):
    background(slide)
    # Left rail and top accent
    rect(slide, 0, 0, 0.11, H, NAVY)
    rect(slide, 0, 0, 0.11, 1.28, TEAL)
    # Section label
    text(slide, 0.62, 0.28, 0.42, 0.26, f"{idx-1:02d}", size=10.5,
         color=GOLD, bold=True, font=MATH_FONT)
    text(slide, 1.07, 0.28, 7.8, 0.28, section, size=10.5, color=TEAL, bold=True)
    # Title / subtitle
    text(slide, 0.62, 0.55, 12.0, 0.62, title, size=27, color=NAVY, bold=True,
         line_spacing=1.0)
    if subtitle:
        text(slide, 0.63, 1.13, 12.0, 0.32, subtitle, size=12.5, color=GRAY)
    # Thin rule
    rect(slide, 0.62, 1.50, 12.08, 0.012, BORDER)
    rect(slide, 0.62, 1.485, 0.88, 0.045, GOLD)
    # Footer
    text(slide, 0.62, 7.14, 8.0, 0.22,
         "S = {1, 2, …, 280}   ·   5 个两两互素的数", size=8.5, color=GRAY_2)
    text(slide, 11.45, 7.14, 1.25, 0.22, f"{idx:02d} / {total:02d}",
         size=8.5, color=GRAY_2, align=PP_ALIGN.RIGHT, font=MATH_FONT)
    # Progress line
    rect(slide, 0, 7.40, W, 0.035, RGBColor(0xDD, 0xE7, 0xF0))
    rect(slide, 0, 7.40, W * idx / total, 0.035, TEAL)


def notes(slide, content):
    slide.notes_slide.notes_text_frame.text = content


# ----------------------------------------------------------------------------
# Data and verification
# ----------------------------------------------------------------------------
def primes_upto(n):
    sieve = [True] * (n + 1)
    sieve[0] = sieve[1] = False
    for i in range(2, int(n ** 0.5) + 1):
        if sieve[i]:
            sieve[i * i:n + 1:i] = [False] * (((n - i * i) // i) + 1)
    return [i for i, ok in enumerate(sieve) if ok]


PRIMES = primes_upto(280)
assert len(PRIMES) == 59

B_VALUES = [
    [1] + PRIMES,
    [4, 9, 25, 49, 121, 169],
    [2 * 131, 3 * 89, 5 * 53, 7 * 37, 11 * 23, 13 * 19],
    [2 * 127, 3 * 83, 5 * 47, 7 * 31, 11 * 19, 13 * 17],
    [2 * 113, 3 * 79, 5 * 43, 7 * 29, 11 * 17],
    [2 * 109, 3 * 73, 5 * 41, 7 * 23, 11 * 13],
]
B_FLAT = [x for group in B_VALUES for x in group]
assert len(B_FLAT) == 88 and len(set(B_FLAT)) == 88
for group in B_VALUES:
    for i in range(len(group)):
        for j in range(i + 1, len(group)):
            assert gcd(group[i], group[j]) == 1

A = [x for x in range(1, 281) if any(x % p == 0 for p in (2, 3, 5, 7))]
assert len(A) == 216


# ----------------------------------------------------------------------------
# Presentation
# ----------------------------------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
prs.core_properties.title = "5 个两两互素的数：n = 217"
prs.core_properties.subject = "S={1,2,…,280} 中保证 5 个两两互素数的最小 n"
prs.core_properties.author = "解题演示"
prs.core_properties.keywords = "组合数学, 互素, 容斥原理, 抽屉原理"

blank = prs.slide_layouts[6]


# ============================== Slide 1 =====================================
s = prs.slides.add_slide(blank)
background(s, NAVY)

# Decorative geometry
rect(s, 8.50, 0.92, 4.36, 4.36, fill=None, line=NAVY_SOFT, line_w=1.2, shape=MSO_SHAPE.OVAL)
rect(s, 10.18, 2.60, 2.28, 2.28, fill=None, line=GOLD, line_w=0.9, shape=MSO_SHAPE.OVAL)
dot(s, 10.17, 1.21, 0.14, GOLD)
dot(s, 12.68, 3.15, 0.10, TEAL)
dot(s, 8.67, 4.96, 0.09, TEAL)
text(s, 8.26, 1.34, 4.55, 2.72, "217", size=190, color=RGBColor(0x12, 0x3A, 0x54),
     bold=True, font=MATH_FONT, align=PP_ALIGN.CENTER, line_spacing=0.85)
text(s, 9.62, 4.36, 3.15, 0.30, "最 小 阈 值", size=10.5, color=GOLD,
     bold=True, align=PP_ALIGN.CENTER)
rich(s, 9.42, 4.75, 3.55, 0.55,
     [("n", {"font": MATH_FONT, "size": 23, "color": WHITE}),
      (" = ", {"font": MATH_FONT, "size": 23, "color": WHITE}),
      ("217", {"font": MATH_FONT, "size": 27, "bold": True, "color": GOLD})],
     align=PP_ALIGN.CENTER)

# Left title block
rect(s, 0.92, 1.03, 0.62, 0.045, GOLD)
text(s, 1.65, 0.93, 5.6, 0.30, "组 合 数 学   ·   极 值 问 题", size=10.5,
     color=GOLD, bold=True)
text(s, 0.88, 1.52, 7.45, 1.20, "五个两两互素的数", size=44, color=WHITE, bold=True,
     line_spacing=1.0)
text(s, 0.90, 2.78, 7.10, 0.48, "从 280 元集合到最小阈值 217", size=19,
     color=RGBColor(0xA8, 0xC9, 0xD8), font=CN_FONT)
rect(s, 0.92, 3.55, 5.25, 0.018, RGBColor(0x2C, 0x61, 0x78))
text(s, 0.92, 3.76, 6.85, 0.76,
     "下界构造  ×  上界分区  ×  抽屉原理", size=13.0,
     color=RGBColor(0xC9, 0xDC, 0xE6))
text(s, 0.92, 6.58, 8.0, 0.30, "解题演示  ｜  8 页", size=9.5,
     color=RGBColor(0x6E, 0x93, 0xA8))

notes(s, "本题要求最小的 n，使得 S={1,2,…,280} 的任意 n 元子集都含有 5 个两两互素的数。"
         "答案先给出：n=217。下面用“构造下界 + 上界保证”的方式完成证明。")


# ============================== Slide 2 =====================================
s = prs.slides.add_slide(blank)
chrome(s, "0 1   ·   问 题 转 化", "不直接找答案：先求最大“坏子集”",
       "把“每个 n 元子集都成功”翻译成“坏子集最大能有多大”", idx=2)

# Original problem card
rect(s, 0.62, 1.72, 7.42, 2.70, fill=WHITE, line=BORDER, shadow=True)
rect(s, 0.62, 1.72, 0.075, 2.70, GOLD)
text(s, 0.92, 1.92, 3.0, 0.28, "原 题", size=10.5, color=TEAL, bold=True)
add_paragraphs(s, 0.92, 2.32, 6.70, 1.65, [
    [("设 ", {}), ("S = {1, 2, …, 280}", {"font": MATH_FONT, "color": NAVY, "bold": True}),
     ("。求最小的自然数 ", {}), ("n", {"font": MATH_FONT, "color": NAVY, "bold": True}),
     ("，使得 ", {}), ("S", {"font": MATH_FONT, "color": NAVY, "bold": True}),
     (" 的每一个 ", {}), ("n", {"font": MATH_FONT, "color": NAVY, "bold": True}),
     (" 元子集都含有 5 个两两互素的数。", {})],
], size=16, color=INK, line_spacing=1.28, space_after=8)
rect(s, 0.92, 3.18, 6.48, 0.012, BORDER)
text(s, 0.92, 3.34, 6.60, 0.32,
     "两两互素：任意两个数的最大公因数都为 1。", size=11.5, color=GRAY)
text(s, 0.92, 3.73, 6.60, 0.32,
     "例：2、3、5、7 两两互素；6、10、15 不是。", size=10.8, color=GRAY)

# Reformulation card
rect(s, 8.30, 1.72, 4.40, 2.70, fill=NAVY, line=None, shadow=True)
text(s, 8.58, 1.92, 3.6, 0.28, "换 个 说 法", size=10.5, color=GOLD, bold=True)
add_paragraphs(s, 8.58, 2.18, 3.82, 2.04, [
    [("坏子集", {"bold": True, "color": GOLD}), ("：不含 5 个两两互素的数。", {"color": WHITE})],
    [("最大坏子集的大小记为 ", {"color": WHITE}), ("M", {"font": MATH_FONT, "bold": True, "color": GOLD}), ("。", {"color": WHITE})],
    [("则 ", {"color": WHITE}), ("n = M + 1", {"font": MATH_FONT, "bold": True, "color": GOLD}), ("。", {"color": WHITE})],
], size=14.0, line_spacing=1.18, space_after=6, anchor=MSO_ANCHOR.MIDDLE)

# Bottom translation band
rect(s, 0.62, 4.75, 12.08, 1.48, fill=TEAL_LIGHT, line=None)
rect(s, 0.62, 4.75, 0.075, 1.48, TEAL)
text(s, 0.94, 4.96, 2.0, 0.28, "核 心 转 化", size=10.5, color=TEAL, bold=True)
rich(s, 0.94, 5.30, 11.40, 0.52,
     [("求 ", {"size": 17, "color": INK}), ("n", {"size": 20, "bold": True, "color": NAVY, "font": MATH_FONT}),
      ("   ⇔   求最大坏子集的大小 ", {"size": 17, "color": INK}),
      ("M", {"size": 20, "bold": True, "color": NAVY, "font": MATH_FONT}),
      ("   ⇔   证明 ", {"size": 17, "color": INK}),
      ("M = 216", {"size": 21, "bold": True, "color": TEAL, "font": MATH_FONT})],
     line_spacing=1.0)
text(s, 0.94, 5.84, 10.4, 0.26, "先构造一个 216 元的坏子集，再证明 217 元时坏子集不可能存在。",
     size=10.5, color=GRAY)

notes(s, "关键是把“任意 n 元子集都成功”转化为“最大坏子集的大小”。"
         "如果最大坏子集有 M 个元素，那么 n=M+1。因此本题只需证明 M=216："
         "既要构造 216 元的坏子集，也要证明 217 元时一定出现 5 个两两互素的数。")


# ============================== Slide 3 =====================================
s = prs.slides.add_slide(blank)
chrome(s, "0 2   ·   证 明 路 线", "两条边界，在 217 处相遇",
       "下界给出反例，上界给出保证；两者夹出唯一阈值", idx=3)

# Left card
rect(s, 0.62, 1.75, 5.86, 3.62, fill=WHITE, line=BORDER, shadow=True)
rect(s, 0.62, 1.75, 5.86, 0.09, GOLD)
text(s, 0.93, 1.98, 3.4, 0.28, "下 界   ·   构 造 反 例", size=10.5, color=GOLD, bold=True)
text(s, 4.55, 1.95, 1.58, 0.30, "n ≥ 217", size=11.5, color=GOLD, bold=True,
     align=PP_ALIGN.RIGHT, font=MATH_FONT)
rich(s, 0.93, 2.34, 3.4, 0.82,
     [("216", {"size": 46, "bold": True, "color": NAVY, "font": MATH_FONT}),
      ("  元坏子集", {"size": 13, "color": GRAY})], line_spacing=0.95)
add_paragraphs(s, 0.95, 3.30, 5.05, 1.85, [
    [("• 取所有被 2、3、5、7 之一整除的数", {})],
    [("• 5 个数必有两个落在同一个素数类", {})],
    [("• 这两个数不互素，所以构造失败", {})],
], size=12.5, color=INK, line_spacing=1.08, space_after=5)

# Right card
rect(s, 6.84, 1.75, 5.86, 3.62, fill=WHITE, line=BORDER, shadow=True)
rect(s, 6.84, 1.75, 5.86, 0.09, TEAL)
text(s, 7.15, 1.98, 3.4, 0.28, "上 界   ·   抽 屉 原 理", size=10.5, color=TEAL, bold=True)
text(s, 10.77, 1.95, 1.58, 0.30, "n ≤ 217", size=11.5, color=TEAL, bold=True,
     align=PP_ALIGN.RIGHT, font=MATH_FONT)
rich(s, 7.15, 2.34, 3.4, 0.82,
     [("217", {"size": 46, "bold": True, "color": NAVY, "font": MATH_FONT}),
      ("  元必然成功", {"size": 13, "color": GRAY})], line_spacing=0.95)
add_paragraphs(s, 7.17, 3.30, 5.05, 1.85, [
    [("• 把 88 个特殊数分成 6 个互素组", {})],
    [("• 任意 217 元子集至少取到其中 25 个", {})],
    [("• 25 > 6 × 4，必有一组取到 5 个", {})],
], size=12.5, color=INK, line_spacing=1.08, space_after=5)

# Bottom banner
rect(s, 0.62, 5.72, 12.08, 0.92, fill=NAVY, line=None)
rich(s, 0.94, 5.94, 11.40, 0.42,
     [("216", {"size": 19, "bold": True, "color": GOLD, "font": MATH_FONT}),
      (" 只是反例   ｜   ", {"size": 13.5, "color": RGBColor(0xC9, 0xDC, 0xE6)}),
      ("217", {"size": 19, "bold": True, "color": GOLD, "font": MATH_FONT}),
      (" 才是保证   ⇒   ", {"size": 13.5, "color": RGBColor(0xC9, 0xDC, 0xE6)}),
      ("n = 217", {"size": 22, "bold": True, "color": WHITE, "font": MATH_FONT})],
     align=PP_ALIGN.CENTER)

notes(s, "证明分两条线。下界：构造一个 216 元的坏子集，说明 n 不能小于 217。"
         "上界：把 88 个特殊数分成 6 个互素组，任意 217 元子集至少从中取 25 个；"
         "25 个数放进 6 个抽屉，必有一个抽屉至少 5 个，而同一抽屉内两两互素。"
         "于是 n=217。")


# ============================== Slide 4 =====================================
s = prs.slides.add_slide(blank)
chrome(s, "0 3   ·   下 界 构 造", "四个素数类，造出 216 元反例",
       "每个数至少带一个“小素数标签”；5 个数挤进 4 个标签", idx=4)

# Left card with bars
rect(s, 0.62, 1.72, 6.14, 4.62, fill=WHITE, line=BORDER, shadow=True)
text(s, 0.92, 1.94, 5.45, 0.30, "取并集  A = A₂ ∪ A₃ ∪ A₅ ∪ A₇", size=14,
     color=NAVY, bold=True, font=CN_FONT)
text(s, 0.93, 2.31, 5.55, 0.26, "Aₚ 表示 S 中所有 p 的倍数", size=10, color=GRAY)

bar_data = [
    ("2", "2 的倍数", 140, 3.40, TEAL, TEAL_LIGHT),
    ("3", "3 的倍数", 93, 2.26, RGBColor(0x2D, 0x9D, 0xD4), RGBColor(0xE7, 0xF3, 0xFA)),
    ("5", "5 的倍数", 56, 1.36, GOLD, GOLD_LIGHT),
    ("7", "7 的倍数", 40, 0.97, RED, RED_LIGHT),
]
y0 = 2.70
for i, (p, label, count, bw, color, light) in enumerate(bar_data):
    yy = y0 + i * 0.68
    pill(s, 0.94, yy, 0.34, 0.34, p, color, size=11)
    text(s, 1.42, yy + 0.02, 1.28, 0.30, label, size=11.5, color=INK)
    rect(s, 2.73, yy + 0.03, 3.30, 0.28, fill=RGBColor(0xEE, 0xF2, 0xF6), line=None,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    rect(s, 2.73, yy + 0.03, bw, 0.28, fill=color, line=None,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(s, 6.02, yy + 0.01, 0.60, 0.30, str(count), size=13, color=color, bold=True,
         align=PP_ALIGN.RIGHT, font=MATH_FONT)
text(s, 0.94, 5.58, 5.48, 0.50,
     "四个集合有重叠，直接相加会重复计数；下一页用容斥去重。",
     size=10.2, color=GRAY)

# Right card: pigeonhole visual
rect(s, 6.98, 1.72, 5.72, 4.62, fill=NAVY, line=None, shadow=True)
text(s, 7.28, 1.95, 4.9, 0.28, "为 什 么 一 定 不 互 素 ？", size=10.5, color=GOLD, bold=True)
text(s, 7.28, 2.27, 4.95, 0.30, "5 个数，只带 4 种素数标签", size=18, color=WHITE, bold=True)

# Five dots above four bins: the two dots over "3" make the repetition explicit.
dot_specs = [(7.76, TEAL_PALE), (8.78, TEAL_PALE), (9.06, GOLD),
             (10.08, TEAL_PALE), (11.24, TEAL_PALE)]
for dx, dc in dot_specs:
    dot(s, dx, 2.85, 0.25, dc, line_color=TEAL if dc == TEAL_PALE else None)
# Four bins
bins = [
    ("2", 7.30, TEAL),
    ("3", 8.46, TEAL),
    ("5", 9.62, TEAL),
    ("7", 10.78, TEAL),
]
for label, bx, bc in bins:
    rect(s, bx, 3.42, 0.92, 0.55, fill=NAVY_2, line=bc, line_w=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.22)
    text(s, bx, 3.47, 0.92, 0.36, label, size=14, color=bc, bold=True,
         align=PP_ALIGN.CENTER, font=MATH_FONT)
text(s, 7.30, 4.13, 4.85, 0.52, "必有两个点落入同一个素数类。", size=12.5, color=WHITE)

# equation block
rect(s, 7.28, 4.82, 5.10, 0.78, fill=RGBColor(0x16, 0x3D, 0x58), line=None)
rich(s, 7.48, 4.97, 4.78, 0.48,
     [("p | a，p | b", {"font": MATH_FONT, "size": 16, "color": GOLD, "bold": True}),
      ("   ⇒   ", {"font": MATH_FONT, "size": 14, "color": WHITE}),
      ("gcd(a, b) ≥ p > 1", {"font": MATH_FONT, "size": 16, "color": WHITE, "bold": True})],
     align=PP_ALIGN.CENTER, line_spacing=1.0)

notes(s, "令 A=A2∪A3∪A5∪A7。A 中每个数至少被 2、3、5、7 中的一个整除。"
         "任取 5 个数，相当于给它们各贴一个能整除它的小素数标签；只有 4 种标签，"
         "由抽屉原理，必有两个数贴同一种标签，于是它们同被这个素数整除，不互素。"
         "所以 A 是一个不含 5 个两两互素数的坏子集。")


# ============================== Slide 5 =====================================
s = prs.slides.add_slide(blank)
chrome(s, "0 3   ·   下 界 构 造", "容斥计数：|A| = 216",
       "四层修正，把四个倍数集合的重叠恰好去干净", idx=5)

cards = [
    ("单集合之和", "329", "140 + 93 + 56 + 40", TEAL_LIGHT, TEAL, "+"),
    ("两两交之和", "133", "46 + 28 + 20 + 18 + 13 + 8", PALE_BLUE, RGBColor(0x2D, 0x9D, 0xD4), "−"),
    ("三三交之和", "21", "9 + 6 + 4 + 2", GOLD_LIGHT, GOLD, "+"),
    ("四重交集", "1", "只有 210 这一个数", RED_LIGHT, RED, "−"),
]
xs = [0.62, 3.74, 6.86, 9.98]
cw = 2.72
for i, (label, val, sub, light, color, sign) in enumerate(cards):
    x = xs[i]
    rect(s, x, 1.78, cw, 2.28, fill=WHITE, line=BORDER, shadow=True)
    rect(s, x, 1.78, cw, 0.42, fill=light, line=None)
    text(s, x, 1.87, cw, 0.24, label, size=10, color=color, bold=True,
         align=PP_ALIGN.CENTER)
    text(s, x, 2.30, cw, 0.62, val, size=35, color=color, bold=True,
         align=PP_ALIGN.CENTER, font=MATH_FONT)
    text(s, x + 0.12, 3.03, cw - 0.24, 0.78, sub, size=8.8, color=GRAY,
         align=PP_ALIGN.CENTER, line_spacing=1.12)
    if i < 3:
        text(s, x + cw + 0.03, 2.50, 0.30, 0.40, sign, size=20, color=color,
             bold=True, align=PP_ALIGN.CENTER, font=MATH_FONT)

# Result band
rect(s, 0.62, 4.42, 12.08, 1.02, fill=NAVY, line=None, shadow=True)
rich(s, 0.90, 4.64, 11.50, 0.50,
     [("|A|  =  ", {"font": MATH_FONT, "size": 20, "color": WHITE}),
      ("329", {"font": MATH_FONT, "size": 26, "color": GOLD, "bold": True}),
      (" − ", {"font": MATH_FONT, "size": 24, "color": WHITE}),
      ("133", {"font": MATH_FONT, "size": 26, "color": GOLD, "bold": True}),
      (" + ", {"font": MATH_FONT, "size": 24, "color": WHITE}),
      ("21", {"font": MATH_FONT, "size": 26, "color": GOLD, "bold": True}),
      (" − ", {"font": MATH_FONT, "size": 24, "color": WHITE}),
      ("1", {"font": MATH_FONT, "size": 26, "color": GOLD, "bold": True}),
      ("  =  ", {"font": MATH_FONT, "size": 24, "color": WHITE}),
      ("216", {"font": MATH_FONT, "size": 30, "color": TEAL_PALE, "bold": True})],
     align=PP_ALIGN.CENTER, line_spacing=1.0)
text(s, 0.93, 5.62, 11.40, 0.28,
     "每一项的数目，都是 280 除以对应最小公倍数后取整。", size=10, color=GRAY,
     align=PP_ALIGN.CENTER)
rect(s, 0.62, 6.02, 12.08, 0.55, fill=TEAL_LIGHT, line=None)
rich(s, 0.90, 6.15, 11.52, 0.30,
     [("所以存在 216 元的坏子集  ⇒  ", {"size": 14, "color": INK}),
      ("n ≥ 217", {"size": 18, "color": TEAL, "bold": True, "font": MATH_FONT})],
     align=PP_ALIGN.CENTER, line_spacing=1.0)

notes(s, "用容斥原理计算 A 的元素个数。单集合共 140+93+56+40=329；"
         "两两交为 46+28+20+18+13+8=133；三三交为 9+6+4+2=21；"
         "四重交为 1。因此 |A|=329−133+21−1=216。"
         "由于 A 中没有 5 个两两互素的数，n 至少为 217。")


# ============================== Slide 6 =====================================
s = prs.slides.add_slide(blank)
chrome(s, "0 4   ·   上 界 构 造", "设计 88 个数的 6 个“互素抽屉”",
       "每组内部两两互素；B 以外的数全部暂不使用", idx=6)

rows, cols = 7, 3
table_x, table_y = 0.62, 1.70
table_w, table_h = 12.08, 4.46
shape = s.shapes.add_table(rows, cols, Inches(table_x), Inches(table_y),
                           Inches(table_w), Inches(table_h))
table = shape.table
table.columns[0].width = Inches(1.05)
table.columns[1].width = Inches(9.62)
table.columns[2].width = Inches(1.41)
table.rows[0].height = Inches(0.46)
for r in range(1, rows):
    table.rows[r].height = Inches(0.66)

headers = ["组", "元素", "个数"]
for c, htxt in enumerate(headers):
    style_cell(table.cell(0, c), NAVY, htxt, size=11, color=WHITE, bold=True,
               align=PP_ALIGN.CENTER if c != 1 else PP_ALIGN.LEFT,
               margins=(0.10, 0.10, 0.02, 0.02))

row_text = [
    ("B₁", "1，以及不超过 280 的全部素数", "60"),
    ("B₂", "2²，3²，5²，7²，11²，13²", "6"),
    ("B₃", "2·131，3·89，5·53，7·37，11·23，13·19", "6"),
    ("B₄", "2·127，3·83，5·47，7·31，11·19，13·17", "6"),
    ("B₅", "2·113，3·79，5·43，7·29，11·17", "5"),
    ("B₆", "2·109，3·73，5·41，7·23，11·13", "5"),
]
for r, (lab, vals, cnt) in enumerate(row_text, start=1):
    base = WHITE if r % 2 == 0 else RGBColor(0xF9, 0xFB, 0xFD)
    style_cell(table.cell(r, 0), base, lab, size=13, color=TEAL, bold=True,
               align=PP_ALIGN.CENTER)
    style_cell(table.cell(r, 1), base, vals, size=11.5 if r <= 2 else 11.0,
               color=INK, align=PP_ALIGN.LEFT)
    style_cell(table.cell(r, 2), base, cnt, size=14, color=NAVY, bold=True,
               align=PP_ALIGN.CENTER, margins=(0.04, 0.04, 0.02, 0.02))

# Callout
rect(s, 0.62, 6.27, 12.08, 0.58, fill=TEAL_LIGHT, line=None)
rect(s, 0.62, 6.27, 0.065, 0.58, TEAL)
rich(s, 0.88, 6.41, 11.55, 0.30,
     [("✓  ", {"size": 12, "color": TEAL, "bold": True}),
      ("同一组内，各数出现的素因子互不相同  ⇒  任意两数互素。", {"size": 12.5, "color": INK}),
      ("   其中 π(280) = 59，所以 |B₁| = 1 + 59 = 60。", {"size": 10.5, "color": GRAY})],
     line_spacing=1.0)

notes(s, "构造 B：B1 是 1 和所有不超过 280 的素数，共 60 个；B2 是 2、3、5、7、11、13 的平方；"
         "B3 到 B6 是精心挑选的两素数乘积。每一组内所有素因子互不相同，所以同一组内任意两数互素。"
         "B 一共 88 个元素，S 中其余 192 个数不在 B 中。")


# ============================== Slide 7 =====================================
s = prs.slides.add_slide(blank)
chrome(s, "0 5   ·   抽 屉 论 证", "217 元集合，至少 25 个数落在 B 里",
       "6 个抽屉，每个最多容纳 4 个；第 25 个必然溢出", idx=7)

# Proportion bar
bar_x, bar_y, bar_w, bar_h = 0.62, 1.82, 12.08, 0.72
w1 = bar_w * 192 / 280
w2 = bar_w - w1
rect(s, bar_x, bar_y, w1, bar_h, fill=RGBColor(0xE6, 0xEE, 0xF5), line=None)
rect(s, bar_x + w1, bar_y, w2, bar_h, fill=TEAL, line=None)
text(s, bar_x, bar_y + 0.21, w1, 0.30, "B 以外的 192 个数", size=11.5, color=GRAY,
     align=PP_ALIGN.CENTER)
text(s, bar_x + w1, bar_y + 0.21, w2, 0.30, "B：88 个", size=12.5, color=WHITE,
     bold=True, align=PP_ALIGN.CENTER)

# Main inequality
rich(s, 0.64, 2.70, 6.35, 0.62,
     [("|X ∩ B|", {"font": MATH_FONT, "size": 19, "bold": True, "color": NAVY}),
      ("  ≥  217 − 192  =  ", {"font": MATH_FONT, "size": 17, "color": INK}),
      ("25", {"font": MATH_FONT, "size": 25, "bold": True, "color": RED})],
     line_spacing=1.0)
text(s, 7.22, 2.76, 5.42, 0.42,
     "任取 217 元子集 X，其中至少 25 个来自 B。", size=12.5, color=GRAY,
     align=PP_ALIGN.RIGHT)

# Six bins
bin_w, bin_h, gap = 1.86, 1.86, 0.17
start_x = 0.62
bin_y = 3.58
for i in range(6):
    bx = start_x + i * (bin_w + gap)
    is_five = (i == 2)
    rect(s, bx, bin_y, bin_w, bin_h, fill=WHITE, line=BORDER, shadow=True)
    rect(s, bx, bin_y, bin_w, 0.38, fill=TEAL if not is_five else GOLD, line=None)
    text(s, bx, bin_y + 0.08, bin_w, 0.24, f"B{i+1}", size=11, color=WHITE,
         bold=True, align=PP_ALIGN.CENTER, font=MATH_FONT)
    text(s, bx, bin_y + 0.52, bin_w, 0.30,
         "5 个" if is_five else "4 个", size=15 if is_five else 12,
         color=RED if is_five else GRAY, bold=True, align=PP_ALIGN.CENTER)
    # dots: 2 columns x 2 rows, fifth in the fifth slot
    for k in range(5 if is_five else 4):
        col = k % 2
        row = k // 2
        dx = bx + 0.47 + col * 0.54
        dy = bin_y + 0.98 + row * 0.32
        fill = RED if (is_five and k == 4) else TEAL_PALE
        dot(s, dx, dy, 0.17, fill, line_color=None if fill == RED else TEAL)
    if is_five:
        text(s, bx + 0.50, bin_y + 1.58, 1.05, 0.22, "第 25 个", size=8.5,
             color=RED, bold=True, align=PP_ALIGN.CENTER)

# Conclusion
rect(s, 0.62, 5.82, 12.08, 0.88, fill=NAVY, line=None)
rich(s, 0.90, 6.02, 11.52, 0.45,
     [("25 > 6 × 4 = 24", {"font": MATH_FONT, "size": 17, "color": GOLD, "bold": True}),
      ("   ⇒   某一组至少 5 个   ⇒   这 5 个数两两互素", {"size": 15, "color": WHITE, "bold": True})],
     align=PP_ALIGN.CENTER, line_spacing=1.0)
text(s, 0.66, 6.86, 12.0, 0.20, "示意：把 25 个点放进 6 个容量为 4 的抽屉，必有一个抽屉出现第 5 个点。",
     size=8.2, color=GRAY_2, align=PP_ALIGN.CENTER)

notes(s, "设 X 是任意 217 元子集。B 只有 88 个元素，B 外有 192 个元素，"
         "所以 X 中至少有 217−192=25 个元素属于 B。"
         "若每个 B_i 至多贡献 4 个，则 6 组总共至多 24 个，矛盾。"
         "因此某一组至少有 5 个元素属于 X。同一组内任意两数互素，"
         "所以这 5 个数就是所要求的 5 个两两互素的数。")


# ============================== Slide 8 =====================================
s = prs.slides.add_slide(blank)
background(s, NAVY)

# Decorative rings
rect(s, 9.55, 1.35, 3.15, 3.15, fill=None, line=NAVY_SOFT, line_w=1.0, shape=MSO_SHAPE.OVAL)
rect(s, 10.45, 2.25, 1.35, 1.35, fill=None, line=GOLD, line_w=0.8, shape=MSO_SHAPE.OVAL)
dot(s, 12.72, 1.55, 0.10, TEAL)

rect(s, 0.90, 0.94, 0.62, 0.045, GOLD)
text(s, 1.62, 0.86, 4.6, 0.28, "0 6   ·   结 论", size=10.5, color=GOLD, bold=True)

text(s, 0.88, 1.38, 8.0, 0.65, "两条边界，锁定唯一答案", size=30, color=WHITE, bold=True)

# Two summary cards
rect(s, 0.88, 2.28, 4.72, 1.48, fill=NAVY_2, line=RGBColor(0x2C, 0x61, 0x78), line_w=0.8)
text(s, 1.12, 2.46, 2.4, 0.26, "下 界", size=10.5, color=GOLD, bold=True)
rich(s, 1.12, 2.74, 4.15, 0.35,
     [("216", {"font": MATH_FONT, "size": 22, "bold": True, "color": WHITE}),
      (" 元坏子集", {"size": 12, "color": RGBColor(0xC9, 0xDC, 0xE6)})])
text(s, 1.12, 3.19, 3.9, 0.28, "⇒   n ≥ 217", size=13, color=GOLD, bold=True,
     font=MATH_FONT)

rect(s, 5.86, 2.28, 4.72, 1.48, fill=NAVY_2, line=RGBColor(0x2C, 0x61, 0x78), line_w=0.8)
text(s, 6.10, 2.46, 2.4, 0.26, "上 界", size=10.5, color=TEAL_PALE, bold=True)
rich(s, 6.10, 2.74, 4.15, 0.35,
     [("217", {"font": MATH_FONT, "size": 22, "bold": True, "color": WHITE}),
      (" 元必含 5 个", {"size": 12, "color": RGBColor(0xC9, 0xDC, 0xE6)})])
text(s, 6.10, 3.19, 3.9, 0.28, "⇒   n ≤ 217", size=13, color=TEAL_PALE, bold=True,
     font=MATH_FONT)

# Big answer
text(s, 0.88, 4.05, 8.35, 1.30, "n = 217", size=72, color=GOLD, bold=True,
     font=MATH_FONT, line_spacing=0.95)
text(s, 0.93, 5.35, 7.95, 0.40, "最小自然数 n 就是 217。", size=17, color=WHITE)

# Small proof completion badge
rect(s, 9.10, 4.50, 2.95, 0.74, fill=RGBColor(0x13, 0x3A, 0x55), line=None,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
text(s, 9.10, 4.66, 2.95, 0.34, "证 明 完 成", size=12, color=TEAL_PALE, bold=True,
     align=PP_ALIGN.CENTER)

# Footer method line
rect(s, 0.90, 6.25, 11.55, 0.012, RGBColor(0x2C, 0x61, 0x78))
text(s, 0.92, 6.46, 11.55, 0.30,
     "关键工具：素因子分类   ·   容斥原理   ·   六组划分   ·   抽屉原理",
     size=10.5, color=RGBColor(0x8F, 0xB4, 0xC8))

notes(s, "下界：216 元坏子集说明 n≥217。上界：任意 217 元子集都含 5 个两两互素的数，说明 n≤217。"
         "两者合起来得到 n=217。证明完成。")


out = r"D:\ai\deepseekharness\作业\5个两两互素的数_解题演示.pptx"
prs.save(out)
print(f"Saved: {out}")
print(f"Slides: {len(prs.slides)}")
