# -*- coding: utf-8 -*-
"""Build the 6-slide deck on top of the SEU template (东南大学PPT.pptx)."""
import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

WORK = r"E:\AI导论hw2\work"
TEMPLATE = os.path.join(WORK, "template.pptx")
OUT = r"E:\AI导论hw2\AI导论hw2 ppt.pptx"
CROP = os.path.join(WORK, "crop")

GREEN = RGBColor(0x44, 0x54, 0x37)
GREEN_DARK = RGBColor(0x2E, 0x3A, 0x22)
YELLOW = RGBColor(0xFF, 0xCC, 0x00)
INK = RGBColor(0x26, 0x26, 0x26)
GREY = RGBColor(0x86, 0x86, 0x86)
CN = "微软雅黑"
MATH = "Cambria Math"

LAY_COVER, LAY_CONTENT, LAY_PHOTO = 0, 8, 12


# --------------------------------------------------------------------------- #
# text helpers
# --------------------------------------------------------------------------- #
def _set_ea(run, name=CN):
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", name)


def _style(run, name, size, color=INK, bold=False, italic=False, spacing=None):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = name
    _set_ea(run)
    f.color.rgb = color
    if spacing:
        run._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))


def _math_tokens(s):
    out, i = [], 0
    while i < len(s):
        if s.startswith("^{", i):
            j = s.index("}", i)
            out.append(("sup", s[i + 2:j]))
            i = j + 1
        elif s.startswith("_{", i):
            j = s.index("}", i)
            out.append(("sub", s[i + 2:j]))
            i = j + 1
        else:
            out.append(("n", s[i]))
            i += 1
    merged = []
    for kind, txt in out:
        if merged and merged[-1][0] == kind:
            merged[-1][1] += txt
        else:
            merged.append([kind, txt])
    return merged


def M(p, s, size, color=INK, bold=False):
    """math run: latin letters italic (upright e/l/n), ^{..} and _{..} scripts"""
    for kind, txt in _math_tokens(s):
        if kind == "n":
            buf, buf_it = "", None
            for ch in txt:
                it = ch.isalpha() and ch not in "eln"
                if buf and it == buf_it:
                    buf += ch
                    continue
                if buf:
                    r = p.add_run()
                    r.text = buf
                    _style(r, MATH, size, color, bold, buf_it)
                buf, buf_it = ch, it
            if buf:
                r = p.add_run()
                r.text = buf
                _style(r, MATH, size, color, bold, buf_it)
        else:
            r = p.add_run()
            r.text = txt
            _style(r, MATH, size, color, bold, True)
            r._r.get_or_add_rPr().set(
                "baseline", "30000" if kind == "sup" else "-25000")


def C(p, s, size, color=INK, bold=False, spacing=None):
    r = p.add_run()
    r.text = s
    _style(r, CN, size, color, bold, False, spacing)


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def para(tf, first=False, before=0, after=6, line=1.32, indent=0,
         align=PP_ALIGN.LEFT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.line_spacing = line
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    p.alignment = align
    if indent:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(indent))))
        pPr.set("indent", "0")
    return p


def bar(slide, x, y, w, color=YELLOW, h=0.032):
    shp = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def photo(slide, name, x, y, w, h):
    return slide.shapes.add_picture(os.path.join(CROP, name), Inches(x),
                                    Inches(y), Inches(w), Inches(h))


def ph_by_idx(slide, idx):
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.idx == idx:
            return sh
    return None


def drop_ph(slide, keep=()):
    for sh in list(slide.shapes):
        if sh.is_placeholder and sh.placeholder_format.idx not in keep:
            sh._element.getparent().remove(sh._element)


def slide_title(slide, text, size=19):
    sh = ph_by_idx(slide, 11)
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    C(p, text, size, GREEN, True, spacing=0.6)


def add_footer(slide, text="2026 高考数学 · 全国二卷 19 题"):
    tf = textbox(slide, 0.72, 6.80, 6.2, 0.32)
    p = para(tf, first=True, after=0, line=1.0)
    C(p, text, 10.5, GREEN, False, spacing=0.5)


def add_pagenum(slide, n):
    tf = textbox(slide, 9.6, 6.80, 3.0, 0.32)
    p = para(tf, first=True, after=0, line=1.0, align=PP_ALIGN.RIGHT)
    fld = etree.SubElement(p._p, qn("a:fld"))
    fld.set("id", "{B4E5B0D2-3F1A-4C77-9A11-1%02d}" % n)
    fld.set("type", "slidenum")
    rPr = etree.SubElement(fld, qn("a:rPr"))
    rPr.set("lang", "zh-CN")
    rPr.set("sz", "1050")
    rPr.set("b", "0")
    fill = etree.SubElement(rPr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr")).set("val", "445437")
    etree.SubElement(rPr, qn("a:latin")).set("typeface", CN)
    etree.SubElement(fld, qn("a:t")).text = str(n)


def clear_slides(prs):
    lst = prs.slides._sldIdLst
    for sldId in list(lst):
        prs.part.drop_rel(sldId.get(qn("r:id")))
        lst.remove(sldId)


# --------------------------------------------------------------------------- #
prs = Presentation(TEMPLATE)
clear_slides(prs)
layouts = prs.slide_masters[0].slide_layouts

# ============================================================== 1 封面 ======
s1 = prs.slides.add_slide(layouts[LAY_COVER])
p = ph_by_idx(s1, 11).text_frame.paragraphs[0]
C(p, "2026 高考数学 · 全国二卷 · 第 19 题", 15.5, GREY, False, spacing=1.6)

sh = ph_by_idx(s1, 10)
tf = sh.text_frame
tf.clear()
p = tf.paragraphs[0]
p.line_spacing = 1.02
C(p, "导数综合题解析", 40, GREEN_DARK, True, spacing=1.5)

p = ph_by_idx(s1, 12).text_frame.paragraphs[0]
C(p, "切线定参 · 同构变形 · 对称差求最值", 14.5, GREEN, False, spacing=1.2)

tf = textbox(s1, 0.75, 5.52, 6.4, 0.4)
p = para(tf, first=True, after=0, line=1.0)
C(p, "从“切线定参”到“对称差定下界”——一题三问的完整板书", 12, GREY, False, 0.4)

tf = textbox(s1, 0.75, 6.80, 6.2, 0.32)
p = para(tf, first=True, after=0, line=1.0)
C(p, "高三一轮复习 · 函数与导数专题", 10.5, GREEN, False, 0.5)

# ============================================================== 2 原题干 ====
s2 = prs.slides.add_slide(layouts[LAY_CONTENT])
slide_title(s2, "原题呈现")
add_footer(s2)
add_pagenum(s2, 2)

tf = textbox(s2, 0.78, 1.40, 7.8, 1.95)
p = para(tf, first=True, after=0, line=1.55)
C(p, "已知函数 ", 17.5)
M(p, "f(x)=xe^{x}+ax+b", 17.5)
C(p, "，曲线 ", 17.5)
M(p, "y=f(x)", 17.5)
C(p, " 在点 ", 17.5)
M(p, "(0, f(0))", 17.5)
C(p, " 处的切线方程为 ", 17.5)
M(p, "y=−2x+1", 17.5)
C(p, "．", 17.5)

tf = textbox(s2, 0.78, 2.86, 7.8, 2.6)
p = para(tf, first=True, after=15, line=1.4)
C(p, "（1）求 ", 17.5)
M(p, "a，b", 17.5)
C(p, "；", 17.5)

p = para(tf, after=15, line=1.4)
C(p, "（2）当 ", 17.5)
M(p, "x>0", 17.5)
C(p, " 时，", 17.5)
M(p, "f(x+m)−f(x)>m", 17.5)
C(p, "，求 ", 17.5)
M(p, "m", 17.5)
C(p, " 的取值范围；", 17.5)

p = para(tf, after=0, line=1.4)
C(p, "（3）当 ", 17.5)
M(p, "x>0", 17.5)
C(p, " 时，", 17.5)
M(p, "f(k+x)+f(k−x)>2f(k)", 17.5)
C(p, "，求 ", 17.5)
M(p, "k", 17.5)
C(p, " 的最小值．", 17.5)

tf = textbox(s2, 0.78, 6.02, 7.8, 0.4)
p = para(tf, first=True, after=0, line=1.0)
C(p, "考查：导数的几何意义 · 单调性转化 · 对称差中线性项的抵消", 11, GREY, False, 0.4)

photo(s2, "p2.jpg", 9.02, 1.36, 3.5, 4.63)

# ============================================================== 3 （1） ====
s3 = prs.slides.add_slide(layouts[LAY_PHOTO])
ph_by_idx(s3, 12).insert_picture(os.path.join(CROP, "p3.jpg"))
drop_ph(s3, keep=(11, 12))
slide_title(s3, "（1）求 a，b：一个切点，一个斜率")
add_footer(s3)
add_pagenum(s3, 3)

tf = textbox(s3, 7.52, 1.90, 5.05, 0.5)
p = para(tf, first=True, after=0, line=1.2)
C(p, "两个条件，两个未知数", 12, GREY, False, 0.9)

tf = textbox(s3, 7.52, 2.84, 5.05, 1.85)
p = para(tf, first=True, after=13, line=1.4)
C(p, "① 点：", 15, GREEN, True, 0.4)
C(p, "切线在 ", 15)
M(p, "x=0", 15)
C(p, " 处给出 ", 15)
M(p, "y=1", 15)
C(p, "，故 ", 15)
M(p, "f(0)=b=1", 15)
C(p, "．", 15)

p = para(tf, after=13, line=1.4)
C(p, "② 斜率：", 15, GREEN, True, 0.4)
M(p, "f′(x)=(x+1)e^{x}+a", 15)

p = para(tf, after=0, line=1.4)
C(p, "③ 代入：", 15, GREEN, True, 0.4)
M(p, "1+a=−2", 15)
C(p, "，得 ", 15)
M(p, "a=−3", 15)
C(p, "．", 15)

bar(s3, 7.52, 4.86, 1.25)
tf = textbox(s3, 7.52, 4.98, 5.05, 0.85)
p = para(tf, first=True, after=0, line=1.3)
C(p, "所以 ", 15, GREEN)
M(p, "f(x)=xe^{x}−3x+1", 19.5, GREEN_DARK, True)

tf = textbox(s3, 7.52, 6.18, 5.05, 0.5)
p = para(tf, first=True, after=0, line=1.25)
C(p, "先看“点”，再看“斜率”：切点本身也在切线上．", 10.5, GREY, False, 0.3)

# ============================================================== 4 （2） ====
s4 = prs.slides.add_slide(layouts[LAY_CONTENT])
slide_title(s4, "（2）求 m 的取值范围")
add_footer(s4)
add_pagenum(s4, 4)

tf = textbox(s4, 0.78, 1.30, 9.2, 4.15)
p = para(tf, first=True, after=4, line=1.4)
C(p, "① 代入整理　", 16, GREEN, True, 0.4)
M(p, "f(x+m)−f(x)=(x+m)e^{x+m}−xe^{x}−3m>m", 16)
p = para(tf, after=14, line=1.35, indent=0.34)
C(p, "⟺ ", 16)
M(p, "g(x+m)−g(x)>4m", 16)
C(p, "，其中 ", 16)
M(p, "g(t)=te^{t}", 16)
C(p, "．", 16)

p = para(tf, after=14, line=1.4)
C(p, "② 看单调性　", 16, GREEN, True, 0.4)
M(p, "g′(t)=(1+t)e^{t}>0", 16)
C(p, "（", 16)
M(p, "t>0", 16)
C(p, "），故 ", 16)
M(p, "g", 16)
C(p, " 在 ", 16)
M(p, "(0,+∞)", 16)
C(p, " 上递增．", 16)

p = para(tf, after=6, line=1.4)
C(p, "③ 分类讨论　", 16, GREEN, True, 0.4)
C(p, "令 ", 16)
M(p, "F(x)=g(x+m)−g(x)", 16)

p = para(tf, after=4, line=1.38, indent=0.34)
M(p, "m>0", 15)
C(p, "：", 15)
M(p, "F′(x)=g′(x+m)−g′(x)>0", 15)
C(p, "，", 15)
M(p, "F", 15)
C(p, " 递增，下确界 ", 15)
M(p, "g(m)−g(0)=me^{m}", 15)
C(p, "；", 15)

p = para(tf, after=4, line=1.38, indent=0.66)
C(p, "需 ", 15)
M(p, "me^{m}≥4m", 15)
C(p, "，即 ", 15)
M(p, "m≥2ln2", 15)
C(p, "．", 15)

p = para(tf, after=4, line=1.38, indent=0.34)
M(p, "m=0", 15)
C(p, "：", 15)
M(p, "0>0", 15)
C(p, "，不成立；", 15)

p = para(tf, after=0, line=1.38, indent=0.34)
M(p, "m<0", 15)
C(p, "：", 15)
M(p, "x→+∞", 15)
C(p, " 时 ", 15)
M(p, "g(x+m)−g(x)→−∞", 15)
C(p, "，不成立．", 15)

bar(s4, 0.78, 5.62, 1.15)
tf = textbox(s4, 0.78, 5.74, 9.2, 0.85)
p = para(tf, first=True, after=0, line=1.3)
C(p, "④ 结论：", 16, GREEN, True, 0.4)
M(p, "m≥2ln2", 20, GREEN_DARK, True)
C(p, "，即 ", 16)
M(p, "m∈[2ln2, +∞)", 16, GREEN_DARK, True)

photo(s4, "p4.jpg", 10.32, 2.42, 2.34, 2.95)

# ============================================================== 5 （3） ====
s5 = prs.slides.add_slide(layouts[LAY_CONTENT])
slide_title(s5, "（3）求 k 的最小值")
add_footer(s5)
add_pagenum(s5, 5)

photo(s5, "p5.jpg", 0.78, 1.36, 2.34, 2.95)

tf = textbox(s5, 3.36, 1.30, 9.1, 4.2)
p = para(tf, first=True, after=9, line=1.3)
C(p, "线性项在对称差中整体抵消，剩下的只是 ", 14.5)
M(p, "te^{t}", 14.5)
C(p, " 型结构．", 14.5)

p = para(tf, after=4, line=1.38)
C(p, "① 作对称差　", 15.5, GREEN, True, 0.4)
M(p, "f(k+x)+f(k−x)−2f(k)", 15.5)
p = para(tf, after=10, line=1.35, indent=0.34)
C(p, "= ", 15.5)
M(p, "(k+x)e^{k+x}+(k−x)e^{k−x}−2ke^{k}", 15.5)

p = para(tf, after=4, line=1.38)
C(p, "② 提出 ", 15.5, GREEN, True, 0.4)
M(p, "e^{k}", 15.5)
C(p, "　", 15.5)
M(p, "= e^{k}[k(e^{x}+e^{−x}−2)+x(e^{x}−e^{−x})]>0", 15.5)

p = para(tf, after=4, line=1.38)
C(p, "③ 分离参数　", 15.5, GREEN, True, 0.4)
C(p, "对一切 ", 15.5)
M(p, "x>0", 15.5)
C(p, " 有 ", 15.5)
M(p, "k>−x(e^{x}−e^{−x})/(e^{x}+e^{−x}−2)≜−φ(x)", 15)

p = para(tf, after=4, line=1.36, indent=0.34)
C(p, "令 ", 15)
M(p, "ψ(x)=x(e^{x}−e^{−x})−2(e^{x}+e^{−x}−2)", 15)
C(p, "，", 15)
M(p, "ψ(0)=ψ′(0)=0", 15)
C(p, "，", 15)

p = para(tf, after=0, line=1.36, indent=0.34)
M(p, "ψ″(x)=x(e^{x}−e^{−x})>0", 15)
C(p, " ⟹ ", 15)
M(p, "ψ(x)>0", 15)
C(p, "，即 ", 15)
M(p, "φ(x)>2", 15)
C(p, "；又 ", 15)
M(p, "x→0^{+}", 15)
C(p, " 时 ", 15)
M(p, "φ(x)→2", 15)
C(p, "，故 ", 15)
M(p, "k≥−2", 15)
C(p, "．", 15)

bar(s5, 3.36, 5.52, 1.15)
tf = textbox(s5, 3.36, 5.64, 9.1, 0.9)
p = para(tf, first=True, after=0, line=1.3)
C(p, "④ 结论：", 16, GREEN, True, 0.4)
M(p, "k", 18, GREEN_DARK)
C(p, " 的最小值为 ", 16, GREEN)
M(p, "−2", 20, GREEN_DARK, True)
C(p, "（此时对一切 ", 13.5, GREEN)
M(p, "x>0", 13.5, GREEN)
C(p, " 恒成立）", 13.5, GREEN)

# ============================================================== 6 小结 =====
s6 = prs.slides.add_slide(layouts[LAY_CONTENT])
slide_title(s6, "答案与小结")
add_footer(s6)
add_pagenum(s6, 6)

tf = textbox(s6, 0.78, 1.40, 8.0, 2.7)
p = para(tf, first=True, after=26, line=1.2)
C(p, "（1）", 14, GREEN, True, 0.4)
M(p, "a=−3，b=1", 19, GREEN_DARK, True)
C(p, "，", 16)
M(p, "f(x)=xe^{x}−3x+1", 16)

p = para(tf, after=26, line=1.2)
C(p, "（2）", 14, GREEN, True, 0.4)
M(p, "m≥2ln2", 19, GREEN_DARK, True)
C(p, "，即 ", 16)
M(p, "m∈[2ln2, +∞)", 16)

p = para(tf, after=0, line=1.2)
C(p, "（3）", 14, GREEN, True, 0.4)
M(p, "k", 19, GREEN_DARK, True)
C(p, " 的最小值为 ", 16)
M(p, "−2", 19, GREEN_DARK, True)
C(p, "．", 16)

bar(s6, 0.78, 4.30, 11.7, RGBColor(0xCC, 0xCC, 0xBE), 0.020)

tf = textbox(s6, 0.78, 4.52, 8.0, 1.95)
p = para(tf, first=True, after=12, line=1.0)
C(p, "方法提炼", 11, GREY, True, 1.4)

p = para(tf, after=10, line=1.35)
C(p, "切线定参", 15, GREEN, True, 0.5)
C(p, " → ", 15, GREY)
C(p, "同构减元", 15, GREEN, True, 0.5)
C(p, " → ", 15, GREY)
C(p, "单调定界", 15, GREEN, True, 0.5)
C(p, " → ", 15, GREY)
C(p, "对称差消项", 15, GREEN, True, 0.5)

p = para(tf, after=0, line=1.35)
C(p, "三问同源：把“恒成立”化成函数的下确界问题．", 13, INK)

tf = textbox(s6, 0.78, 6.02, 11.7, 0.4)
p = para(tf, first=True, after=0, line=1.15)
C(p, "易错点：b 由切点定、m<0 别漏、分离参数后别忘了取下确界", 12, GREY, False, 0.3)

photo(s6, "p6.jpg", 9.02, 1.42, 3.5, 1.72)

prs.save(OUT)
print("saved", OUT, os.path.getsize(OUT) // 1024, "KB")

