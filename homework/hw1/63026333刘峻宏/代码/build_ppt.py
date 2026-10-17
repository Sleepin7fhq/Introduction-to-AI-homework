# -*- coding: utf-8 -*-
"""
build_ppt.py
============
生成答辩级幻灯片《2024 新课标 I 卷 第 19 题 · 可分数列》。

设计说明
--------
* 画布 16:9（13.333 x 7.5 英寸），纯 python-pptx 绘制，无外部模板。
* 全文源文件、验证脚本均由 DeepSeek Harness (dsh) 在 headless 模式下产出。
* 数学排版：自行实现的极简解析器，支持  _{...} 下标 / ^{...} 上标，
  通过 DrawingML 的 baseline 属性实现真正的上下标（不是拼接 Unicode 字符）。
* 所有插图由 gen_figs.py（matplotlib）生成，按等比缩放嵌入。
"""

import os
import re
import struct

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ------------------------------------------------------------------
# 路径
# ------------------------------------------------------------------
HERE = os.path.dirname(os.path.abspath(__file__))
FIGDIR = os.path.join(HERE, "figures")
OUTDIR = os.path.dirname(HERE)
OUT = os.path.join(OUTDIR, "04_2024新课标I卷第19题_可分数列解答.pptx")

# ------------------------------------------------------------------
# 视觉系统
# ------------------------------------------------------------------
FONT = "Microsoft YaHei UI"
MONO = "Consolas"

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF6, 0xF8, 0xFC)
PANEL = RGBColor(0xED, 0xF2, 0xFA)
LINE = RGBColor(0xD7, 0xE1, 0xEC)

INK = RGBColor(0x13, 0x22, 0x39)
BODY = RGBColor(0x3E, 0x4D, 0x66)
MUTED = RGBColor(0x7B, 0x88, 0x9B)

PRIMARY = RGBColor(0x1B, 0x3A, 0x6B)
BLUE = RGBColor(0x2E, 0x6F, 0xD1)
ACC = RGBColor(0xD9, 0x77, 0x06)
RED = RGBColor(0xD6, 0x45, 0x45)
TEAL = RGBColor(0x0F, 0x8A, 0x80)
PURPLE = RGBColor(0x7C, 0x5C, 0xD6)
GREY = RGBColor(0x9A, 0xA5, 0xB5)

SW, SH = 13.333, 7.5          # 画布宽高（英寸）
ML, MR = 0.78, 0.78           # 左右页边距
CW = SW - ML - MR             # 内容宽度 ≈ 11.77
PAGE = {"n": 0}

FOOT_LEFT = "2024 新课标Ⅰ卷 · 第 19 题  |  (i, j)-可分数列"


# ------------------------------------------------------------------
# 基础工具
# ------------------------------------------------------------------
def png_size(path):
    """读取 PNG 宽高（像素），避免依赖 Pillow。"""
    with open(path, "rb") as f:
        head = f.read(26)
    w, h = struct.unpack(">II", head[16:24])
    return w, h


def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE,
         adj=None, shadow=False):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    if adj is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = adj
    shp.text_frame.word_wrap = True
    return shp


def textbox(slide, x, y, w, h, wrap=True, anchor=MSO_ANCHOR.TOP):
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tf


_TOK = re.compile(r"_\{([^}]*)\}|\^\{([^}]*)\}")


def emit(tf, text, size=14, color=BODY, bold=False, align=PP_ALIGN.LEFT,
         italic=False, line_spacing=1.25, space_after=6, space_before=0,
         font=FONT, first=True):
    """把带 _{} / ^{} 标记的文本写入文本框，可含 \\n 换行。"""
    for i, chunk in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if (first and i == 0) else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        p.space_before = Pt(space_before)
        pieces = _TOK.split(chunk)
        idx = 0
        while idx < len(pieces):
            plain = pieces[idx]
            sub = pieces[idx + 1] if idx + 1 < len(pieces) else None
            sup = pieces[idx + 2] if idx + 2 < len(pieces) else None
            idx += 3
            if plain:
                r = p.add_run()
                r.text = plain
                f = r.font
                f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
                f.color.rgb = color
                f._rPr.set("altLang", "zh-CN")
            for piece, base in ((sub, "-22000"), (sup, "+26000")):
                if piece is None:
                    continue
                r = p.add_run()
                r.text = piece
                f = r.font
                f.name, f.size, f.bold = font, Pt(size * 0.78), bold
                f.color.rgb = color
                f._rPr.set("baseline", base)
                f._rPr.set("altLang", "zh-CN")
    return tf


def txt(slide, x, y, w, h, text, **kw):
    return emit(textbox(slide, x, y, w, h), text, **kw)


def pic(slide, name, x, y, max_w, max_h, bg=True, pad=0.14, align="left"):
    """等比缩放插图，居中于给定可用区域。"""
    path = os.path.join(FIGDIR, name)
    pw, ph = png_size(path)
    ratio = pw / ph
    w = min(max_w, max_h * ratio)
    h = w / ratio
    if align == "left":
        cx = x + (max_w - w) / 2
    elif align == "right":
        cx = x + max_w - w - 0.0
    else:
        cx = x + (max_w - w) / 2
    ox, oy = cx, y + (max_h - h) / 2
    if bg:
        rect(slide, ox - pad, oy - pad, w + 2 * pad, h + 2 * pad,
             fill=WHITE, line=LINE, lw=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.06)
    slide.shapes.add_picture(path, Inches(ox), Inches(oy), Inches(w), Inches(h))


def card(slide, x, y, w, h, title, body, accent=BLUE, tsize=13.5, bsize=11.5,
         fill=WHITE, topbar=True, tcolor=None, bcolor=BODY):
    rect(slide, x, y, w, h, fill=fill, line=LINE, lw=0.75,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.045)
    if topbar:
        rect(slide, x, y, w, 0.055, fill=accent, shape=MSO_SHAPE.RECTANGLE)
    cur = y + 0.30
    if title:
        txt(slide, x + 0.26, cur, w - 0.52, 0.34, title, size=tsize, bold=True,
            color=tcolor or INK, space_after=4)
        cur += 0.40
    if body:
        txt(slide, x + 0.26, cur, w - 0.52, h - (cur - y) - 0.18, body,
            size=bsize, color=bcolor, line_spacing=1.28)


def chip(slide, x, y, w, h, text, size=12, fg=PRIMARY, bg=None, radius=0.32, bold=True):
    bg = bg or PANEL
    rect(slide, x, y, w, h, fill=bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=radius)
    tf = slide.shapes[-1].text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, fg


def keybox(slide, x, y, w, h, text, size=15, fg=WHITE, fill=PRIMARY, radius=0.12):
    """强调结论的实心色块。"""
    rect(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=radius)
    tf = slide.shapes[-1].text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.2)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    for i, chunk in enumerate(text.split("\n")):
        pp = p if i == 0 else tf.add_paragraph()
        pp.alignment = PP_ALIGN.CENTER
        for plain, sub, sup in _chunks(chunk):
            if plain:
                r = pp.add_run()
                r.text = plain
                r.font.name, r.font.size, r.font.bold = FONT, Pt(size), True
                r.font.color.rgb = fg
            for piece, base in ((sub, "-22000"), (sup, "+26000")):
                if piece is None:
                    continue
                r = pp.add_run()
                r.text = piece
                r.font.name, r.font.size, r.font.bold = FONT, Pt(size * 0.76), True
                r.font.color.rgb = fg
                r.font._rPr.set("baseline", base)


def _chunks(chunk):
    pieces = _TOK.split(chunk)
    out = []
    i = 0
    while i < len(pieces):
        out.append((pieces[i],
                    pieces[i + 1] if i + 1 < len(pieces) else None,
                    pieces[i + 2] if i + 2 < len(pieces) else None))
        i += 3
    return out


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def head(slide, kicker, title, sub=None, accent=PRIMARY):
    """统一的页眉：小标签 + 标题 + 强调下划线。"""
    PAGE["n"] += 1
    txt(slide, ML, 0.42, CW, 0.3, kicker, size=11, color=accent, bold=True, space_after=0)
    txt(slide, ML, 0.72, CW, 0.62, title, size=27, color=INK, bold=True, space_after=0)
    rect(slide, ML, 1.42, 0.62, 0.055, fill=accent, shape=MSO_SHAPE.RECTANGLE)
    if sub:
        txt(slide, ML + 0.86, 1.30, CW - 0.86, 0.34, sub, size=12, color=MUTED)
    footer(slide)


def footer(slide):
    n = PAGE["n"]
    rect(slide, ML, SH - 0.60, CW, 0.012, fill=LINE, shape=MSO_SHAPE.RECTANGLE)
    txt(slide, ML, SH - 0.50, CW - 1.0, 0.26, FOOT_LEFT, size=8.5, color=MUTED)
    if n >= 1:
        txt(slide, SW - MR - 0.7, SH - 0.50, 0.7, 0.26, "%02d" % n, size=9,
            color=MUTED, bold=True, align=PP_ALIGN.RIGHT)


def bullet(slide, x, y, w, items, size=13, gap=0.42, dot=True, color=BODY,
           title_color=INK, line_spacing=1.24):
    """要点列表，支持 (小标题, 说明) 二元组。"""
    cy = y
    for it in items:
        if isinstance(it, tuple):
            t, b = it
            if dot:
                rect(slide, x + 0.02, cy + 0.09, 0.085, 0.085, fill=ACC,
                     shape=MSO_SHAPE.OVAL)
            txt(slide, x + 0.26, cy - 0.02, w - 0.26, 0.3, t, size=size, bold=True,
                color=title_color, space_after=1)
            h = emit(textbox(slide, x + 0.26, cy + 0.24, w - 0.26, 0.3), b,
                     size=size - 0.5, color=color, line_spacing=line_spacing)
            cy += 0.30 + gap + 0.02
        else:
            if dot:
                rect(slide, x + 0.04, cy + 0.10, 0.07, 0.07, fill=ACC,
                     shape=MSO_SHAPE.OVAL)
            txt(slide, x + 0.28, cy - 0.03, w - 0.28, 0.3, it, size=size, color=color,
                space_after=0)
            cy += gap
    return cy


# ==================================================================
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)

# ==================================================================
# P1 · 封面
# ==================================================================
s = new_slide(prs)
rect(s, 0, 0, SW, 0.14, fill=PRIMARY, shape=MSO_SHAPE.RECTANGLE)
rect(s, 0, SH - 0.14, SW, 0.14, fill=ACC, shape=MSO_SHAPE.RECTANGLE)
rect(s, SW - 4.55, 0.14, 4.55, SH - 0.28, fill=BG, shape=MSO_SHAPE.RECTANGLE)

txt(s, ML, 1.30, 7.6, 0.34, "DeepSeek Harness (dsh) 求解实录", size=13, color=ACC, bold=True)
txt(s, ML, 1.76, 7.6, 1.5, "可 分 数 列\n结构化新定义题的完整求解", size=40, color=PRIMARY,
    bold=True, line_spacing=1.12, space_after=0)
txt(s, ML, 3.42, 7.7, 0.9,
    "2024 年普通高等学校招生全国统一考试\n数学（新课标 Ⅰ 卷）第 19 题 · 全卷压轴 · 17 分",
    size=15.5, color=BODY, line_spacing=1.32)

rect(s, ML, 4.52, 0.62, 0.055, fill=ACC, shape=MSO_SHAPE.RECTANGLE)
txt(s, ML, 4.86, 7.8, 0.9,
    "题目自称“(i, j)-可分数列”，是一个当场给出、当场考查的新定义。\n"
    "本卷用一条等价翻译成抓手，用模 4 做分工，用两族构造完成计数下界，\n"
    "最后用 Python 回溯穷举把结论在机器里复算了一遍。",
    size=13, color=BODY, line_spacing=1.42)

txt(s, ML, 6.16, 8.0, 0.3, "学号 63026333 · 刘峻宏", size=13.5, color=INK, bold=True)
txt(s, ML, 6.48, 8.0, 0.3, "工具链：dsh(DeepSeek headless) × Python 回溯验证 × matplotlib 程序化制图",
    size=10.5, color=MUTED)

# 右侧信息卡
x0 = SW - 4.05
for i, (t, b, c) in enumerate([
    ("第 (1) 问 · 枚举定位", "把 1…6 中全部四项等差数列列尽，取补集得到 (1,2),(1,6),(5,6)", BLUE),
    ("第 (2) 问 · 显式构造", "删去 2 与 13，给出 m 组下标集合并逐组验证公差", TEAL),
    ("第 (3) 问 · 计数下界", "两族构造合计 m^{2}+m+1 对，推出 P_{m} > 1/8", ACC),
]):
    card(s, x0, 1.62 + i * 1.62, 3.28, 1.42, t, b, accent=c, tsize=13, bsize=10.5)

keybox(s, x0, 6.16, 3.28, 0.86, "结 论\nP_{m} > 1/8 对一切正整数 m 成立", size=13.5,
       fill=PRIMARY, radius=0.14)

# ==================================================================
# P2 · 解题全景
# ==================================================================
s = new_slide(prs)
head(s, "Route Map", "解题全景：从新定义到概率下界",
     "五步走：翻译成下标语言 → 逐问突破 → 机器复核")

steps = [
    ("01", "读懂新定义", "“每组 4 个数成等差”\n翻译成下标的精确覆盖", BLUE),
    ("02", "第 (1) 问", "m = 1，6 项中删 2 项\n穷举 + 补集 → 3 组答案", TEAL),
    ("03", "第 (2) 问", "(2, 13) 的显式 m 组构造\n头 3 组 + 尾部 m−3 块", PURPLE),
    ("04", "第 (3) 问", "模 4 分类 → 两族构造\n合计 m²+m+1 对", ACC),
    ("05", "机器复核", "回溯穷举 m = 1…5\n理论值与实际完全吻合", RED),
]
w = 2.16
gap = 0.16
x0 = ML + 0.02
for i, (no, t, b, c) in enumerate(steps):
    x = x0 + i * (w + gap)
    rect(s, x, 2.16, w, 1.92, fill=WHITE, line=LINE, lw=0.9,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05)
    rect(s, x, 2.16, w, 0.075, fill=c, shape=MSO_SHAPE.RECTANGLE)
    txt(s, x + 0.18, 2.34, w - 0.36, 0.3, no, size=13, color=c, bold=True)
    txt(s, x + 0.18, 2.72, w - 0.36, 0.34, t, size=13.5, color=INK, bold=True)
    txt(s, x + 0.18, 3.14, w - 0.36, 0.8, b, size=10.8, color=BODY, line_spacing=1.28)
    if i < len(steps) - 1:
        rect(s, x + w + 0.02, 3.06, gap - 0.04, 0.03, fill=LINE, shape=MSO_SHAPE.RECTANGLE)

# 主线：一条核心等价
keybox(s, ML + 0.02, 4.42, CW - 0.04, 0.72,
       "贯穿全题的一条等价：a_{t} = a_{1} + (t−1)d 是仿射对应，d ≠ 0 —— "
       "「四个项值成等差」 ⇔ 「四个下标成等差」",
       size=14.5, fill=PRIMARY, radius=0.14)

rows = [
    ("关键难点", "“可不可分”看似是数列值的问题，其实与公差 d 完全无关，只看下标"),
    ("破题抓手", "连续 4 个下标天然成等差 ⇒ 只要把下标切成若干段长度均为 4 的倍数即可"),
    ("为什么还不够", "只用第一族 (1 mod 4, 2 mod 4)，在 m ≥ 3 时 10/91 < 1/8，必须补第二族"),
]
cy = 5.42
for k, v in rows:
    rect(s, ML + 0.02, cy, 1.9, 0.34, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.3)
    txt(s, ML + 0.14, cy + 0.055, 1.66, 0.26, k, size=10.5, color=PRIMARY, bold=True)
    txt(s, ML + 2.06, cy + 0.05, CW - 2.1, 0.28, v, size=12, color=BODY)
    cy += 0.46

# ==================================================================
# P3 · 原题与来源
# ==================================================================
s = new_slide(prs)
head(s, "Problem & Source", "原题 · 出处 · 设定",
     "来源：2024 年普通高等学校招生全国统一考试（新课标 Ⅰ 卷）数学，第 19 题，17 分")

# 左：原题
rect(s, ML, 1.72, 7.62, 4.30, fill=WHITE, line=LINE, lw=1.0,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.035)
txt(s, ML + 0.30, 1.94, 7.0, 0.3, "原 题 全 文", size=12.5, color=PRIMARY, bold=True)
emit(textbox(s, ML + 0.30, 2.34, 7.02, 0.3),
     "设 m 为正整数，数列 a_{1}, a_{2}, …, a_{4m+2} 是公差不为 0 的等差数列。"
     "若从中删去两项 a_{i} 和 a_{j}（i < j）后，剩余的 4m 项可被平均分为 m 组，"
     "且每组的 4 个数都能构成等差数列，则称数列 a_{1}, a_{2}, …, a_{4m+2} 是 (i, j)-可分数列。",
     size=12.3, color=BODY, line_spacing=1.34, space_after=8)
qs = [
    "(1) 写出所有的 (i, j)，1 ≤ i < j ≤ 6，使数列 a_{1}, a_{2}, …, a_{6} 是 (i, j)-可分数列；",
    "(2) 当 m ≥ 3 时，证明：数列 a_{1}, a_{2}, …, a_{4m+2} 是 (2, 13)-可分数列；",
    "(3) 从 1, 2, …, 4m+2 中一次任取两个数 i 和 j（i < j），记数列 a_{1}, a_{2}, …, a_{4m+2} "
    "是 (i, j)-可分数列的概率为 P_{m}，证明：P_{m} > 1/8。",
]
cy = 3.40
for i, q in enumerate(qs):
    txt(s, ML + 0.30, cy, 7.02, 0.74, q, size=12.1, color=INK, line_spacing=1.30)
    cy += 0.80 + (0.10 if i < 2 else 0.0)

# 右：来源卡
xr = ML + 7.86
wr = CW - 7.86
rect(s, xr, 1.72, wr, 4.30, fill=BG, line=LINE, lw=0.85,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.045)
txt(s, xr + 0.26, 1.94, wr - 0.52, 0.3, "题 源 信 息", size=12.5, color=PRIMARY, bold=True)
meta = [
    ("考试名称", "2024 年普通高等学校招生全国统一考试"),
    ("使用卷种", "新课标 Ⅰ 卷（浙江、江苏、山东、广东等省份使用）"),
    ("科目 / 题号", "数学 · 第 19 题（解答题压轴）"),
    ("分值", "17 分，三小问递进"),
    ("试卷时间", "2024 年 6 月 7 日 15:00—17:00"),
    ("题型定位", "新定义结构题：即时概念 + 存在性构造 + 概率计数"),
]
cy = 2.30
for k, v in meta:
    txt(s, xr + 0.26, cy, wr - 0.52, 0.24, k, size=9.6, color=MUTED, bold=True)
    txt(s, xr + 0.26, cy + 0.20, wr - 0.52, 0.42, v, size=10.8, color=INK, line_spacing=1.18)
    cy += 0.60

# 底部：为什么值得
rect(s, ML, 6.16, CW, 0.56, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.14)
txt(s, ML + 0.30, 6.28, CW - 0.6, 0.36,
    "为什么值得单独做一遍：它把「数列 · 集合划分 · 同余分类 · 组合计数 · 概率估计」五个模块串成一条链，"
    "且没有任何现成模型可套——必须先自己造出翻译工具。",
    size=11.4, color=BODY, line_spacing=1.2)

# ==================================================================
# P4 · 命题透视
# ==================================================================
s = new_slide(prs)
head(s, "Anatomy", "命题透视：这道题在考什么",
     "把一道新定义题拆成四个能力点——每一问各占其一")

items = [
    ("① 即时学习与转化", "“4 个数成等差” ⇔ “4 个下标成等差”。\n把数列语言翻译成纯粹的组合语言，"
     "公差 d 从此不再出现。", BLUE),
    ("② 穷举的完备性", "第 (1) 问要「写出所有」，必须先论证 1…6 中的四项等差数列只有 3 条，"
     "再取补集——枚举也要给出上界。", TEAL),
    ("③ 构造性证明", "第 (2) 问不要求计数，只要求「一定能做到」。\n于是给出 m 组显式下标集合即可，"
     "关键在于尾部的块数。", PURPLE),
    ("④ 计数与概率估计", "第 (3) 问只需下界：构造一族一定可分的 (i, j)，数出对数，"
     "与总对数 C(4m+2, 2) 相除。", ACC),
]
w = (CW - 3 * 0.22) / 2
for i, (t, b, c) in enumerate(items):
    x = ML + (i % 2) * (w + 0.22)
    y = 1.78 + (i // 2) * 2.32
    card(s, x, y, w, 2.10, t, b, accent=c, tsize=14, bsize=11.8)

# 底部一行：作答规格
rect(s, ML, 6.28, CW, 0.66, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.1)
txt(s, ML + 0.34, 6.42, CW - 0.68, 0.4,
    "评分视角：第 (1) 问看结论是否穷尽；第 (2) 问看分组是否显式且能覆盖任意 m ≥ 3；"
    "第 (3) 问看是否给出了「可分数对」的构造族并正确计数——只给一族会卡在 m ≥ 3。",
    size=12, color=BODY, line_spacing=1.24)

# ==================================================================
# P5 · 核心引理
# ==================================================================
s = new_slide(prs)
head(s, "Lemma", "核心引理：把题目翻译成中文",
     "这是全题唯一的技术性前提，也是唯一必须写双向的一步")

# 左：叙述
rect(s, ML, 1.74, 4.72, 1.18, fill=WHITE, line=LINE, lw=0.9,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.08)
rect(s, ML, 1.74, 0.055, 1.18, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
txt(s, ML + 0.24, 1.90, 4.3, 0.28, "引 理", size=12.5, color=PRIMARY, bold=True)
emit(textbox(s, ML + 0.24, 2.22, 4.34, 0.3),
     "互异下标 p, q, r, s：\n"
     "a_{p}, a_{q}, a_{r}, a_{s} 重排后成等差  ⇕  p, q, r, s 重排后成等差",
     size=12.4, color=INK, line_spacing=1.34, space_after=2)

txt(s, ML, 3.12, 4.7, 0.3, "证 明 要 点", size=13, color=PRIMARY, bold=True)
bullet(s, ML, 3.48, 4.72, [
    ("⇒ 方向", "若四个值成等差、公差为 D，则 D = d(σ_{k+1} − σ_{k})；因 d ≠ 0，"
     "相邻下标之差恒为 D/d —— 下标本身成等差。"),
    ("⇐ 方向", "若下标重排后公差为 r ≠ 0，则相邻项差恒为 dr —— 值自然成等差。"),
    ("d ≠ 0 用在哪", "它让 t ↦ a_{t} 成为一一映射，并把「等距」按比例双向搬运；"
     "d = 0 时所有项相等，等价失效。"),
], size=12.0, gap=0.80)

# 右：图
pic(s, "fig_lemma.png", ML + 4.98, 1.72, CW - 4.98, 2.1, align="right")

keybox(s, ML + 4.98, 3.86, CW - 4.98, 0.72,
       "结论：从第 (1) 问起，「可不可分」只是 1…4m+2 上的集合划分问题，公差 d 彻底消失。",
       size=12.6, fill=TEAL, radius=0.14)

rect(s, ML + 4.98, 4.78, CW - 4.98, 1.54, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.07)
txt(s, ML + 5.22, 4.98, CW - 5.46, 0.3, "一个立刻能用的推论", size=12.5, color=PRIMARY, bold=True)
txt(s, ML + 5.22, 5.32, CW - 5.46, 0.85,
    "任意 4 个连续下标 k, k+1, k+2, k+3 一定构成四项等差数列。\n"
    "因此只要能把剩余下标切成若干段、每段长度都是 4 的倍数，就自动完成了分组。",
    size=11.8, color=BODY, line_spacing=1.30)

# ==================================================================
# P6 · 第 (1) 问
# ==================================================================
s = new_slide(prs)
head(s, "Part 1", "第 (1) 问：六个下标里的三条等差",
     "m = 1 ⇒ 原数列 6 项；删去 2 项后剩余 4 个下标必须构成四项等差数列")

# 左：推导
txt(s, ML, 1.76, 5.6, 0.3, "推 导", size=13, color=PRIMARY, bold=True)
bullet(s, ML, 2.12, 5.62, [
    ("设四个下标为 x, x+r, x+2r, x+3r", "由 x + 3r ≤ 6 且 r ≥ 1：若 r ≥ 2 则 x + 3r ≥ 7，"
     "矛盾 ⇒ 只能 r = 1，x = 1, 2, 3 —— 全部四项等差数列只有 3 条。"),
    ("取补集得到被删的两项", "{1,2,3,4} → 删 {5,6}；{2,3,4,5} → 删 {1,6}；"
     "{3,4,5,6} → 删 {1,2}。"),
    ("完备性说明", "上面的枚举已穷尽 1…6 中全部四项等差数列，因此没有第四种答案。"),
], size=12.4, gap=0.72)

keybox(s, ML, 5.30, 5.62, 0.86, "答 案\n(i, j) = (1, 2), (1, 6), (5, 6)",
       size=16, fill=PRIMARY, radius=0.12)

# 右：图
pic(s, "fig_q1.png", ML + 5.94, 1.72, CW - 5.94, 3.05, align="right")
rect(s, ML + 5.94, 4.94, CW - 5.94, 1.06, fill=WHITE, line=LINE, lw=0.85,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.12)
txt(s, ML + 6.18, 5.10, CW - 6.42, 0.75,
    "机器复核：m = 1 时回溯穷举 15 对 (i, j)，可分的恰为上述 3 对，与推导一致。",
    size=11.6, color=BODY, line_spacing=1.28)

# ==================================================================
# P7 · 第 (2) 问
# ==================================================================
s = new_slide(prs)
head(s, "Part 2", "第 (2) 问：删去 2 和 13，给出 m 组划分",
     "m ≥ 3 时首项管 14 个数、尾部管其余 —— 关键在于把 14 交给第三组")

txt(s, ML, 1.74, 5.9, 0.3, "显 式 分 组 方 案", size=13, color=PRIMARY, bold=True)
bullet(s, ML, 2.10, 5.92, [
    ("前三组（公差 3）", "{1, 4, 7, 10}  {3, 6, 9, 12}  {5, 8, 11, 14}\n"
     "这三个集合并起来，正好是 1…14 中删去 2 与 13 后剩下的 12 个数。"),
    ("尾部 m − 3 块（公差 1）", "T_{t} = {15+4t, 16+4t, 17+4t, 18+4t}，t = 0, 1, …, m−4\n"
     "覆盖 15, 16, …, 4m+2，共 4(m−3) 个数、m−3 块。"),
    ("组数核对", "3 + (m−3) = m 组，每组 4 个数，共 4m 个数，恰好用尽。"),
], size=12.3, gap=0.74)

rect(s, ML, 5.06, 5.92, 0.92, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.13)
txt(s, ML + 0.28, 5.20, 5.4, 0.3, "易错点：尾部不能从 14 开始", size=12, color=RED, bold=True)
txt(s, ML + 0.28, 5.46, 5.4, 0.4,
    "14 到 4m+2 的长度是 4(m−3)+1，不是 4 的倍数，会多出一个数；"
    "把 14 并入第三组后才干净。",
    size=11.2, color=BODY, line_spacing=1.24)

pic(s, "fig_q2.png", ML + 6.24, 1.70, CW - 6.24, 2.42, align="right")

rect(s, ML + 6.24, 4.28, CW - 6.24, 1.58, fill=WHITE, line=LINE, lw=0.85,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.07)
txt(s, ML + 6.48, 4.46, CW - 6.72, 0.3, "为什么对一切 m ≥ 3 都成立", size=12.5, color=PRIMARY, bold=True)
txt(s, ML + 6.48, 4.80, CW - 6.72, 0.86,
    "前三组与 m 无关，永远合法；尾部从 15 起按「连续 4 个一块」切，长度恒为 4 的倍数。"
    "m = 3 时尾部为空，只剩前三组，仍然成立。",
    size=11.6, color=BODY, line_spacing=1.30)

# ==================================================================
# P8 · 为什么是模 4
# ==================================================================
s = new_slide(prs)
head(s, "Part 3 · Step 1", "第 (3) 问之①：为什么想到模 4",
     "不是拍脑袋看余数，而是「连续 4 个一组」这个动作逼出来的整除条件")

bullet(s, ML, 1.76, 4.9, [
    ("一个朴素但高效的策略", "删去 i < j 后，剩余下标天然断成三段："
     "[1, i−1]、[i+1, j−1]、[j+1, 4m+2]。"),
    ("如果每段都能按连续 4 项分完", "那么三段长度都必须是 4 的倍数，即\n"
     "i − 1 ≡ 0，j − i − 1 ≡ 0，4m + 2 − j ≡ 0  (mod 4)。"),
    ("解出来只有一种余数搭配", "由第一式得 i ≡ 1 (mod 4)；代入第二式得 j ≡ 2 (mod 4)；"
     "第三式给出同一个 j ≡ 2，相容。"),
], size=12.4, gap=0.78)

keybox(s, ML, 4.90, 4.9, 0.78,
       "第一族：(i, j) ≡ (1, 2)  (mod 4)", size=14.5, fill=BLUE, radius=0.14)

pic(s, "fig_mod3.png", ML + 5.22, 1.70, CW - 5.22, 3.1, align="right")

rect(s, ML + 5.22, 5.00, CW - 5.22, 0.68, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.14)
txt(s, ML + 5.46, 5.14, CW - 5.7, 0.42,
    "三段长度能否被 4 整除，正好由 i、j 的模 4 余数控制 —— 模 4 不是巧合，是被整除条件逼出来的。",
    size=11.2, color=BODY, line_spacing=1.22)

# ==================================================================
# P9 · 第一族构造与计数
# ==================================================================
s = new_slide(prs)
head(s, "Part 3 · Step 2", "第 (3) 问之②：第一族的构造与计数",
     "i = 4r+1，j = 4s+2（0 ≤ r ≤ s ≤ m）—— 三段各自是 4 的倍数，秒分组")

# 左：三段
txt(s, ML, 1.78, 4.7, 0.3, "删 去 之 后 的 三 段", size=12.5, color=PRIMARY, bold=True)
segs = [
    ("前段  [1, 4r]", "4r 项", "分成 r 组", BLUE),
    ("中段  [4r+2, 4s+1]", "4(s−r) 项", "分成 s−r 组", ACC),
    ("后段  [4s+3, 4m+2]", "4(m−s) 项", "分成 m−s 组", TEAL),
]
cy = 2.16
for name, num,grp, c in segs:
    rect(s, ML, cy, 4.66, 0.62, fill=WHITE, line=LINE, lw=0.85,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.14)
    rect(s, ML, cy, 0.055, 0.62, fill=c, shape=MSO_SHAPE.RECTANGLE)
    txt(s, ML + 0.24, cy + 0.13, 2.2, 0.36, name, size=11.6, color=INK, bold=True)
    txt(s, ML + 2.46, cy + 0.13, 1.0, 0.36, num, size=11.2, color=c, bold=True)
    txt(s, ML + 3.42, cy + 0.15, 1.1, 0.34, grp, size=11, color=MUTED)
    cy += 0.74

txt(s, ML, 4.50, 4.7, 0.3, "组 数 合 计", size=12.5, color=PRIMARY, bold=True)
txt(s, ML, 4.84, 4.7, 0.3, "r + (s − r) + (m − s) = m 组", size=13.5, color=INK, bold=True)
pic(s, "fig_mod4.png", ML, 5.30, 4.9, 1.35)

# 右：计数
xr = ML + 5.06
wr = CW - 5.06
rect(s, xr, 1.78, wr, 2.30, fill=WHITE, line=LINE, lw=0.9,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05)
txt(s, xr + 0.28, 1.98, wr - 0.56, 0.3, "计 数：这族有多少对？", size=12.5, color=PRIMARY, bold=True)
emit(textbox(s, xr + 0.28, 2.38, wr - 0.56, 0.3),
     "参数范围：0 ≤ r ≤ s ≤ m。\n"
     "对固定的 s，r 可以取 0, 1, …, s，共 s+1 个。\n\n"
     "T_{1} = Σ_{s=0}^{m} (s+1) = (m+1)(m+2)/2",
     size=12.2, color=INK, line_spacing=1.36, space_after=4)

rect(s, xr, 4.22, wr, 1.92, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05)
txt(s, xr + 0.28, 4.42, wr - 0.56, 0.3, "与总对数比较", size=12.5, color=PRIMARY, bold=True)
emit(textbox(s, xr + 0.28, 4.80, wr - 0.56, 0.3),
     "总对数：C(4m+2, 2) = (2m+1)(4m+1)\n\n"
     "P_{m} ≥ T_{1} / C(4m+2, 2) = (m+1)(m+2) / [(4m+2)(4m+1)]",
     size=12.2, color=INK, line_spacing=1.34)

# ==================================================================
# P10 · 第一族的极限
# ==================================================================
s = new_slide(prs)
head(s, "Part 3 · Step 3", "第 (3) 问之③：只靠第一族是不够的",
     "把不等式化成一元二次，正根 ≈ 2.27 —— 也就是说只覆盖到 m = 1, 2")

# 左推导
txt(s, ML, 1.78, 5.5, 0.3, "解 一 元 二 次 不 等 式", size=12.5, color=PRIMARY, bold=True)
emit(textbox(s, ML, 2.16, 5.5, 0.3),
     "要 P_{m} > 1/8，只需 8(m+1)(m+2) > (4m+2)(4m+1)\n\n"
     "8m^{2} + 24m + 16 > 16m^{2} + 12m + 2\n"
     "⇔ 4m^{2} − 6m − 7 < 0\n\n"
     "正根：(3 + √37) / 4 ≈ 2.27  ⇒  m = 1, 2 时成立，m ≥ 3 时失效。",
     size=13.2, color=INK, line_spacing=1.42, space_after=6)

rect(s, ML, 4.60, 5.5, 1.08, fill=WHITE, line=RED, lw=1.1,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.1)
txt(s, ML + 0.28, 4.76, 4.98, 0.3, "m = 3 的反例（也是典型陷阱）", size=12, color=RED, bold=True)
emit(textbox(s, ML + 0.28, 5.10, 4.98, 0.3),
     "T_{1} = 10，总对数 C(14, 2) = 91，10/91 ≈ 0.1099 < 1/8 = 0.125。\n"
     "所以必须再补一族构造，否则第 (3) 问没有真正证完。",
     size=11.6, color=BODY, line_spacing=1.30)

# 右图
pic(s, "fig_fail.png", ML + 5.82, 1.74, CW - 5.82, 3.34, align="right")

txt(s, ML + 5.82, 5.26, CW - 5.82, 0.9,
    "图中「空心圈」标出 m = 3：只用第一族（橙线）已经掉到 1/8 以下；"
    "两族合并（蓝线）则始终远高于阈值。",
    size=11.8, color=BODY, line_spacing=1.30)

# ==================================================================
# P11 · 第二族
# ==================================================================
s = new_slide(prs)
head(s, "Part 3 · Step 4", "第 (3) 问之④：补上对称的第二族",
     "i = 4b+2，j = 4a+1（0 ≤ b < a ≤ m，且 d = a − b ≥ 2）—— 中间段用「错位四带」抢救")

# 左：构造
txt(s, ML, 1.76, 5.6, 0.3, "构 造 要 点", size=12.5, color=PRIMARY, bold=True)
bullet(s, ML, 2.12, 5.62, [
    ("两头照旧", "[1, 4b] 与 [4a+3, 4m+2] 长度仍为 4 的倍数，按连续 4 项分组。"),
    ("中间段出问题", "[4b+1, 4a+2] 长 4d+2，删去两端后剩 4d 个数 —— "
     "不再是简单的连续分块，需要错位。"),
    ("四带错位法", "把 4d 个数按 4d = 4·d 排成 4 带 d 列，每带第 p 个位置合起来成 "
     "公差为 d 的四项等差数列；末尾多出的 4b+4d+2 正好补第一组的缺口。"),
], size=12.2, gap=0.72)

rect(s, ML, 5.06, 5.62, 0.92, fill=WHITE, line=RED, lw=1.0,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.1)
txt(s, ML + 0.28, 5.20, 5.1, 0.3, "必须剔除 d = 1 的那 m 对", size=12, color=RED, bold=True)
emit(textbox(s, ML + 0.28, 5.44, 5.1, 0.3),
     "d = 1 时 j − i = 4d − 1 = 3，中间段没有足够错位空间，构造失效；"
     "这样的对恰有 (a, b) = (1,0), (2,1), …, (m, m−1) 共 m 对。",
     size=11.4, color=BODY, line_spacing=1.28)

# 右：图 + 计数
pic(s, "fig_band.png", ML + 5.94, 1.70, CW - 5.94, 2.26, align="right")
rect(s, ML + 5.94, 4.10, CW - 5.94, 1.80, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05)
txt(s, ML + 6.18, 4.28, CW - 6.42, 0.3, "第 二 族 计 数", size=12.5, color=PRIMARY, bold=True)
emit(textbox(s, ML + 6.18, 4.62, CW - 6.42, 0.3),
     "全部 b < a 的对：m(m+1)/2\n"
     "扣掉 d = 1 的 m 对\n\n"
     "T_{2} = m(m+1)/2 − m = m(m−1)/2",
     size=11.8, color=INK, line_spacing=1.30)

# ==================================================================
# P12 · 收官
# ==================================================================
s = new_slide(prs)
head(s, "Finale", "第 (3) 问之⑤：合并两族，完成不等式",
     "两族余数类型 (1,2) 与 (2,1) 互不相交，可以放心相加")

# 左：合并
cards = [
    ("第一族  T_{1}", "(m+1)(m+2)/2", "(i, j) ≡ (1, 2) (mod 4)", BLUE),
    ("第二族  T_{2}", "m(m−1)/2", "(i, j) ≡ (2, 1) (mod 4)，且 d ≥ 2", ACC),
]
cy = 1.86
for i, (t, v, note, c) in enumerate(cards):
    rect(s, ML, cy, 4.7, 1.28, fill=WHITE, line=LINE, lw=0.9,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.08)
    rect(s, ML, cy, 0.055, 1.28, fill=c, shape=MSO_SHAPE.RECTANGLE)
    txt(s, ML + 0.26, cy + 0.18, 2.4, 0.3, t, size=12.5, color=c, bold=True)
    txt(s, ML + 0.26, cy + 0.58, 4.2, 0.4, v, size=17, color=INK, bold=True)
    txt(s, ML + 0.26, cy + 1.00, 4.3, 0.26, note, size=10.4, color=MUTED)
    cy += 1.44

rect(s, ML, 4.76, 4.7, 0.62, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.16)
emit(textbox(s, ML + 0.26, 4.90, 4.2, 0.3),
     "T_{1} + T_{2} = m^{2} + m + 1",
     size=15, color=PRIMARY, bold=True)

# 右：不等式链条
xr = ML + 5.02
wr = CW - 5.02
rect(s, xr, 1.86, wr, 3.52, fill=WHITE, line=LINE, lw=0.9,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.045)
txt(s, xr + 0.3, 2.06, wr - 0.6, 0.3, "完 成 不 等 式", size=12.5, color=PRIMARY, bold=True)
emit(textbox(s, xr + 0.3, 2.46, wr - 0.6, 0.3),
     "总对数：C(4m+2, 2) = (2m+1)(4m+1)\n\n"
     "P_{m} ≥ (m^{2}+m+1) / [(2m+1)(4m+1)]\n\n"
     "比较分子分母：\n"
     "8(m^{2}+m+1) − (2m+1)(4m+1)\n"
     "= 8m^{2}+8m+8 − (8m^{2}+6m+1)\n"
     "= 2m + 7 > 0",
     size=12.6, color=INK, line_spacing=1.36, space_after=4)

keybox(s, xr, 5.54, wr, 1.06, "因此对任意正整数 m\nP_{m} > 1/8", size=18, fill=PRIMARY, radius=0.1)

rect(s, ML, 5.56, 4.7, 1.04, fill=WHITE, line=LINE, lw=0.85,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.1)
txt(s, ML + 0.26, 5.72, 4.2, 0.75,
    "整道第 (3) 问的量级是：构造出大约 m² 个好对象，"
    "而总对象大约是 8m² 个 —— 恰好越过了 1/8 这条线。",
    size=11.4, color=BODY, line_spacing=1.28)

# ==================================================================
# P13 · 机器验证
# ==================================================================
s = new_slide(prs)
head(s, "Verification", "机器复核：把结论在代码里再算一遍",
     "verify.py（由 dsh 编写）用回溯搜索独立验证 m = 1…5 的每一个 (i, j)")

rows = [
    ("m", "n = 4m+2", "可分对数（穷举）", "m²+m+1（理论）", "总对数", "比例", "≥ 1/8"),
    ("1", "6", "3", "3", "15", "0.2000", "✓"),
    ("2", "10", "7", "7", "45", "0.1556", "✓"),
    ("3", "14", "13", "13", "91", "0.1429", "✓"),
    ("4", "18", "21", "21", "153", "0.1373", "✓"),
    ("5", "22", "31", "31", "231", "0.1342", "✓"),
]
tw = 6.3
cols = [0.45, 0.88, 1.30, 1.32, 0.80, 0.92, 0.60]
x0 = ML
y0 = 1.86
rh = 0.42
for ri, row in enumerate(rows):
    cx = x0
    for ci, val in enumerate(row):
        fill = PRIMARY if ri == 0 else (WHITE if ri % 2 else PANEL)
        rect(s, cx, y0 + ri * rh, cols[ci], rh, fill=fill,
             line=LINE if ri == 0 else WHITE, lw=0.6, shape=MSO_SHAPE.RECTANGLE)
        col = WHITE if ri == 0 else (TEAL if (ci == 6 and ri > 0) else INK)
        txt(s, cx, y0 + ri * rh + 0.09, cols[ci], 0.26, val,
            size=11 if ri else 10.6, color=col, bold=(ri == 0 or ci == 6),
            align=PP_ALIGN.CENTER)
        cx += cols[ci]

txt(s, ML, 4.68, 6.3, 0.3, "其余自动检查（全部通过）", size=12.5, color=PRIMARY, bold=True)
bullet(s, ML, 5.02, 6.3, [
    "m = 1 的穷举结果为 [(1,2), (1,6), (5,6)]，与第 (1) 问答案逐条吻合。",
    "m = 3 时 (2, 13) 可分，回溯找到的分组正是 {1,4,7,10}、{3,6,9,12}、{5,8,11,14}。",
    "反过来也说明：m = 1…5 上「可分数对」恰好就是两族构造给出的那些——下界其实是精确的。",
], size=11.6, gap=0.40)

pic(s, "fig_data.png", ML + 6.66, 1.72, CW - 6.66, 3.20, align="right")
rect(s, ML + 6.66, 5.06, CW - 6.66, 1.42, fill=WHITE, line=LINE, lw=0.85,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.09)
txt(s, ML + 6.88, 5.22, CW - 7.12, 0.3, "橙点与蓝柱顶端完全重合", size=11.8, color=PRIMARY, bold=True)
txt(s, ML + 6.88, 5.54, CW - 7.12, 0.8,
    "橙点是理论值 m²+m+1，蓝柱是穷举得到的真实可分对数——两者在 m = 1…5 上无一例外地相等。",
    size=11, color=BODY, line_spacing=1.28)

# ==================================================================
# P14 · 复盘
# ==================================================================
s = new_slide(prs)
head(s, "Retrospective", "复盘：可复用的三步套路与三个坑",
     "把这套做法命名为「模 4 分段 — 构造族计数法」")

txt(s, ML, 1.76, 6.0, 0.3, "三 步 套 路", size=12.5, color=PRIMARY, bold=True)
methods = [
    ("01 翻译", "用仿射对应把「值的问题」变成「下标的集合划分问题」，公差 d 消失。"),
    ("02 找不变量", "「连续 4 项一组」⇔ 每段长度被 4 整除 ⇒ i ≡ 1、j ≡ 2 (mod 4)。"),
    ("03 构造+计数", "构造一族一定可分的下标对，数出对数，除以总对数得到概率下界。"),
]
cy = 2.14
for i, (t, b) in enumerate(methods):
    chip(s, ML, cy, 1.46, 0.34, t, size=10.6, fg=WHITE,
         bg=[BLUE, TEAL, ACC][i], bold=True)
    txt(s, ML + 1.64, cy + 0.03, 4.36, 0.7, b, size=11.4, color=BODY, line_spacing=1.26)
    cy += 0.70

txt(s, ML, 4.42, 6.0, 0.3, "三 个 真 实 易 错 点", size=12.5, color=RED, bold=True)
bullet(s, ML, 4.80, 6.0, [
    ("尾部块数差一", "第 (2) 问尾部若从 14 开始切，长度 4(m−3)+1 不是 4 的倍数。"),
    ("以为一族就够", "只靠第一族在 m ≥ 3 时已落在 1/8 之下——最典型的「看似证完」。"),
    ("忘了 d = 1 的洞", "第二族不扣除 j − i = 3 的那 m 对，计数会虚增。"),
], size=11.8, gap=0.42, title_color=RED)

xr = ML + 6.5
wr = CW - 6.5
rect(s, xr, 1.76, wr, 2.44, fill=WHITE, line=LINE, lw=0.9,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.045)
txt(s, xr + 0.26, 1.94, wr - 0.52, 0.3, "一 句 话 本 质", size=12.5, color=PRIMARY, bold=True)
txt(s, xr + 0.26, 2.34, wr - 0.52, 1.6,
    "删去两个下标后，剩余下标能否被\n四项等差数列精确覆盖——\n与公差无关，只与下标有关。",
    size=14.5, color=INK, bold=True, line_spacing=1.36)

rect(s, xr, 4.34, wr, 2.40, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.045)
txt(s, xr + 0.26, 4.52, wr - 0.52, 0.3, "一 个 自 然 追 问", size=12.5, color=PRIMARY, bold=True)
emit(textbox(s, xr + 0.26, 4.90, wr - 0.52, 0.3),
     "若把「4 项一组」改成「k 项一组」？\n"
     "表层答案：模 4 换成模 k，i ≡ 1、j ≡ 2 (mod k)，对称族排除 j − i = k − 1。\n"
     "真正的答案还要再分层：k 项等差数列 x, x+r, …, x+(k−1)r 的余数分布取决于 "
     "gcd(r, k) —— 应改为「模 k 分层 + 用 gcd(r, k) 区分公差类型」。",
     size=10.8, color=BODY, line_spacing=1.30)

# ==================================================================
# P15 · 工具链
# ==================================================================
s = new_slide(prs)
head(s, "Toolchain", "工具链与文件清单",
     "全程由 DeepSeek Harness 在 headless 模式下完成推理，Python 负责验证与制图")

items = [
    ("DeepSeek Harness (dsh)", "headless profile，一次调用完成：解三问 → 写 solution.md → "
     "写 verify.py → 运行验证 → 写 analysis.md。耗时约 830 秒。", BLUE),
    ("Python 回溯验证 verify.py", "仅用标准库。剪枝搜索判定任意 (m, i, j) 是否可分，"
     "并穷举 m = 1…5 的全部下标对。", TEAL),
    ("matplotlib 制图 gen_figs.py", "本卷 8 张配图全部由脚本生成，中文字体 Microsoft YaHei，"
     "尺寸与配色统一，可重复复现。", ACC),
    ("python-pptx 排版 build_ppt.py", "幻灯片本体也由代码生成：真正的 DrawingML 上下标，"
     "改一处配色即可全局生效。", PURPLE),
]
w = (CW - 0.26) / 2
for i, (t, b, c) in enumerate(items):
    x = ML + (i % 2) * (w + 0.26)
    y = 1.82 + (i // 2) * 1.72
    card(s, x, y, w, 1.54, t, b, accent=c, tsize=13, bsize=11)

rect(s, ML, 5.14, CW, 1.24, fill=WHITE, line=LINE, lw=0.9,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.08)
txt(s, ML + 0.28, 5.30, CW - 0.56, 0.3, "调 用 方 式（可复现）", size=12, color=PRIMARY, bold=True)
txt(s, ML + 0.28, 5.62, CW - 0.56, 0.7,
    "$ node @deepseek-ai/dsh/lib/bin.js --profile headless \"$(cat prompt.txt)\"\n"
    "$ python verify.py        # 回溯穷举复核\n"
    "$ python build_ppt.py     # 生成本幻灯片",
    size=10.5, color=BODY, line_spacing=1.24, font=MONO)

txt(s, ML, 6.52, CW, 0.3, "学号 63026333 · 刘峻宏   |   原题来源：2024 年新课标 Ⅰ 卷数学第 19 题",
    size=11, color=MUTED)

# ==================================================================
prs.save(OUT)
print("SAVED:", OUT)
print("SLIDES:", len(prs.slides.__iter__.__self__._sldIdLst))
