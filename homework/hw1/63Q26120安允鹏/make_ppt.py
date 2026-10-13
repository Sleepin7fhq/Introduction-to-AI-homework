#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成《2025 数学压轴大题》讲解 PPT —— 共 9 页。

输出：ai_agent_demo_数学压轴v2_sample.pptx（不覆盖已有文件）

做法：
  * 直接以用户提供的《东南大学PPT.pptx》为底板，保留它的母版 / 版式 / 主题配色 /
    背景照片 / 校徽 / 页码，只删掉原有的 91 页，再按同一套版式重建 9 页。
  * 不在页面上另加卡片、圆角矩形、装饰圆，正文只用文字与一张函数图象。
  * 封面用 "标题幻灯片" 版式，内容页用 "仅标题（无主副标题）" 版式，
    末页用 "末尾幻灯片" 版式，列出本题用到的数学思想与方法。

运行：  .venv/bin/python make_ppt.py
"""

import os
import re

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
os.environ.setdefault("MPLCONFIGDIR", os.path.join(BASE_DIR, "_mplcache"))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

TEMPLATE = os.path.join(BASE_DIR, "template_src.pptx")
OUTPUT = os.path.join(BASE_DIR, "ai_agent_demo_数学压轴v2_sample.pptx")
FIG = os.path.join(BASE_DIR, "_probe", "fig_f1.png")

# ---- 取自模板主题的配色 -------------------------------------------------
GREEN = RGBColor(0x44, 0x54, 0x37)   # accent1：模板的深绿
INK = RGBColor(0x26, 0x26, 0x26)     # 正文
GREY = RGBColor(0x76, 0x76, 0x76)    # 提示性小字
SOFT = RGBColor(0x50, 0x50, 0x50)    # 次级正文

LAYOUT_TITLE = 0     # 标题幻灯片
LAYOUT_BODY = 8      # 仅标题（无主副标题）
LAYOUT_END = 15      # 末尾幻灯片


# ======================================================================
# 生成函数图象
# ======================================================================
def build_figure():
    import numpy as np
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    os.makedirs(os.path.dirname(FIG), exist_ok=True)

    x = np.linspace(0.0, np.pi / 4, 400)
    y = 5 * np.cos(x) - np.cos(5 * x)
    xm, ym = np.pi / 6, 3 * np.sqrt(3)

    fig, ax = plt.subplots(figsize=(5.75, 4.30), dpi=220)
    ax.plot(x, y, color="#445437", lw=2.6, solid_capstyle="round")
    ax.axvline(xm, color="#B8912B", lw=1.1, ls=(0, (4, 3)))
    ax.plot([xm], [ym], "o", color="#B8912B", ms=6.5, zorder=5)
    ax.annotate(r"$(\frac{\pi}{6},\ 3\sqrt{3})$",
                xy=(xm, ym), xytext=(xm + 0.10, ym - 0.62),
                fontsize=12.5, color="#3A3A3A")

    ax.set_xlim(0, np.pi / 4 + 0.05)
    ax.set_ylim(3.75, 5.55)
    ax.set_xticks([0, np.pi / 12, np.pi / 6, np.pi / 4])
    ax.set_xticklabels(["0", r"$\frac{\pi}{12}$", r"$\frac{\pi}{6}$", r"$\frac{\pi}{4}$"],
                       fontsize=11.5)
    ax.set_yticks([4, 3 * np.sqrt(3)])
    ax.set_yticklabels(["4", r"$3\sqrt{3}$"], fontsize=11.5)
    ax.tick_params(colors="#5A5A5A", length=3)

    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color("#9A9A9A")
        ax.spines[side].set_linewidth(0.9)

    ax.grid(axis="y", color="#E2E2E2", lw=0.8)
    ax.set_axisbelow(True)
    ax.set_xlabel("x", fontsize=12.5, color="#3A3A3A", labelpad=2)
    ax.set_ylabel("f(x)", fontsize=12.5, color="#3A3A3A", labelpad=2)

    fig.tight_layout(pad=0.35)
    fig.savefig(FIG, transparent=True)
    plt.close(fig)
    print("figure ->", FIG)


# ======================================================================
# 文字工具
# ======================================================================
SUP_RE = re.compile(r"\^([^^]+)\^")


def _segments(text, base):
    """把 'cos^3^x' 解析成 [(文本, 样式), ...]，^...^ 内部为上标。"""
    out, pos = [], 0
    for m in SUP_RE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], dict(base)))
        sup = dict(base)
        sup["sup"] = True
        out.append((m.group(1), sup))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], dict(base)))
    return out


def add_para(tf, text, size=19, color=INK, bold=False,
             space_before=0, space_after=12, line_spacing=1.28,
             align=PP_ALIGN.LEFT, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line_spacing
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)

    base = {"size": size, "color": color, "bold": bold}
    for seg, opt in _segments(text, base):
        run = p.add_run()
        run.text = seg
        f = run.font
        s = opt.get("size", size)
        if opt.get("sup"):
            s = max(9.0, round(s * 0.62, 1))
            f._rPr.set("baseline", "30000")
        f.size = Pt(s)
        f.bold = opt.get("bold", bold)
        f.color.rgb = opt.get("color", color)
    return p


def body_box(slide, left=0.90, top=1.22, width=11.50, height=5.20):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    # 关掉自动缩放，避免 PowerPoint 打开时改动排版
    body_pr = tf._txBody.find(qn("a:bodyPr"))
    if body_pr is not None:
        for tag in ("a:normAutofit", "a:spAutoFit"):
            el = body_pr.find(qn(tag))
            if el is not None:
                body_pr.remove(el)
    return tf


def ph_by_idx(slide, idx):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    raise KeyError(idx)


def set_placeholder_box(ph, left, top, width, height):
    """显式写死占位符的 off/ext（直接改 XML，避免 setter 生成 (0,0) 的 off）。"""
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls
    sp_pr = ph._element.spPr
    xfrm = sp_pr.find(qn("a:xfrm"))
    if xfrm is not None:
        sp_pr.remove(xfrm)
    xfrm = parse_xml(
        '<a:xfrm %s><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
        % (nsdecls("a"), Inches(left), Inches(top), Inches(width), Inches(height)))
    sp_pr.insert(0, xfrm)


def fill_placeholder(slide, idx, items):
    tf = ph_by_idx(slide, idx).text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        add_para(tf, it.get("text", ""),
                 size=it.get("size", 18),
                 color=it.get("color", INK),
                 bold=it.get("bold", False),
                 space_before=it.get("space_before", 0),
                 space_after=it.get("space_after", 6),
                 line_spacing=it.get("line_spacing", 1.2),
                 first=first)
        first = False


def set_header(slide, text):
    tf = ph_by_idx(slide, 11).text_frame
    tf.word_wrap = True
    add_para(tf, text, size=20, color=GREEN, bold=True,
             space_after=0, line_spacing=1.0, first=True)


# ======================================================================
# 删除模板原有的全部幻灯片（保留母版 / 版式 / 主题）
# ======================================================================
def strip_slides(prs):
    sld_id_lst = prs.slides._sldIdLst
    n = len(list(sld_id_lst))
    for sld_id in list(sld_id_lst):
        prs.part.drop_rel(sld_id.get(qn("r:id")))
        sld_id_lst.remove(sld_id)
    print("removed %d template slides" % n)


def prune_template(prs, keep_layouts, keep_masters=1):
    """删掉模板里用不到的母版与版式（它们带着约 5MB 从不显示的大图）。"""
    pres_el = prs.part._element
    masters = list(prs.slide_masters)

    master_lst = pres_el.find(qn("p:sldMasterIdLst"))
    dropped_m = 0
    for sld_master_id in list(master_lst):
        r_id = sld_master_id.get(qn("r:id"))
        if prs.part.rels[r_id].target_part is masters[0].part:
            continue
        prs.part.drop_rel(r_id)
        master_lst.remove(sld_master_id)
        dropped_m += 1

    master = masters[0]
    lay_lst = master.element.find(qn("p:sldLayoutIdLst"))
    dropped_l = 0
    for sld_layout_id in list(lay_lst):
        r_id = sld_layout_id.get(qn("r:id"))
        if str(master.part.rels[r_id].target_part.partname) in keep_layouts:
            continue
        master.part.drop_rel(r_id)
        lay_lst.remove(sld_layout_id)
        dropped_l += 1
    print("pruned %d master(s), %d unused layout(s)" % (dropped_m, dropped_l))


# ======================================================================
# 各页内容
# ======================================================================
def slide_cover(prs, layouts):
    s = prs.slides.add_slide(layouts[LAYOUT_TITLE])
    fill_placeholder(s, 11, [{"text": "高考数学专题 · 新课标Ⅰ卷", "size": 20,
                              "color": SOFT, "space_after": 0}])
    fill_placeholder(s, 10, [{"text": "2025 数学压轴大题", "size": 40,
                              "color": GREEN, "bold": True, "space_after": 0}])
    fill_placeholder(s, 12, [{"text": "第 19 题（17 分） · 用高中知识完整解答", "size": 13,
                              "color": GREY, "space_after": 0}])
    return s


def slide_problem(prs, layouts):
    s = prs.slides.add_slide(layouts[LAYOUT_BODY])
    set_header(s, "题目呈现")
    tf = body_box(s)
    add_para(tf, "2025 年新课标Ⅰ卷 · 第 19 题（解答题，17 分）",
             size=14, color=GREY, space_after=16, first=True)
    add_para(tf, "设函数 f(x) = 5cos x − cos 5x。", size=21, bold=True,
             color=INK, space_after=18)
    add_para(tf, "(1) 求 f(x) 在区间 [0, π/4] 上的最大值；", size=19, space_after=14)
    add_para(tf, "(2) 给定 θ∈(0, π) 和 a∈R，证明：存在 y∈[a−θ, a+θ]，"
                 "使得 cos y ≤ cos θ；", size=19, space_after=14)
    add_para(tf, "(3) 设 b∈R。若存在 φ∈R，使得 5cos x − cos(5x+φ) ≤ b "
                 "对任意 x∈R 恒成立，求 b 的最小值。", size=19, space_after=14)
    return s


def slide_p1_derivative(prs, layouts):
    s = prs.slides.add_slide(layouts[LAYOUT_BODY])
    set_header(s, "第(1)问：先求导，看单调性")
    tf = body_box(s, left=0.90, top=1.28, width=5.45, height=5.10)
    add_para(tf, "思路：求导 → 定单调 → 找最大值", size=13, color=GREY,
             space_after=12, first=True)
    add_para(tf, "f′(x) = −5sin x + 5sin 5x", size=18, space_after=2)
    add_para(tf, "　　　= 10 cos 3x · sin 2x", size=20, bold=True, color=GREEN,
             space_after=14)
    add_para(tf, "x∈(0, π/4) 时 sin 2x > 0，符号只由 cos 3x 决定：",
             size=17, space_after=10)
    add_para(tf, "0 < x < π/6 时 f 递增；π/6 < x < π/4 时 f 递减。",
             size=17, space_after=14)
    add_para(tf, "最大值 f(π/6) = 5cos(π/6) − cos(5π/6) = 3√3",
             size=19, bold=True, color=GREEN, space_after=12)
    add_para(tf, "端点：f(0) = 4，f(π/4) = 3√2，都小于 3√3。",
             size=13, color=GREY, space_after=0)

    # 右侧配一张函数图象
    s.shapes.add_picture(FIG, Inches(6.52), Inches(1.32),
                         width=Inches(5.92), height=Inches(4.43))
    return s


def slide_p1_identity(prs, layouts):
    s = prs.slides.add_slide(layouts[LAYOUT_BODY])
    set_header(s, "第(1)问：换个角度，得到全局上界")
    tf = body_box(s)
    add_para(tf, "思路：把 5cos x − cos 5x 化成 cos x 的多项式", size=13,
             color=GREY, space_after=16, first=True)
    add_para(tf, "由五倍角公式　cos 5x = 16cos^5^x − 20cos^3^x + 5cos x，得",
             size=18, space_after=12)
    add_para(tf, "f(x) = 5cos x − cos 5x = 4cos^3^x (5 − 4cos^2^x)。",
             size=21, bold=True, color=GREEN, space_after=16)
    add_para(tf, "令 t = |cos x| ∈ [0, 1]，则 f(x) ≤ 4t^3^(5 − 4t^2^) = g(t)。",
             size=18, space_after=12)
    add_para(tf, "g′(t) = 20t^2^(3 − 4t^2^)，在 t = √3/2 处最大，g(√3/2) = 3√3。",
             size=18, space_after=16)
    add_para(tf, "所以对一切 x∈R，都有 f(x) ≤ 3√3。",
             size=20, bold=True, color=GREEN, space_after=10)
    add_para(tf, "—— 第(1)问的结论由此升级为全局上界，第(3)问直接可用。",
             size=13, color=GREY, space_after=0)
    return s


def slide_p2_proof(prs, layouts):
    s = prs.slides.add_slide(layouts[LAYOUT_BODY])
    set_header(s, "第(2)问：反证法 + 解集结构")
    tf = body_box(s)
    add_para(tf, "思路：先把 cos y > cos θ 的解集看清楚", size=13, color=GREY,
             space_after=16, first=True)
    add_para(tf, "cos y > cos θ ⇔ y ∈ (2kπ−θ, 2kπ+θ)，k∈Z；",
             size=19, space_after=10)
    add_para(tf, "这些开区间长度都是 2θ，并且两两不相交。",
             size=18, space_after=14)
    add_para(tf, "反设 [a−θ, a+θ] 内的点都满足 cos y > cos θ，"
                 "则该区间连通、长度恰为 2θ，",
             size=18, space_after=10)
    add_para(tf, "只能整体落在某个 (2kπ−θ, 2kπ+θ) 内，"
                 "迫使 a > 2kπ 且 a < 2kπ，矛盾。",
             size=18, space_after=16)
    add_para(tf, "故存在 y∈[a−θ, a+θ]，使得 cos y ≤ cos θ。",
             size=20, bold=True, color=GREEN, space_after=0)
    return s


def slide_p3_lower(prs, layouts):
    s = prs.slides.add_slide(layouts[LAYOUT_BODY])
    set_header(s, "第(3)问：b 至少要多大")
    tf = body_box(s)
    add_para(tf, "思路：用第(2)问造出一个点，逼出 b 的下界", size=13,
             color=GREY, space_after=16, first=True)
    add_para(tf, "记 h(x) = 5cos x − cos(5x+φ)。由 h(x+π) = −h(x) 知 |h(x)| ≤ b。",
             size=18, space_after=12)
    add_para(tf, "在(2)中取 θ = 5π/6、a = φ：存在 y∈[φ−5π/6, φ+5π/6]，",
             size=18, space_after=6)
    add_para(tf, "使 cos y ≤ cos(5π/6) = −√3/2。",
             size=18, space_after=12)
    add_para(tf, "取 x = (y−φ)/5，则 |x| ≤ π/6 且 5x+φ = y，于是",
             size=18, space_after=10)
    add_para(tf, "h(x) = 5cos x − cos y ≥ 5cos(π/6) + √3/2 = 3√3，",
             size=20, bold=True, color=GREEN, space_after=10)
    add_para(tf, "即 b ≥ 3√3。", size=22, bold=True, color=GREEN, space_after=0)
    return s


def slide_p3_upper(prs, layouts):
    s = prs.slides.add_slide(layouts[LAYOUT_BODY])
    set_header(s, "第(3)问：3√3 能取到")
    tf = body_box(s)
    add_para(tf, "思路：找一个 φ，让 b = 3√3 真的成立", size=13, color=GREY,
             space_after=18, first=True)
    add_para(tf, "取 φ = 0，则 5cos x − cos(5x+φ) = 5cos x − cos 5x",
             size=19, space_after=6)
    add_para(tf, "= 4cos^3^x (5 − 4cos^2^x) ≤ 3√3（第(1)问的恒等变形）。",
             size=19, bold=True, color=GREEN, space_after=16)
    add_para(tf, "所以 b = 3√3 满足题目的恒成立要求。", size=19, space_after=18)
    add_para(tf, "再结合上面得到的 b ≥ 3√3，两边夹住：", size=18, color=SOFT,
             space_after=16)
    add_para(tf, "b 的最小值 = 3√3", size=32, bold=True, color=GREEN,
             space_after=14)
    add_para(tf, "等号在 x = ±π/6 + 2kπ 处取到，例如 x = π/6。",
             size=13, color=GREY, space_after=0)
    return s


def slide_method(prs, layouts):
    s = prs.slides.add_slide(layouts[LAYOUT_BODY])
    set_header(s, "回看这道题：三个关键动作")
    tf = body_box(s)
    add_para(tf, "把三问连起来看，它其实是一条完整的推理链", size=13,
             color=GREY, space_after=18, first=True)
    add_para(tf, "① 求导看单调 —— 拿到第(1)问的局部最大值；", size=19, space_after=14)
    add_para(tf, "② 五倍角恒等变形 —— 把局部结论升级为对一切 x∈R 的上界；",
             size=19, space_after=14)
    add_para(tf, "③ 反证 + 解集结构 —— 由第(2)问造出下界。", size=19, space_after=18)
    add_para(tf, "第(2)问给下界，第(1)问给上界，一夹就得到最小值。",
             size=20, bold=True, color=GREEN, space_after=16)
    add_para(tf, "易错点：把第(1)问的最大值直接当作本题答案，忽视三问之间的联系。",
             size=13, color=GREY, space_after=0)
    return s


def slide_summary(prs, layouts):
    """末页：列出本题解题过程中用到的数学思想与方法。"""
    s = prs.slides.add_slide(layouts[LAYOUT_END])
    fill_placeholder(s, 12, [{"text": "2025 新课标Ⅰ卷 · 第 19 题", "size": 18,
                              "color": SOFT, "space_after": 0}])
    fill_placeholder(s, 11, [{"text": "总结：解题思想与方法", "size": 32,
                              "color": GREEN, "bold": True, "space_after": 0}])

    items = [
        ("化归与转化", "把 cos 5x 化为 cos x 的多项式（五倍角公式）"),
        ("换元法", "令 t = |cos x| ∈ [0,1]，三角最值化归为多项式最值"),
        ("导数法", "求导定单调，解决 [0, π/4] 上的局部最大值"),
        ("反证法", "借助 cos y > cos θ 的解集结构证明存在性"),
        ("放缩法", "用 |cos x| ≤ 1 控制上界"),
        ("夹逼思想", "下界来自(2)，上界来自(1)，两边夹出最小值"),
        ("构造与特殊值", "取 θ = 5π/6 造下界，取 φ = 0 验证可取"),
    ]
    body = [{"text": "%s：%s" % (name, desc), "size": 13, "color": SOFT,
             "space_after": 2.5, "line_spacing": 1.12} for name, desc in items]
    body.append({"text": "本题结论：b 的最小值为 3√3", "size": 13.5,
                 "color": GREEN, "bold": True, "space_after": 0,
                 "line_spacing": 1.12})
    fill_placeholder(s, 17, body)

    # 版式里这段只有 0.31 in 高，整块上移并加高，给方法清单留出空间
    set_placeholder_box(ph_by_idx(s, 12), 0.73, 1.90, 6.34, 0.42)
    set_placeholder_box(ph_by_idx(s, 11), 0.73, 2.40, 6.34, 0.66)
    set_placeholder_box(ph_by_idx(s, 17), 0.73, 3.18, 6.34, 2.05)
    return s


# ======================================================================
def main():
    build_figure()

    prs = Presentation(TEMPLATE)
    layouts = prs.slide_masters[0].slide_layouts
    strip_slides(prs)

    slide_cover(prs, layouts)
    slide_problem(prs, layouts)
    slide_p1_derivative(prs, layouts)
    slide_p1_identity(prs, layouts)
    slide_p2_proof(prs, layouts)
    slide_p3_lower(prs, layouts)
    slide_p3_upper(prs, layouts)
    slide_method(prs, layouts)
    slide_summary(prs, layouts)

    prune_template(prs, keep_layouts={
        "/ppt/slideLayouts/slideLayout1.xml",    # 标题幻灯片
        "/ppt/slideLayouts/slideLayout9.xml",    # 仅标题（无主副标题）
        "/ppt/slideLayouts/slideLayout16.xml",   # 末尾幻灯片
    })

    prs.save(OUTPUT)
    print("saved ->", OUTPUT, "(%d slides)" % len(prs.slides._sldIdLst))


if __name__ == "__main__":
    main()
