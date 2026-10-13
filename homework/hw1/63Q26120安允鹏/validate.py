#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""对生成结果做结构化校验：页数、文字、几何、溢出、包完整性。"""

import os
import sys
import zipfile
import unicodedata

from pptx import Presentation
from pptx.util import Emu
from lxml import etree

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    BASE, "ai_agent_demo_数学压轴v2_sample.pptx")

EMU_IN = 914400.0
SLIDE_W, SLIDE_H = 13.3333, 7.5
FOOTER_Y = 6.75   # 页脚/页码所在的横带，正文不应压到它


def char_w(ch, size):
    """返回字符宽度（pt）的粗略估计。"""
    if ch == "\t":
        return size
    o = ord(ch)
    if o < 128:
        return size * (0.28 if ch == " " else 0.55)
    if unicodedata.east_asian_width(ch) in ("W", "F"):
        return size
    return size * 0.62


def est_lines(text, size, avail_pt):
    if not text:
        return 0
    total = sum(char_w(c, size) for c in text)
    return max(1, int(total / avail_pt + 0.999))


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type is not None and sh.shape_type == 6:  # group
            for sub in walk(sh.shapes):
                yield sub


def main():
    prs = Presentation(OUT)
    print("slides:", len(prs.slides))
    print("size  : %.3f x %.3f in" % (prs.slide_width / EMU_IN, prs.slide_height / EMU_IN))

    problems = []
    for i, slide in enumerate(prs.slides, 1):
        lay = slide.slide_layout
        print("\n" + "=" * 72)
        print("P%-2d  layout=%s" % (i, lay.name))

        for sh in walk(slide.shapes):
            L = None if sh.left is None else sh.left / EMU_IN
            T = None if sh.top is None else sh.top / EMU_IN
            W = None if sh.width is None else sh.width / EMU_IN
            H = None if sh.height is None else sh.height / EMU_IN

            if L is not None:
                if L < -0.02 or T < -0.02 or L + W > SLIDE_W + 0.02 or T + H > SLIDE_H + 0.02:
                    problems.append("P%d %s 超出页面: (%.2f,%.2f,%.2f,%.2f)" % (i, sh.name, L, T, W, H))

            if sh.shape_type is not None and str(sh.shape_type).startswith("PICTURE"):
                print("   [图] %-22s (%.2f,%.2f) %.2fx%.2f" % (sh.name, L, T, W, H))
                continue

            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue

            tf = sh.text_frame
            avail_w = (W - 0.10) * 72 if W else 0
            total_h = 0.0
            for p in tf.paragraphs:
                if not p.runs:
                    continue
                size = max((r.font.size.pt for r in p.runs if r.font.size), default=18)
                txt = "".join(r.text for r in p.runs)
                ls = p.line_spacing if p.line_spacing else 1.2
                n = est_lines(txt, size, avail_w)
                sa = p.space_after.pt if p.space_after is not None else 0
                sb = p.space_before.pt if p.space_before is not None else 0
                total_h += n * size * ls + sa + sb
                print("      · %-58s  %2d行 %4.0fpt" % (txt[:56], n, size))
            total_in = total_h / 72.0
            bottom = (T or 0) + total_in
            flag = ""
            if H is not None and total_in > H + 0.05:
                flag = "  <<< 超出文本框 %.2f in" % (total_in - H)
                problems.append("P%d %s 文本超框 %.2f in" % (i, sh.name, total_in - H))
            if bottom > FOOTER_Y:
                flag += "  <<< 压到页脚 (bottom=%.2f)" % bottom
                problems.append("P%d %s 压到页脚 bottom=%.2f" % (i, sh.name, bottom))
            print("      -> 估算文本高 %.2f in, 框高 %s%s"
                  % (total_in, ("%.2f in" % H) if H else "继承", flag))

    print("\n" + "=" * 72)
    if problems:
        print("发现 %d 个问题：" % len(problems))
        for p in problems:
            print("  -", p)
    else:
        print("几何/溢出检查：未发现问题")

    # ---- 包完整性 ----
    z = zipfile.ZipFile(OUT)
    bad = []
    for n in z.namelist():
        if n.endswith((".xml", ".rels")):
            try:
                etree.fromstring(z.read(n))
            except Exception as e:
                bad.append((n, str(e)))
    print("包内 part 数：%d；XML 解析失败：%d" % (len(z.namelist()), len(bad)))
    for n, e in bad:
        print("   !", n, e)
    slides = [n for n in z.namelist() if n.startswith("ppt/slides/slide")]
    media = [n for n in z.namelist() if n.startswith("ppt/media/")]
    print("slide part=%d, media=%d, 文件大小=%.1f MB"
          % (len(slides), len(media), os.path.getsize(OUT) / 1e6))


if __name__ == "__main__":
    main()
