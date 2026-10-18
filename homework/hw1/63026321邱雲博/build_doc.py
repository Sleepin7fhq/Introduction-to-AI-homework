# -*- coding: utf-8 -*-
"""生成 2026 新高考Ⅰ卷第14题 详解作业文档"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HEI, SONG, LAT = "黑体", "宋体", "Times New Roman"
DARK = RGBColor(0x1F, 0x1F, 0x1F)
GRAY = RGBColor(0x59, 0x59, 0x59)
ACC = RGBColor(0x1F, 0x39, 0x64)

q = 1.5 ** (1 / 3)
q2 = q * q
x = 4 / (q2 * (1 + q + q2))
NUM = {
    "q": f"{q:.4f}",
    "x": f"{x:.4f}",
}
print("q =", q, "x =", x, "B2 =", x * q2 * (1 + q + q2), "B3 =", x * q2 * (1 + q + q2) * q ** 3)


# ---------- 基础工具 ----------
def set_fonts(run, ea=SONG, latin=LAT):
    run.font.name = latin
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:eastAsia"), ea)
    rFonts.set(qn("w:cs"), latin)


def add_math(p, s, size, bold=False):
    """数学片段：字母斜体，数字/运算符正体，支持 _{..} ^{..}"""
    for tok in re.split(r"(_\{[^}]*\}|\^\{[^}]*\})", s):
        if not tok:
            continue
        sub = tok.startswith("_{")
        sup = tok.startswith("^{")
        body = tok[2:-1] if (sub or sup) else tok
        for m in re.finditer(r"[A-Za-z]+|[^A-Za-z]+", body):
            t = m.group(0)
            r = p.add_run(t)
            r.font.size = Pt(size * (0.88 if (sub or sup) else 1.0))
            r.bold = bold
            r.italic = bool(re.fullmatch(r"[A-Za-z]+", t))
            if sub:
                r.font.subscript = True
            if sup:
                r.font.superscript = True
            set_fonts(r)


def add_rich(p, s, size=12, bold=False, ea=SONG, color=None):
    for i, part in enumerate(s.split("$")):
        if i % 2 == 1:
            start = len(p.runs)
            add_math(p, part, size, bold)
            if color is not None:
                for r in p.runs[start:]:
                    r.font.color.rgb = color
        elif part:
            r = p.add_run(part)
            r.font.size = Pt(size)
            r.bold = bold
            r.italic = False
            if color is not None:
                r.font.color.rgb = color
            set_fonts(r, ea=ea)


LINE = 1.31
AFTER = 3


def para(container, text, size=12, bold=False, indent=True, align=None,
         before=0, after=AFTER, line=LINE, ea=SONG, color=None, left=0.0, keep=False, pb=False):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if indent:
        pf.first_line_indent = Pt(size * 2)
    if left:
        pf.left_indent = Cm(left)
    if align is not None:
        p.alignment = align
    if keep:
        pf.keep_with_next = True
    if pb:
        pf.page_break_before = True
    add_rich(p, text, size, bold, ea, color)
    return p


def bottom_border(p, color="BFBFBF", sz=6):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), "4")
    b.set(qn("w:color"), color)
    pBdr.append(b)
    pPr.append(pBdr)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def table_borders(t, color="808080", sz=4):
    tblPr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        borders.append(e)
    tblPr.append(borders)


def cell_margins(t, top=60, bottom=60, left=110, right=110):
    tblPr = t._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)


def fixed_layout(t, widths_cm):
    t.autofit = False
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for row in t.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
    for i, w in enumerate(widths_cm):
        t.columns[i].width = Cm(w)


def keep_table_together(t):
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
    for row in t.rows[:-1]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True


def cell_text(cell, text, size=11, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, ea=SONG, color=None):
    p = cell.paragraphs[0]
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = 1.25
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)
    add_rich(p, text, size, bold, ea, color)


def box(container, lines, fill="F7F7F7", border="C9C9C9", size=12, bold=False,
        align=WD_ALIGN_PARAGRAPH.LEFT, width=16.4, spacing=1.45):
    t = container.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t, color=border, sz=8)
    cell_margins(t, top=170, bottom=170, left=220, right=220)
    fixed_layout(t, [width])
    c = t.cell(0, 0)
    shade(c, fill)
    first = True
    for ln in lines:
        if first:
            p = c.paragraphs[0]
            first = False
        else:
            p = c.add_paragraph()
        p.alignment = align
        pf = p.paragraph_format
        pf.line_spacing = spacing
        pf.space_before = Pt(3)
        pf.space_after = Pt(3)
        add_rich(p, ln, size, bold)


def add_field(p, instr, size=9, color=GRAY):
    r = p.add_run()
    r.font.size = Pt(size)
    r.font.color.rgb = color
    set_fonts(r)
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "end")
    r._r.append(f1)
    r._r.append(it)
    r._r.append(f2)


# ---------- 文档 ----------
doc = Document()

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.1), Cm(2.0)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.header_distance, sec.footer_distance = Cm(1.2), Cm(1.1)

normal = doc.styles["Normal"]
normal.font.name = LAT
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)
normal.paragraph_format.line_spacing = LINE
normal.paragraph_format.space_after = Pt(4)

for name, size, before, after in (("Heading 1", 14.5, 13, 6), ("Heading 2", 12.5, 8, 4)):
    st = doc.styles[name]
    st.font.name = LAT
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = ACC
    st.font.italic = False
    rPr = st.element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.append(rf)
    rf.set(qn("w:ascii"), LAT)
    rf.set(qn("w:hAnsi"), LAT)
    rf.set(qn("w:eastAsia"), HEI)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.line_spacing = 1.2
    st.paragraph_format.keep_with_next = True

# 页眉
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_rich(hp, "2026 年新高考Ⅰ卷数学 · 第 14 题详解", 9, color=GRAY)
bottom_border(hp, "D0D0D0")

# 页脚
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("— ")
r.font.size = Pt(9)
r.font.color.rgb = GRAY
set_fonts(r)
add_field(fp, "PAGE")
r = fp.add_run(" —")
r.font.size = Pt(9)
r.font.color.rgb = GRAY
set_fonts(r)

# 标题
t = para(doc, "2026 年新高考Ⅰ卷数学 第 14 题详解", 19, True,
         indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=2, ea=HEI, color=DARK, keep=True)
para(doc, "2026 年普通高等学校招生全国统一考试（新高考Ⅰ卷） 三、填空题 第 14 题 · 5 分",
     10.5, False, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=4, ea=SONG, color=GRAY)
info = para(doc, "姓名：＿＿＿＿＿　班级：＿＿＿＿＿　学号：＿＿＿＿＿　得分：＿＿＿＿",
            10, False, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=8, ea=SONG, color=GRAY)
bottom_border(info, "9E9E9E", 6)

# 一、原题
para(doc, "一、原题重现", 14.5, True, indent=False, before=4, after=6, ea=HEI, color=ACC, keep=True)
box(doc, [
    "设实数 $q$ 满足：存在数列 ${a_{n}}$，使得对于任意 $n∈N^{*}$，均有",
    "$a_{1}+a_{2}+⋯+a_{3n}=n^{2}+n$，",
    "且 ${a_{n}}$ 中有某连续 9 项 $a_{k},a_{k+1},⋯,a_{k+8}$ 是公比为 $q$ 的等比数列，",
    "则 $q$ 的最大值为＿＿＿＿＿＿．",
], fill="F5F6F8", border="C9CDD4", size=12)

para(doc, "【出处】2026 年普通高等学校招生全国统一考试数学（新高考Ⅰ卷）填空题第 14 题，为本卷填空题的最后一题，5 分。"
          "本题以“每三项之和”的递推约束为载体，考查等比数列的性质、分组求和与分类讨论。",
     10, False, indent=False, before=6, after=4, ea=SONG, color=GRAY)

# 二、审题
para(doc, "二、审题：把条件翻译成“三项一组”", 14.5, True, indent=False, before=8, after=6, ea=HEI, color=ACC)
para(doc, "题设只约束“每连续三项之和”，对每一项本身不作要求。因此解题的第一步，是让“9 个连续项”与“每 3 项一组”的分块对齐。")

para(doc, "1．条件只锁定“组和”", 12.5, True, indent=False, before=6, after=4, ea=HEI, color=DARK, keep=True)
para(doc, "记第 $n$ 组的三项和为 $B_{n}=a_{3n-2}+a_{3n-1}+a_{3n}$（$n≥1$）。设 $S_{n}=a_{1}+a_{2}+⋯+a_{n}$（约定 $S_{0}=0$），由 $S_{3n}=n^{2}+n$ 得",
     after=2)
para(doc, "$B_{n}=S_{3n}-S_{3n-3}=(n^{2}+n)-[(n-1)^{2}+(n-1)]=2n$。",
     indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
para(doc, "即各组之和依次为 $2,4,6,8,⋯$，是公差为 $2$ 的等差数列，且恒为正数。每一项的具体取值可以任意，只有“三项一组”的组和被锁定——这正是本题的突破口。")

para(doc, "2．相隔 3 个位置的两段，其和之比恰为 $q^{3}$", 12.5, True, indent=False, before=6, after=4, ea=HEI, color=DARK, keep=True)
para(doc, "把连续三项看成“一段”，整段后移 3 个位置相当于段内各项都乘以 $q^{3}$，故“位置相差 3 的两段”之和之比恰为 $q^{3}$。9 项窗口里含几个完整组，由 $k$ 除以 3 的余数决定：", after=6)

tbl = doc.add_table(rows=4, cols=3)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
table_borders(tbl, "9AA5B1", 6)
cell_margins(tbl, 70, 70, 110, 110)
fixed_layout(tbl, [4.4, 7.6, 4.4])
hdr = ["$k$ 除以 3 的余数", "9 项中含有的完整组", "完整组个数"]
for i, h in enumerate(hdr):
    cell_text(tbl.rows[0].cells[i], h, 10.5, True, WD_ALIGN_PARAGRAPH.CENTER, HEI)
    shade(tbl.rows[0].cells[i], "E8ECF1")
body = [
    ("1（$k=3m+1$）", "第 $m+1$、$m+2$、$m+3$ 组", "3"),
    ("2（$k=3m+2$）", "第 $m+2$、$m+3$ 组", "2"),
    ("0（$k=3m$）", "第 $m+1$、$m+2$ 组", "2"),
]
for r_i, row in enumerate(body, start=1):
    for c_i, v in enumerate(row):
        cell_text(tbl.rows[r_i].cells[c_i], v, 10.5, False, WD_ALIGN_PARAGRAPH.CENTER)
keep_table_together(tbl)

para(doc, "三种余数对应三种情形，其中“3 个完整组”的情形正是本题的“陷阱”，下面逐一讨论。",
     11, False, indent=False, before=6, after=2, ea=SONG, color=GRAY)

# 三、解答
para(doc, "三、规范解答", 14.5, True, indent=False, before=10, after=6, ea=HEI, color=ACC)

para(doc, "第一步　求每三项之和", 12.5, True, indent=False, before=2, after=4, ea=HEI, color=DARK, keep=True)
para(doc, "设 $S_{n}=a_{1}+a_{2}+⋯+a_{n}$，$S_{0}=0$。对任意正整数 $n$，", after=2)
para(doc, "$B_{n}=S_{3n}-S_{3n-3}=2n$，即 $a_{3n-2}+a_{3n-1}+a_{3n}=2n$。",
     indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)

para(doc, "第二步　按 $k$ 除以 3 的余数分类讨论", 12.5, True, indent=False, before=6, after=4, ea=HEI, color=DARK, keep=True)
para(doc, "设这 9 项的首项 $a_{k}=x$，则这 9 项依次为 $x,xq,xq^{2},⋯,xq^{8}$。")

para(doc, "情形一　$k=3m+1$（$m$ 为非负整数）", 12, True, indent=False, before=6, after=3, ea=HEI, color=DARK, keep=True)
para(doc, "这 9 项恰好是第 $m+1$、$m+2$、$m+3$ 组，组内三项分别对应等比数列的第 1 至 3 项、第 4 至 6 项、第 7 至 9 项，故三组之和依次为", after=2)
para(doc, "$x(1+q+q^{2})$，$xq^{3}(1+q+q^{2})$，$xq^{6}(1+q+q^{2})$。",
     indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
para(doc, "因为 $1+q+q^{2}=(q+1/2)^{2}+3/4>0$，而第 $m+1$ 组之和 $2m+2>0$，所以 $x≠0$，可以对组和作比：", after=2)
para(doc, "$q^{3}=B_{m+2}/B_{m+1}=(m+2)/(m+1)$，且 $q^{3}=B_{m+3}/B_{m+2}=(m+3)/(m+2)$。",
     indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
para(doc, "两式合并得 $(m+2)^{2}=(m+1)(m+3)$，即 $m^{2}+4m+4=m^{2}+4m+3$，矛盾。故 $k$ 除以 3 余 1 的情形不可能出现。")

para(doc, "情形二　$k=3m+2$（$m≥0$）", 12, True, indent=False, before=6, after=3, ea=HEI, color=DARK, keep=True)
para(doc, "9 项为 $a_{3m+2},⋯,a_{3m+10}$。其中第 3 至 5 项是完整的第 $m+2$ 组，第 6 至 8 项是完整的第 $m+3$ 组，两组之和分别为 $2m+4$、$2m+6$，故", after=2)
para(doc, "$q^{3}=B_{m+3}/B_{m+2}=(2m+6)/(2m+4)=(m+3)/(m+2)=1+1/(m+2)≤3/2$，",
     indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
para(doc, "等号当且仅当 $m=0$，即 $k=2$。")

para(doc, "情形三　$k=3m$（$m≥1$）", 12, True, indent=False, before=6, after=3, ea=HEI, color=DARK, keep=True)
para(doc, "9 项为 $a_{3m},⋯,a_{3m+8}$。其中第 2 至 4 项是完整的第 $m+1$ 组，第 5 至 7 项是完整的第 $m+2$ 组，两组之和分别为 $2m+2$、$2m+4$，故", after=2)
para(doc, "$q^{3}=B_{m+2}/B_{m+1}=(2m+4)/(2m+2)=(m+2)/(m+1)=1+1/(m+1)≤3/2$，",
     indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
para(doc, "等号当且仅当 $m=1$，即 $k=3$。")

para(doc, "综合三种情形，$q^{3}≤3/2$ 恒成立。又各组之和为正，故 $q^{3}>0$，从而 $q>0$，于是", before=4, after=2)
para(doc, "$q≤∛(3/2)$。", indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)

para(doc, "第三步　验证等号可以取到（存在性）", 12.5, True, indent=False, before=8, after=4, ea=HEI, color=DARK, keep=True)
para(doc, "取 $k=2$（属于情形二，$m=0$），令 $q=∛(3/2)$，$x=4/[q^{2}(1+q+q^{2})]$，按如下方式构造数列：", after=2)
para(doc, "$a_{i}=xq^{i-2}\\ (i=2,3,⋯,10)$；$a_{1}=2-a_{2}-a_{3}$，$a_{11}+a_{12}=8-a_{10}$；当 $n≥5$ 时，$a_{3n-2}=2n$，$a_{3n-1}=a_{3n}=0$。".replace("\\ ", " "),
     indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
para(doc, "此时 $B_{1}=2$；$B_{2}=xq^{2}(1+q+q^{2})=4$；$B_{3}=xq^{5}(1+q+q^{2})=q^{3}·B_{2}=6$；$B_{4}=8$；$n≥5$ 时 $B_{n}=2n$。"
          "即 $B_{n}=2n$ 对一切正整数 $n$ 成立，从而", before=2, after=2)
para(doc, "$S_{3n}=2(1+2+⋯+n)=n(n+1)=n^{2}+n$，",
     indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
para(doc, "满足题设条件；同时 $a_{2},a_{3},⋯,a_{10}$ 是公比为 $∛(3/2)$ 的等比数列。故 $q=∛(3/2)$ 确实可以取到。")

# 四、结论
para(doc, "四、结论", 14.5, True, indent=False, before=10, after=6, ea=HEI, color=ACC)
box(doc, ["$q$ 的最大值为 $∛(3/2)$，即 $∛12/2$，约为 $1.1447$。"],
    fill="EAF1F8", border="8FA8C8", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.6)

# 五、数值检验
para(doc, "五、数值检验", 14.5, True, indent=False, before=12, after=6, ea=HEI, color=ACC, keep=True)
para(doc, "取 $q=∛(3/2)≈" + NUM["q"] + "$，则 $x=a_{2}≈" + NUM["x"] + "$。逐项验算，各组的和确实为 $2n$：", after=6, keep=True)

rows = [
    ("$q^{3}$", "$3/2$", "1.5000"),
    ("$B_{2}=a_{4}+a_{5}+a_{6}$", "$xq^{2}(1+q+q^{2})$", "4.0000"),
    ("$B_{3}=a_{7}+a_{8}+a_{9}$", "$q^{3}·B_{2}$", "6.0000"),
    ("$S_{9}$", "$B_{1}+B_{2}+B_{3}$", "$12=3^{2}+3$"),
]
tbl2 = doc.add_table(rows=len(rows) + 1, cols=3)
tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
table_borders(tbl2, "9AA5B1", 6)
cell_margins(tbl2, 45, 45, 110, 110)
keep_table_together(tbl2)
fixed_layout(tbl2, [5.0, 6.4, 5.0])
for i, h in enumerate(["检验对象", "计算式", "数值"]):
    cell_text(tbl2.rows[0].cells[i], h, 10.5, True, WD_ALIGN_PARAGRAPH.CENTER, HEI)
    shade(tbl2.rows[0].cells[i], "E8ECF1")
for r_i, row in enumerate(rows, start=1):
    for c_i, v in enumerate(row):
        cell_text(tbl2.rows[r_i].cells[c_i], v, 10.5, False, WD_ALIGN_PARAGRAPH.CENTER)

# 六、易错点
para(doc, "六、易错点提醒", 14.5, True, indent=False, before=12, after=6, ea=HEI, color=ACC)
errors = [
    "只盯着“9 项求和”而不分块。9 项本身没有求和公式，必须借“每 3 项一组”的组和，才能“透视”出 $q^{3}$。",
    "漏掉情形一。$k=3m+1$ 时 9 项内恰有 3 个完整组，看似条件最充分，实则必然矛盾；漏掉这一支就会误得 $∛2$。",
    "只用一个完整组作比。窗口内只含 1 个完整组时，关于 $q$ 只有一个方程，对 $q$ 没有限制，得不到上界。",
    "忘记验证存在性。$q^{3}≤3/2$ 只是必要条件，必须构造出具体的 ${a_{n}}$ 说明 $q=∛(3/2)$ 可以取到。",
    "忽略 $q$ 的符号。由组和之比为正得 $q^{3}>0$，故 $q>0$，求最大值时不必考虑负数。",
]
for i, e in enumerate(errors, start=1):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = LINE
    pf.space_before = Pt(1)
    pf.space_after = Pt(2)
    pf.left_indent = Cm(1.0)
    pf.first_line_indent = Cm(-1.0)
    add_rich(p, f"{i}．{e}", 12)

# 七、拓展
para(doc, "七、拓展思考（选做）", 14.5, True, indent=False, before=12, after=6, ea=HEI, color=ACC)
para(doc, "题目中的“9”并非随意取定。若把“连续 9 项”改为“连续 $L$ 项”，结论会随 $L$ 变化：")
ext = [
    "$L≤7$：总有某个起点使窗口内只含 1 个完整组，此时 $q$ 没有上界；",
    "$L=8$：最大值为 $∛2$（取 $k=1$，8 项恰含第 1、2 组，$q^{3}=B_{2}/B_{1}=2$）；",
    "$L=9$ 或 $10$：最大值为 $∛(3/2)$；",
    "$L≥11$：任何起点都含至少 3 个完整组，而 $2m+2$、$2m+4$、$2m+6$ 不可能成等比数列，故这样的 $q$ 不存在。",
]
for i, e in enumerate(ext, start=1):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = LINE
    pf.space_before = Pt(1)
    pf.space_after = Pt(2)
    pf.left_indent = Cm(1.1)
    pf.first_line_indent = Cm(-1.1)
    add_rich(p, f"（{i}）{e}", 12)
para(doc, "可见 $L=9$ 恰是“临界值”：$k$ 除以 3 余 1 时窗口内正好放下 3 个完整组，相邻两组之比 $2$ 与 $3/2$ 不可能同时成立，"
          "“$k=1$、$q=∛2$”的方案被排除，最大值才由 $∛2$ 降为 $∛(3/2)$。", before=4)

doc.save("2026年新高考Ⅰ卷数学第14题详解.docx")
print("saved")
