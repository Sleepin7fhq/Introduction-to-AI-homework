# -*- coding: utf-8 -*-
"""Generate the PPT for the m-interesting sequence problem."""
import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------- palette ----------
DARK   = RGBColor(0x1F, 0x38, 0x64)
ACCENT = RGBColor(0x2E, 0x74, 0xB5)
LIGHT  = RGBColor(0xDE, 0xEB, 0xF7)
LIGHT2 = RGBColor(0xEF, 0xF5, 0xFB)
TEXT   = RGBColor(0x21, 0x21, 0x21)
GRAY   = RGBColor(0x5A, 0x5A, 0x5A)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
RED    = RGBColor(0xC0, 0x00, 0x00)
GOLD   = RGBColor(0xB4, 0x79, 0x1B)

FONT = "Microsoft YaHei"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SLIDE_W = 13.333
SLIDE_H = 7.5

_PAGE = {"n": 0}


def char_units(s):
    u = 0.0
    for ch in s:
        o = ord(ch)
        if ch in "　" or o >= 0x2E80 or 0x3000 <= o <= 0x303F or 0xFF00 <= o <= 0xFFEF:
            u += 1.0
        elif o < 0x80:
            u += 0.54
        else:
            u += 1.0
    return u


def est_lines(text, size, width_in):
    units = char_units(text)
    capacity = max(4.0, width_in * 72.0 / size)
    return max(1, int(math.ceil(units / capacity)))


def est_height(text, size, width_in, line_spacing=1.15, pad=0.06):
    return est_lines(text, size, width_in) * size * line_spacing / 72.0 + pad


def add_rect(slide, x, y, w, h, fill, line=None, line_w=1.0):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    return sh


def add_text(slide, x, y, w, h, text, size, bold=False, color=TEXT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False,
             font=FONT, line_spacing=1.15, space_after=0, space_before=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    if space_before:
        p.space_before = Pt(space_before)
    if space_after:
        p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return tb, tf


def new_slide(title, subtitle=None):
    _PAGE["n"] += 1
    slide = prs.slides.add_slide(BLANK)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    add_rect(slide, 0, 0, SLIDE_W, 1.06, DARK)
    add_rect(slide, 0, 1.06, SLIDE_W, 0.07, ACCENT)
    add_text(slide, 0.6, 0.12, 11.2, 0.82, title, 26, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_text(slide, 0.62, 0.72, 11.5, 0.32, subtitle, 12,
                 color=RGBColor(0xBF, 0xD0, 0xE8))
    add_text(slide, 12.35, 7.05, 0.85, 0.35, str(_PAGE["n"]), 10, color=GRAY,
             align=PP_ALIGN.RIGHT)
    return slide


def body(slide, blocks, y0=1.34, bottom_limit=7.05):
    y = y0
    for b in blocks:
        gap = b.get("gap", 7)
        y += gap / 72.0
        text = b["text"]
        size = b.get("size", 18)
        bold = b.get("bold", False)
        color = b.get("color", TEXT)
        align = b.get("align", PP_ALIGN.LEFT)
        italic = b.get("italic", False)
        indent = b.get("indent", 0.0)
        ls = b.get("line_spacing", 1.15)
        width_in = 12.1 - indent
        if b.get("box", False):
            h = b.get("box_h")
            if h is None:
                h = est_height(text, size, width_in - 0.5, ls, pad=0.22)
            sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                        Inches(0.6 + indent), Inches(y),
                                        Inches(12.1 - indent), Inches(h))
            sh.fill.solid()
            sh.fill.fore_color.rgb = b.get("box_fill", LIGHT2)
            sh.line.color.rgb = b.get("box_line", ACCENT)
            sh.line.width = Pt(1.2)
            sh.shadow.inherit = False
            tf = sh.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = Inches(0.22)
            tf.margin_right = Inches(0.22)
            tf.margin_top = Inches(0.04)
            tf.margin_bottom = Inches(0.04)
            p = tf.paragraphs[0]
            p.alignment = align
            p.line_spacing = ls
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = color
            r.font.name = FONT
            y += h
        else:
            h = b.get("h")
            if h is None:
                h = est_height(text, size, width_in, ls)
            add_text(slide, 0.6 + indent, y, width_in, h + 0.15, text, size,
                     bold=bold, color=color, align=align, italic=italic,
                     line_spacing=ls)
            y += h
        if y > bottom_limit:
            print("WARN overflow slide", _PAGE["n"], "y=", round(y, 2))
    return y


# ============================================================ slide 1 : title
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SLIDE_W, SLIDE_H, DARK)
add_rect(s, 0, 3.30, SLIDE_W, 0.045, GOLD)
add_text(s, 0.8, 1.55, 11.7, 1.1, "有趣的数列问题", 46, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER)
add_text(s, 0.8, 2.80, 11.7, 0.5, "m-有趣数列：最小数 m₀ 与数列计数", 22,
         color=RGBColor(0xBF, 0xD0, 0xE8), align=PP_ALIGN.CENTER)
add_text(s, 0.8, 3.80, 11.7, 1.3, "完整推理  ·  引理证明  ·  构造与计数", 17,
         color=RGBColor(0x9F, 0xB8, 0xDA), align=PP_ALIGN.CENTER)
add_text(s, 0.8, 6.55, 11.7, 0.5, "数学原理 · 逻辑清晰 · 结论严谨", 13,
         color=RGBColor(0x7F, 0x9B, 0xC4), align=PP_ALIGN.CENTER)

# ============================================================ slide 2 : problem
s = new_slide("问题陈述", "定义与目标")
body(s, [
    {"text": "定义（m-有趣数列）", "size": 20, "bold": True, "color": ACCENT, "gap": 2},
    {"text": "称满足下列 (i)、(ii) 的实数列  a₁, a₂, …, aₘ  (m ≥ 10)  为一个 m-有趣数列：", "size": 19, "gap": 6},
    {"text": "(i)  对任意 k = 1, 2, …, m−1，都有  |aₖ₊₁ − aₖ| = 1；", "size": 19, "gap": 9, "indent": 0.4},
    {"text": "(ii)  存在正整数  i₁ < i₂ < ⋯ < i₁₀ ≤ m，使  aᵢ₁, aᵢ₂, …, aᵢ₁₀ 成等比数列，", "size": 19, "gap": 9, "indent": 0.4},
    {"text": "且该等比数列的公比 q 满足  q ≠ ±1。", "size": 19, "indent": 0.4},
    {"text": "目标", "size": 20, "bold": True, "color": ACCENT, "gap": 18},
    {"text": "① 求使 m-有趣数列存在的最小正整数 m₀；", "size": 19, "gap": 7},
    {"text": "② 确定不同的 m₀-有趣数列的个数。", "size": 19},
])

# ============================================================ slide 3 : observations
s = new_slide("关键观察：整数步行走", "把条件 (i) 翻译为「±1 行走」")
body(s, [
    {"text": "观察 1（同余于模 1）", "size": 20, "bold": True, "color": ACCENT, "gap": 2},
    {"text": "|aₖ₊₁ − aₖ| = 1  ⟹  aₖ₊₁ − aₖ = ±1，累加得  aₖ = a₁ + 整数。", "size": 19, "gap": 6},
    {"text": "于是整个数列落在同一「模 1 剩余类」上：任意两项之差都是整数。", "size": 18, "color": GRAY, "gap": 4},
    {"text": "观察 2（等比子列之差）", "size": 20, "bold": True, "color": ACCENT, "gap": 14},
    {"text": "令  bⱼ = aᵢⱼ  (j = 1, …, 10) 为等比数列的 10 项，则  bⱼ − b₁ ∈ ℤ。", "size": 19, "gap": 6},
    {"text": "记相邻差  dⱼ = bⱼ₊₁ − bⱼ  (j = 1, …, 9)。由等比性：", "size": 19, "gap": 6},
    {"text": "dⱼ = bⱼ (q − 1) = b₁ qʲ⁻¹ (q − 1) = d₁ qʲ⁻¹，其中 d₁ = b₁(q − 1)。", "size": 20, "bold": True, "box": True, "box_h": 0.66, "gap": 9, "align": PP_ALIGN.CENTER},
    {"text": "因为 b₁ ≠ 0 且 q ≠ 1，所以 d₁ ≠ 0；又 dⱼ 是行走两值之差，故 dⱼ ∈ ℤ。", "size": 18, "color": GRAY, "gap": 12},
])

# ============================================================ slide 4 : q rational
s = new_slide("公比 q 必为有理数", "由整数差反推出 q 的结构")
body(s, [
    {"text": "第一步：q 是有理数", "size": 20, "bold": True, "color": ACCENT, "gap": 2},
    {"text": "d₁ 与 d₂ = d₁ q 都是非零整数 ⟹  q = d₂ / d₁ ∈ ℚ。", "size": 19, "gap": 6},
    {"text": "令  q = u / v，其中 u, v ∈ ℤ，v ≥ 1，gcd(u, v) = 1。", "size": 19, "gap": 8},
    {"text": "（等比数列公比须非零，故 q ≠ 0 ⟹ u ≠ 0；题设 q ≠ ±1 ⟹ |u| ≠ v。）", "size": 16, "color": GRAY, "gap": 5},
    {"text": "第二步：确定 d₁ 的结构", "size": 20, "bold": True, "color": ACCENT, "gap": 14},
    {"text": "dⱼ = d₁ uʲ⁻¹ / vʲ⁻¹ ∈ ℤ  对 j = 1, …, 9 成立。", "size": 19, "gap": 6},
    {"text": "由 gcd(u, v) = 1 得  vʲ⁻¹ | d₁（每个 j），特别取 j = 9： v⁸ | d₁。", "size": 19, "gap": 6},
    {"text": "故存在非零整数 c，使  d₁ = c·v⁸。", "size": 19, "gap": 6},
    {"text": "于是  |dⱼ| = |c| · |u|ʲ⁻¹ · v⁹⁻ʲ  (j = 1, …, 9)。", "size": 20, "bold": True, "box": True, "box_h": 0.66, "gap": 9, "align": PP_ALIGN.CENTER},
])

# ============================================================ slide 5 : lemma
s = new_slide("引理：差的和的最小值", "本引理给出关键下界，附完整证明")
body(s, [
    {"text": "引理.", "size": 20, "bold": True, "color": RED, "gap": 2},
    {"text": "设 A, B 为正整数，A ≠ B 且 gcd(A, B) = 1。令", "size": 19, "gap": 4},
    {"text": "T(A, B) = B⁸ + A·B⁷ + A²·B⁶ + ⋯ + A⁷·B + A⁸  （共 9 项）。", "size": 19, "bold": True, "box": True, "box_h": 0.6, "gap": 6, "align": PP_ALIGN.CENTER},
    {"text": "则  T(A, B) ≥ 511，且等号成立 ⟺ {A, B} = {1, 2}。", "size": 19, "gap": 10},
    {"text": "证明", "size": 18, "bold": True, "color": ACCENT, "gap": 10},
    {"text": "① T 关于 A、B 对称，且随 A、B 严格单调递增（每项 Aʲ⁻¹ B⁹⁻ʲ 均如此）。", "size": 17, "gap": 5},
    {"text": "② 若 min(A, B) ≥ 2，则 T(A, B) ≥ T(2, 2) = 9 · 2⁸ = 2304 > 511。", "size": 17, "gap": 6},
    {"text": "③ 若 min(A, B) = 1，不妨 A = 1，则 B ≥ 2，且", "size": 17, "gap": 6},
    {"text": "T(1, B) = 1 + B + B² + ⋯ + B⁸ ≥ 1 + 2 + 2² + ⋯ + 2⁸ = 2⁹ − 1 = 511，", "size": 17, "indent": 0.4, "gap": 4},
    {"text": "等号成立当且仅当 B = 2。  ∎", "size": 17, "indent": 0.4},
])

# ============================================================ slide 6 : lower bound
s = new_slide("下界：m ≥ 512", "把引理接到行走长度上")
body(s, [
    {"text": "行走长度与差的关系", "size": 20, "bold": True, "color": ACCENT, "gap": 2},
    {"text": "每一步值至多改变 1，由三角不等式：", "size": 19, "gap": 6},
    {"text": "|bⱼ₊₁ − bⱼ| = |aᵢⱼ₊₁ − aᵢⱼ| ≤ iⱼ₊₁ − iⱼ，即  iⱼ₊₁ − iⱼ ≥ |dⱼ|。", "size": 19, "gap": 5},
    {"text": "对 j = 1, …, 9 累加：", "size": 19, "gap": 9},
    {"text": "i₁₀ − i₁ = (i₂−i₁) + ⋯ + (i₁₀−i₉) ≥ |d₁| + |d₂| + ⋯ + |d₉| = |c| · T(|u|, v)。", "size": 19, "gap": 6},
    {"text": "由引理  T(|u|, v) ≥ 511，且 |c| ≥ 1，故  i₁₀ − i₁ ≥ 511。", "size": 19, "gap": 8},
    {"text": "又 1 ≤ i₁ < i₁₀ ≤ m，于是", "size": 19, "gap": 9},
    {"text": "m ≥ i₁₀ ≥ i₁ + 511 ≥ 1 + 511 = 512。", "size": 22, "bold": True, "color": RED, "box": True, "box_h": 0.72, "gap": 6, "align": PP_ALIGN.CENTER},
    {"text": "结论：任何 m-有趣数列都满足 m ≥ 512。", "size": 18, "color": GRAY, "gap": 12},
])

# ============================================================ slide 7 : construction
s = new_slide("构造：m₀ = 512", "达到下界的例子")
body(s, [
    {"text": "取  aₖ = k  (k = 1, 2, …, 512)。", "size": 20, "bold": True, "gap": 4},
    {"text": "验证 (i)：|aₖ₊₁ − aₖ| = |(k+1) − k| = 1，满足。", "size": 19, "gap": 9},
    {"text": "验证 (ii)：取下标", "size": 19, "gap": 9},
    {"text": "i₁, i₂, …, i₁₀ = 1, 2, 4, 8, 16, 32, 64, 128, 256, 512，", "size": 19, "indent": 0.4, "gap": 4},
    {"text": "对应值", "size": 19, "gap": 4},
    {"text": "a₁, a₂, a₄, a₈, … = 1, 2, 4, 8, 16, 32, 64, 128, 256, 512，", "size": 19, "indent": 0.4, "gap": 4},
    {"text": "构成公比 q = 2 的等比数列。", "size": 19, "indent": 0.4},
    {"text": "故  m₀ = 512。", "size": 24, "bold": True, "color": RED, "box": True, "box_h": 0.74, "gap": 18, "align": PP_ALIGN.CENTER},
])

# ============================================================ slide 8 : tightness
s = new_slide("计数：512-有趣数列的紧性", "m = 512 时下界处处取等")
body(s, [
    {"text": "当 m = 512 时，下界链条必须处处取等：", "size": 19, "gap": 2},
    {"text": "①  i₁₀ − i₁ = 511  ⟹  |c| = 1 且 {|u|, v} = {1, 2}", "size": 19, "gap": 9},
    {"text": "⟹  q = u/v ∈ { 2, −2, 1/2, −1/2 }。", "size": 19, "indent": 0.4, "gap": 4},
    {"text": "②  i₁ = 1，i₁₀ = 512（首末恰好占满）。", "size": 19, "gap": 9},
    {"text": "③  每段  iⱼ₊₁ − iⱼ = |dⱼ|（逐段严格单调，绝不「回头」）。", "size": 19, "gap": 9},
    {"text": "因此：一旦给定 (q, c)，等比子列 bⱼ 与各段行走方向都被唯一确定，", "size": 18, "gap": 9},
    {"text": "整个数列  a₁, …, a₅₁₂ 也随之唯一确定。", "size": 18, "bold": True, "gap": 4},
    {"text": "所以 512-有趣数列恰由 (q, c) 取值决定：q 有 4 种、c 有 ±1 两种。", "size": 18, "gap": 11},
])

# ============================================================ slide 9 : the 8 sequences
s = new_slide("八个 512-有趣数列", "4 个单调型 + 4 个折返型")
body(s, [
    {"text": "单调型（q = 2 或 q = 1/2，c 决定走向）", "size": 18, "bold": True, "color": ACCENT, "gap": 2},
    {"text": "① aₖ = k          q = 2      子列 (1, 2, 4, …, 512)", "size": 17, "gap": 5, "indent": 0.3},
    {"text": "② aₖ = −k         q = 2      子列 (−1, −2, −4, …, −512)", "size": 17, "gap": 3, "indent": 0.3},
    {"text": "③ aₖ = k − 513    q = 1/2    子列 (−512, −256, …, −1)", "size": 17, "gap": 3, "indent": 0.3},
    {"text": "④ aₖ = 513 − k    q = 1/2    子列 (512, 256, …, 1)", "size": 17, "gap": 3, "indent": 0.3},
    {"text": "折返型（q = −2 或 q = −1/2，各取 c = ±1，共 4 个）", "size": 18, "bold": True, "color": ACCENT, "gap": 12},
    {"text": "⑤⑥ q = −2：  bⱼ = (−1)ʲ · c · 2ʲ⁻¹ / 3，下标 iⱼ = 2ʲ⁻¹；", "size": 17, "gap": 5, "indent": 0.3},
    {"text": "⑦⑧ q = −1/2：bⱼ = (−1)ʲ · c · 2¹⁰⁻ʲ / 3，下标 iⱼ = 513 − 2¹⁰⁻ʲ；", "size": 17, "gap": 3, "indent": 0.3},
    {"text": "（j = 1, …, 10；c = ±1 各给一个数列，段间按单调直线行走连接）", "size": 15, "color": GRAY, "gap": 3, "indent": 0.3},
    {"text": "合计：4 + 4 = 8 个，两两不同（首项、走向或分段长度不同）。", "size": 18, "bold": True, "color": RED, "gap": 12},
])

# ============================================================ slide 10 : conclusion
s = new_slide("结论", "问题答案")
body(s, [
    {"text": "最小正整数 m₀", "size": 20, "bold": True, "color": ACCENT, "gap": 16},
    {"text": "m₀ = 512", "size": 40, "bold": True, "color": DARK, "box": True, "box_h": 1.0, "box_fill": LIGHT2, "gap": 6, "align": PP_ALIGN.CENTER},
    {"text": "不同的 m₀-有趣数列个数", "size": 20, "bold": True, "color": ACCENT, "gap": 22},
    {"text": "8 个", "size": 40, "bold": True, "color": DARK, "box": True, "box_h": 1.0, "box_fill": LIGHT2, "gap": 6, "align": PP_ALIGN.CENTER},
    {"text": "（q = 2, −2, 1/2, −1/2 四种公比 × 两种符号 c = ±1）", "size": 15, "color": GRAY, "gap": 16, "align": PP_ALIGN.CENTER},
])

prs.save("有趣的数列问题.pptx")
print("saved: 有趣的数列问题.pptx")
print("slides:", _PAGE["n"])
