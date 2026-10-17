# -*- coding: utf-8 -*-
"""生成《2026 新课标 II 卷第 17 题 · 解三角形》演示文稿（16:9，python-pptx）"""
import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
FIG = os.path.join(ROOT, "图表")
OUT_DIR = os.path.join(ROOT, "..", "PowerPoint")
os.makedirs(OUT_DIR, exist_ok=True)
OUT = os.path.join(OUT_DIR, "2026新课标II卷第17题_解三角形.pptx")

# ---------------- 主题 ----------------
NAVY = RGBColor(0x14, 0x21, 0x3D)
INK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x64, 0x70, 0x85)
RED = RGBColor(0xC0, 0x39, 0x2B)
GOLD = RGBColor(0xBA, 0x84, 0x1A)
TEAL = RGBColor(0x12, 0x76, 0x80)
BLUE = RGBColor(0x26, 0x54, 0x9C)
GREEN = RGBColor(0x2E, 0x7D, 0x52)
CREAM = RGBColor(0xF7, 0xF3, 0xEA)
LINE = RGBColor(0xD6, 0xD0, 0xC4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xFD, 0xFB, 0xF6)
NIGHT = RGBColor(0x1E, 0x29, 0x3B)

EN = "Georgia"
CN = "微软雅黑"

W, H = 13.333, 7.5
PAD = 0.62

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]
page_no = [0]


# ---------------- 基础工具 ----------------
def _mk(tag):
    return etree.fromstring(
        '<a:%s xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>' % tag)


def _fix_fit(tb):
    tx = tb._element.find(qn("p:txBody"))
    if tx is None:
        return
    bp = tx.find(qn("a:bodyPr"))
    if bp is None:
        return
    for tag in ("a:spAutoFit", "a:normAutofit"):
        el = bp.find(qn(tag))
        if el is not None:
            bp.remove(el)
    if bp.find(qn("a:noAutofit")) is None:
        bp.append(_mk("noAutofit"))


def add_slide():
    page_no[0] += 1
    s = prs.slides.add_slide(BLANK)
    bg = s.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = PAPER
    return s


def text(slide, x, y, w, h, s, size=18, color=INK, bold=False, font=CN,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line=1.15, space_after=2):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for a in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, a, 0)
    tf.vertical_anchor = anchor
    for i, ln in enumerate(str(s).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        p.space_before = Pt(0)
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color
        rPr = r._r.get_or_add_rPr()
        rPr.append(rPr.makeelement(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}ea", {"typeface": CN}))
    _fix_fit(tb)
    return tb


def rect(slide, x, y, w, h, fill=None, line_color=None, line_w=1.0, radius=None, shadow=False):
    if radius is None:
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                     Inches(w), Inches(h))
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(line_w)
    if not shadow:
        try:
            spPr = shp._element.spPr
            if spPr.find(qn("a:effectLst")) is None:
                spPr.append(_mk("effectLst"))
        except Exception:
            pass
    shp.text_frame.word_wrap = True
    return shp


def hline(slide, x1, y1, x2, y2, color=LINE, w=1.0):
    from pptx.enum.shapes import MSO_CONNECTOR
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    return c


def picture(slide, name, x, y, w=None, h=None):
    p = os.path.join(FIG, name)
    if w and h:
        return slide.shapes.add_picture(p, Inches(x), Inches(y), Inches(w), Inches(h))
    if w:
        return slide.shapes.add_picture(p, Inches(x), Inches(y), width=Inches(w))
    return slide.shapes.add_picture(p, Inches(x), Inches(y), height=Inches(h))


def page(kicker, title, sub=None):
    s = add_slide()
    rect(s, 0, 0, 0.16, H, fill=NAVY)
    text(s, PAD, 0.40, 10.6, 0.30, kicker, 12, GOLD, bold=True)
    text(s, PAD, 0.66, 11.6, 0.58, title, 26, NAVY, bold=True)
    y = 1.26
    if sub:
        text(s, PAD, 1.22, 11.9, 0.34, sub, 13.5, MUTED)
        y = 1.62
    hline(s, PAD, y, W - PAD, y, LINE, 1.2)
    # 页脚
    hline(s, PAD, H - 0.52, W - PAD, H - 0.52, LINE, 0.75)
    text(s, PAD, H - 0.44, 9.6, 0.28, "2026 年新课标 II 卷 · 数学 · 第 17 题", 10.5, MUTED)
    text(s, W - PAD - 1.2, H - 0.44, 1.2, 0.28, "%02d" % page_no[0], 11, MUTED,
         font=EN, align=PP_ALIGN.RIGHT)
    return s


def fblock(slide, x, y, w, lines, size=14, color=INK, title=None, lh=0.42,
           accent=TEAL, gap=0.20):
    h = lh * len(lines) + (0.44 if title else 0.20)
    rect(slide, x, y, w, h, fill=WHITE, line_color=LINE, line_w=1.0)
    rect(slide, x, y, 0.05, h, fill=accent)
    yy = y + 0.18
    if title:
        text(slide, x + 0.24, yy - 0.02, w - 0.40, 0.26, title, 11.5, accent, bold=True)
        yy += 0.26
    for ln in lines:
        text(slide, x + 0.26, yy, w - 0.44, lh, ln, size, color)
        yy += lh
    return y + h + gap


def card(slide, x, y, w, h, title, body, accent=TEAL, size=12.5, tsize=14.5):
    rect(slide, x, y, w, h, fill=WHITE, line_color=LINE, line_w=1.0)
    rect(slide, x, y, w, 0.055, fill=accent)
    text(slide, x + 0.20, y + 0.16, w - 0.40, 0.28, title, tsize, NAVY, bold=True)
    text(slide, x + 0.20, y + 0.54, w - 0.40, h - 0.62, body, size, INK, line=1.30)


def table(slide, x, y, w, rows, col_w, size=12, head=None, accents=None, rh=0.52):
    """轻量表格：rows 为字符串列表的列表"""
    cy = y
    if head:
        rect(slide, x, cy, w, rh, fill=CREAM)
        cx = x
        for j, cell in enumerate(head):
            text(slide, cx + 0.16, cy + 0.14, col_w[j] - 0.2, 0.3, cell, size - 0.5, NAVY, bold=True)
            cx += col_w[j]
        cy += rh
    for i, row in enumerate(rows):
        if i % 2 == 0:
            rect(slide, x, cy, w, rh, fill=WHITE)
        else:
            rect(slide, x, cy, w, rh, fill=RGBColor(0xFB, 0xF9, 0xF4))
        cx = x
        for j, cell in enumerate(row):
            col = INK
            if accents and j in accents:
                col = accents[j]
            text(slide, cx + 0.16, cy + 0.13, col_w[j] - 0.2, 0.32, cell, size, col)
            cx += col_w[j]
        cy += rh
    hline(slide, x, cy, x + w, cy, LINE, 0.8)
    return cy


# ============================================================
# 1 封面
# ============================================================
s = add_slide()
rect(s, 0, 0, W, H, fill=NAVY)
for i in range(8):
    rect(s, W - 4.1 + i * 0.05, 5.72 - i * 0.10, 3.9 - i * 0.44, 0.016,
         fill=RGBColor(0x2B, 0x3A, 0x5C))
text(s, 1.05, 1.30, 11.5, 0.34,
     "2026 年全国普通高等学校招生统一考试 · 新课标 II 卷 · 数学 第 17 题", 14.5, GOLD, bold=True)
text(s, 1.05, 1.78, 11.4, 1.90, "解三角形：\n从角的恒等式到周长", 40, WHITE, bold=True, line=1.26)
hline(s, 1.05, 4.12, 4.2, 4.12, GOLD, 2.0)
text(s, 1.05, 4.38, 11.0, 0.80,
     "一个恒等式  ·  两个公式  ·  两问一体", 17, RGBColor(0xC9, 0xD2, 0xE2))
rect(s, 1.05, 5.34, 6.20, 0.86, fill=RGBColor(0x24, 0x33, 0x55), line_color=GOLD, line_w=1.2)
text(s, 1.24, 5.50, 5.9, 0.30, "答案", 11.5, GOLD, bold=True)
text(s, 1.24, 5.74, 5.9, 0.36, "钝角三角形；周长 3 + √2 ≈ 4.4142", 17, WHITE, bold=True)
text(s, 1.05, 6.62, 9.0, 0.30, "原题与来源 · 提示词 · 代码 · 演示文稿", 12, RGBColor(0x9A, 0xA6, 0xBE))

# ============================================================
# 2 原题与来源
# ============================================================
s = page("ORIGINAL PROBLEM", "原题与来源", "题面一字未改；后面所有推导都从这三行出发。")
rect(s, PAD, 1.86, W - 2 * PAD, 3.30, fill=WHITE, line_color=LINE, line_w=1.2)
rect(s, PAD, 1.86, 0.075, 3.30, fill=RED)
text(s, PAD + 0.34, 2.04, 11.4, 0.40,
     "在 △ABC 中，已知 cos B = 3/4，cos²(A+C) + sin A·sin C = 1。", 19, INK)
text(s, PAD + 0.34, 2.62, 11.4, 1.05,
     "(1) 证明：△ABC 为钝角三角形；\n(2) 若 △ABC 的面积为 √7/4，求 △ABC 的周长。",
     19, INK, line=1.62)
text(s, PAD + 0.34, 3.96, 11.4, 0.40,
     "（解答题，满分 15 分）", 12.5, MUTED)
text(s, PAD, 5.32, 7.2, 0.80,
     "出处：2026 年普通高等学校招生全国统一考试（新课标 II 卷）\n数学 第 17 题 · 原卷序号 17", 12, MUTED, line=1.4)
card(s, 8.10, 5.22, 4.60, 1.34, "参考答案",
     "钝角三角形；周长 3 + √2", GOLD, 13)
text(s, PAD, 6.42, 12.0, 0.40,
     "本题的关键是把 A + C 与 B 的互补关系，翻译成余弦互为相反数 —— 一条恒等式贯穿两问。",
     11.5, INK, line=1.3)

# ============================================================
# 3 条件梳理
# ============================================================
s = page("STEP 0", "先把条件翻译成可用的式子", "三角形的三个角有关系 A + B + C = π，这是所有推导的起点。")
y = 1.88
y = fblock(s, PAD, y, 6.55, [
    "A + B + C = π  ⇒  A + C = π − B",
    "⇒ cos(A+C) = cos(π−B) = −cos B = −3/4",
], title="① 角的互补关系", size=13, lh=0.42, accent=TEAL)
y = fblock(s, PAD, y, 6.55, [
    "cos²(A+C) + sin A·sin C = 1",
    "⇒ sin A·sin C = 1 − (9/16) = 7/16",
], title="② 已知恒等式", size=13, lh=0.42, accent=GOLD)
y = fblock(s, PAD, y, 6.55, [
    "cos(A+C) = cos A·cos C − sin A·sin C",
    "⇒ cos A·cos C = cos(A+C) + sin A·sin C",
], title="③ 和角公式（把乘积拆开）", size=13, lh=0.42, accent=BLUE)
text(s, PAD, y + 0.02, 6.55, 0.60,
     "三条式子各用一次，就是第 (1) 问的全部。", 12.5, INK)
picture(s, "fig2_supplement.png", 7.55, 1.88, h=4.55)

# ============================================================
# 4 (1) 三步
# ============================================================
s = page("PART 1 · 证明", "第 (1) 问：三步判断钝角", "目标不是求角，而是定出 cos A·cos C 的符号。")
y = 1.84
y = fblock(s, PAD, y, 7.60, [
    "cos(A+C) = −cos B = −3/4",
    "sin A·sin C = 1 − cos²(A+C) = 1 − 9/16 = 7/16",
], title="第 1 步  把已知恒等式拆成两个乘积", size=14, lh=0.42, accent=TEAL, gap=0.18)
y = fblock(s, PAD, y, 7.60, [
    "cos A·cos C = cos(A+C) + sin A·sin C",
    "            = −3/4 + 7/16 = −12/16 + 7/16 = −5/16",
], title="第 2 步  用和角公式反解 cos A·cos C", size=14, lh=0.42, accent=GOLD, gap=0.18)
y = fblock(s, PAD, y, 7.60, [
    "cos A·cos C = −5/16 < 0 ⇒ cos A 与 cos C 异号",
    "又 cos B = 3/4 > 0 ⇒ B 为锐角 ⇒ A、C 中必有一个钝角",
], title="第 3 步  判断符号 ⇒ 结论", size=14, lh=0.42, accent=RED, gap=0.20)
text(s, PAD + 0.04, y + 0.04, 7.6, 0.38, "⇒ △ABC 为钝角三角形", 18, RED, bold=True)
card(s, 8.42, 1.86, 4.30, 2.30, "为什么这样想",
     "题目只给角的关系、不给边长，\n"
     "所以应当把“钝角”转化为\n“某两个余弦之积为负”。\n"
     "而乘积恰好可由和角公式从\n"
     "已知的两个乘积反解。", TEAL, 12)
card(s, 8.42, 4.44, 4.30, 1.85, "易错提醒",
     "① 把 cos(A+C) 误算为 +cos B；\n② 由 cos(A+C)<0 直接断定 A+C 为钝角，\n"
     "   从而误判“哪一个是钝角”。", GOLD, 12)

# ============================================================
# 5 (1) 定位钝角 + 图1
# ============================================================
s = page("PART 1 · 深入", "到底是哪一个角钝？", "再补一步：解出 cos(A−C)，就能把三个角都定下来。")
y = 1.82
y = fblock(s, PAD, y, 6.00, [
    "cos(A−C) = cos A·cos C + sin A·sin C",
    "         = −5/16 + 7/16 = 1/8",
], title="① 再求一个和差余弦", size=13.5, lh=0.42, accent=TEAL, gap=0.18)
y = fblock(s, PAD, y, 6.00, [
    "A+C = arccos(−3/4) ≈ 138.59°",
    "|A−C| = arccos(1/8) ≈ 82.82°",
    "⇒ A ≈ 110.70°，C ≈ 27.89°，B ≈ 41.41°",
    "（三内角和恰为 180°；A、C 可互换）",
], title="② 两式相加、相减还原每个角", size=13.5, lh=0.42, accent=GOLD, gap=0.18)
card(s, PAD, y, 6.00, 1.00, "结论",
     "A 为钝角（≈110.70°），B、C 为锐角。", RED, 12.5)
picture(s, "fig1_triangle.png", 7.30, 1.84, h=4.62)

# ============================================================
# 6 (2) 思路图
# ============================================================
s = page("PART 2 · 思路", "第 (2) 问：一条链式推导", "从面积出发，五次代换直达周长；每一步只用一条公式。")
picture(s, "fig4_chain.png", PAD, 1.80, w=12.05)
y = 5.70
y = fblock(s, PAD, y, 3.85, [
    "面积 → ac；正弦定理 → b；",
    "余弦定理 + 配方 → a+c。",
], size=12, lh=0.40, accent=TEAL)
y2 = fblock(s, 4.62, 5.70, 3.85, [
    "全程没有解出 A、C 的具体值，",
    "只用到 sin A·sin C = 7/16。",
], size=12, lh=0.40, accent=GOLD)
fblock(s, 8.60, 5.70, 4.10, [
    "把 (a+c) 整体当作未知量，",
    "是本题最省力的关键一步。",
], size=12, lh=0.40, accent=RED)

# ============================================================
# 7 (2) 面积 → ac
# ============================================================
s = page("PART 2 · 第 1 步", "由面积得到 ac = 2", "面积公式把“面积”换成“两边夹一角”，再代入 sin B。")
y = 1.88
y = fblock(s, PAD, y, 6.15, [
    "cos B = 3/4，且 B ∈ (0, π) ⇒ sin B > 0",
    "sin B = √(1 − cos²B) = √(1 − 9/16) = √7/4",
], title="① 先求 sin B", size=14, lh=0.44, accent=TEAL)
y = fblock(s, PAD, y, 6.15, [
    "S = (1/2)·a·c·sin B",
    "√7/4 = (1/2)·a·c·(√7/4)",
    "⇒ a·c = 2",
], title="② 代入面积公式", size=14, lh=0.44, accent=GOLD)
rect(s, PAD, y + 0.04, 6.15, 0.70, fill=CREAM, line_color=GOLD, line_w=1.2)
text(s, PAD + 0.24, y + 0.20, 5.7, 0.36, "得到第一个关键量：ac = 2", 16, RED, bold=True)
card(s, 7.05, 1.88, 5.65, 1.95, "注意",
     "sin B 必须取正号：B 是三角形的内角，B ∈ (0, π)，\n正弦恒为正；这也是后面所有开方的依据。", TEAL, 12.5)
card(s, 7.05, 4.02, 5.65, 1.95, "为什么先求 ac",
     "已知条件里有两个“乘积”：sin A·sin C 与 ac。\n"
     "两者由正弦定理相连（ac = 4R²·sin A·sin C），\n"
     "所以先算出 ac，才能求 R、进而求 b。", GOLD, 12.5)

# ============================================================
# 8 (2) 正弦定理 → b
# ============================================================
s = page("PART 2 · 第 2 步", "用正弦定理把 ac 变成 b = √2", "正弦定理把“边”换成“角的正弦”，正好接上第 (1) 问的乘积。")
y = 1.86
y = fblock(s, PAD, y, 6.15, [
    "a = 2R·sin A，c = 2R·sin C  （R 为外接圆半径）",
    "⇒ ac = 4R²·sin A·sin C = 4R²·(7/16) = 7R²/4",
], title="① 正弦定理（两边同乘）", size=13.5, lh=0.42, accent=TEAL)
y = fblock(s, PAD, y, 6.15, [
    "7R²/4 = 2  ⇒  R² = 8/7",
], title="② 代入 ac = 2", size=13.5, lh=0.42, accent=GOLD)
y = fblock(s, PAD, y, 6.15, [
    "b = 2R·sin B = 2·√(8/7)·(√7/4)",
    "  = (2·√8·√7) / (4·√7) = √8/2 = √2",
], title="③ 再算 b", size=13.5, lh=0.42, accent=BLUE)
rect(s, PAD, y + 0.04, 6.15, 0.66, fill=CREAM, line_color=BLUE, line_w=1.2)
text(s, PAD + 0.24, y + 0.18, 5.7, 0.34, "b = √2 —— 第二个关键量", 16, RED, bold=True)
picture(s, "fig3_circumcircle.png", 7.75, 1.86, h=4.50)

# ============================================================
# 9 (2) 余弦定理 → a+c
# ============================================================
s = page("PART 2 · 第 3 步", "余弦定理 + 配方，得到 a + c = 3", "不求 a、c 各自的值，而是直接求它们的和。")
y = 1.84
y = fblock(s, PAD, y, 6.20, [
    "b² = a² + c² − 2ac·cos B",
    "2 = a² + c² − 2·2·(3/4) = a² + c² − 3",
    "⇒ a² + c² = 5",
], title="① 余弦定理", size=13.5, lh=0.42, accent=TEAL, gap=0.18)
y = fblock(s, PAD, y, 6.20, [
    "(a+c)² = a² + c² + 2ac = 5 + 2·2 = 9",
    "a + c = 3    （取正号，边长为正）",
], title="② 完全平方公式（关键一步）", size=13.5, lh=0.42, accent=GOLD, gap=0.18)
y = fblock(s, PAD, y, 6.20, [
    "周长 = a + b + c = (a+c) + b = 3 + √2",
], title="③ 直接得到周长", size=13.5, lh=0.42, accent=RED, gap=0.18)
rect(s, PAD, y, 6.20, 0.62, fill=CREAM, line_color=RED, line_w=1.4)
text(s, PAD + 0.24, y + 0.15, 5.7, 0.34, "答案：3 + √2 ≈ 4.4142", 16.5, RED, bold=True)
card(s, 7.10, 1.86, 5.60, 2.10, "为什么不解 a、c",
     "解 a、c 需要联立 ac=2 与 a+c=3，得到 t²−3t+2=0，\n"
     "即 a、c 为 1 与 2 —— 但要区分谁对谁，还需判断 A、C 的大小，\n"
     "而周长只依赖 a+c，所以完全不必分开解。", TEAL, 12)
card(s, 7.10, 4.20, 5.60, 2.20, "两步检查",
     "① (a+c)² = 9 ⇒ a+c = 3（舍去 −3）；\n"
     "② 三边 1、2、√2 满足 1+√2 > 2，三角形存在；\n"
     "   且 a=2 对应的角 A 为钝角，与 (1) 一致。", GOLD, 12)

# ============================================================
# 10 结论 + 验算
# ============================================================
s = page("ANSWER", "结论与自洽验算", "答案是 3 + √2；用三边回代，五条条件全部成立。")
rect(s, PAD, 1.84, 4.55, 2.30, fill=NAVY)
text(s, PAD + 0.30, 2.02, 3.9, 0.30, "第 (2) 问答案", 12, GOLD, bold=True)
text(s, PAD + 0.30, 2.36, 3.9, 0.62, "周长 = 3 + √2", 30, WHITE, bold=True)
text(s, PAD + 0.30, 3.10, 3.95, 0.85,
     "≈ 4.4142135624\n三边为 1、√2、2（含一个钝角）", 12.5, RGBColor(0xC9, 0xD2, 0xE2), line=1.35)
card(s, PAD, 4.34, 4.55, 1.90, "第 (1) 问答案",
     "△ABC 为钝角三角形：\ncos A·cos C = −5/16 < 0，A 为钝角（≈110.70°）。", RED, 12.5)
picture(s, "fig5_check.png", 5.34, 1.86, w=7.40)

# ============================================================
# 11 数值验证
# ============================================================
s = page("VERIFICATION", "数值验证：Decimal 40 位复核",
         "solve.py 独立计算 17 项、全部通过；下表列出与前两问直接相关的 11 项。")
rows = [
    ["cos(A+C) = −cos B", "−0.75 = −3/4", "✓"],
    ["sin A·sin C", "0.4375 = 7/16", "✓"],
    ["cos A·cos C", "−0.3125 = −5/16", "✓"],
    ["cos(A−C)", "0.125 = 1/8", "✓"],
    ["A + B + C", "110.7048° + 41.4096° + 27.8856° = 180°", "✓"],
    ["a·c = 2S / sin B", "2.000000000000000000000000000000", "✓"],
    ["R² = ac / (4 sin A sin C)", "1.142857142857142857142857142857", "✓"],
    ["b = 2R·sin B", "1.414213562373095048801688724210", "✓"],
    ["a² + c²", "5.000000000000000000000000000000", "✓"],
    ["a + c", "3.000000000000000000000000000000", "✓"],
    ["周长 a + b + c", "4.414213562373095048801688724210", "✓"],
]
table(s, PAD, 1.84, 12.10, rows, [3.70, 7.30, 1.10],
      size=11.5, head=["检验项", "计算值（Decimal 40 位）", "结论"],
      accents={2: GREEN}, rh=0.415)

# ============================================================
# 12 提示词
# ============================================================
s = page("PROMPT", "提示词设计", "完整提示词见「其他材料/提示词.txt」。")
card(s, PAD, 1.86, 5.90, 4.42, "提示词要点",
     "1. 逐字引用原题，注明出处与分值；\n\n"
     "2. 先分析“题目给了什么、要求什么”，再动笔；\n\n"
     "3. 每一步注明依据的公式，并说明动机；\n\n"
     "4. 明确要求“不解 a、c，只求 a+c”这类优化思路；\n\n"
     "5. 要求给出末位验算（三边回代），并用 40 位小数复核；\n\n"
     "6. 输出 16:9 文稿，插图自绘。", TEAL, 12)
card(s, 7.02, 1.86, 5.68, 4.42, "验收标准",
     "· 结论与官方参考答案一致（3 + √2）；\n\n"
     "· 第 (1) 问不能只证“有一个钝角”，\n"
     "  还要说明钝角是 A（或 C）；\n\n"
     "· 每个开方都交代取舍理由；\n\n"
     "· 三边回代后五条条件全部成立；\n\n"
     "· 版面无溢出、字号统一、图表可读。", GOLD, 12)

# ============================================================
# 13 代码
# ============================================================
s = page("CODE", "代码清单", "全部代码在「其他材料/代码」中，无需联网即可复现。")
code = """# solve.py —— Decimal(40) 高精度复核
from decimal import Decimal as D, getcontext
getcontext().prec = 40

cosB = D(3) / D(4)
sinB = (1 - cosB**2).sqrt()          # √7/4

# 第 (1) 问：证明钝角
cosAC = -cosB                        # cos(A+C) = −3/4
sinA_sinC = 1 - cosAC**2             # 7/16
cosA_cosC = cosAC + sinA_sinC        # −5/16 < 0 ⇒ 钝角

# 第 (2) 问：求周长
ac = 2 * (D(7).sqrt() / 4) / sinB    # ac = 2
R2 = ac / (4 * sinA_sinC)            # R² = 8/7
b = 2 * R2.sqrt() * sinB             # b = √2
a2c2 = b * b + 2 * ac * cosB         # a² + c² = 5
print((a2c2 + 2 * ac).sqrt() + b)    # 3 + √2"""
rect(s, PAD, 1.84, 12.10, 4.30, fill=NIGHT)
text(s, PAD + 0.26, 1.96, 11.7, 4.05, code, 9.5,
     RGBColor(0xE6, 0xED, 0xF7), font="Consolas", line=1.24)
text(s, PAD, 6.24, 12.0, 0.36,
     "辅助脚本：figs.py（Pillow 插图）、build_ppt.py（本幻灯片）；运行环境 Python 3 + python-pptx + Pillow。",
     10.5, MUTED)

# ============================================================
# 14 结语
# ============================================================
s = add_slide()
rect(s, 0, 0, W, H, fill=NAVY)
text(s, 1.05, 1.55, 10.5, 0.32, "TAKEAWAY", 13.5, GOLD, bold=True, font=EN)
text(s, 1.05, 1.94, 11.2, 2.05,
     "一道三角形题的关键，\n往往只是把“角的关系”翻译成“边的乘积”。", 30, WHITE, bold=True, line=1.34)
hline(s, 1.05, 4.30, 4.4, 4.30, GOLD, 2.0)
text(s, 1.05, 4.54, 11.2, 1.55,
     "A + C = π − B  →  cos 互为相反数  →  cos A·cos C 的符号；\n"
     "面积  →  ac  →  R  →  b  →  余弦定理  →  a + c  →  周长。", 15.5,
     RGBColor(0xC9, 0xD2, 0xE2), line=1.6)
text(s, 1.05, 6.55, 9.0, 0.30,
     "2026 新课标 II 卷 · 第 17 题 · 周长 3 + √2", 12, RGBColor(0x9A, 0xA6, 0xBE))

prs.save(OUT)
print("saved:", OUT)
print("slides:", len(prs.slides))
