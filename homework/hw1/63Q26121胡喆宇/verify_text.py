# -*- coding: utf-8 -*-
from pptx import Presentation

p = Presentation(r'D:\AI\pptdemo\高考压轴题完整评分解答.pptx')
lines = []
lines.append('slides: %d' % len(p.slides))
for i, s in enumerate(p.slides):
    lines.append('\n===== SLIDE %d (layout=%r) =====' % (i+1, s.slide_layout.name))
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            lines.append('  [%s]' % sh.text_frame.text.strip().replace('\n', ' ⏎ '))
open(r'D:\AI\pptdemo\verify_text.txt', 'w', encoding='utf-8').write('\n'.join(lines))
print('done', len(lines))
