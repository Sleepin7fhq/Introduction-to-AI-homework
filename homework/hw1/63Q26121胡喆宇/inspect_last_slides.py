# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.oxml.ns import qn

p = Presentation(r'D:\AI\pptdemo\东南大学PPT.pptx')
lines = []
n = len(p.slides)
lines.append('total slides: %d' % n)
for idx in range(max(0, n-8), n):
    s = p.slides[idx]
    lines.append('\n===== SLIDE %d (layout=%r) =====' % (idx, s.slide_layout.name))
    for sh in s.shapes:
        typ = str(sh.shape_type)
        pos = '(%.2f,%.2f) size=(%.2f,%.2f)' % (sh.left/914400, sh.top/914400, sh.width/914400, sh.height/914400)
        name = getattr(sh, 'name', '')
        txt = ''
        if sh.has_text_frame and sh.text_frame.text.strip():
            txt = ' text=%r' % sh.text_frame.text.strip()[:40]
        img = ''
        if sh.shape_type == 13:  # PICTURE
            try:
                img = ' img=%s' % sh.image.filename
            except Exception as e:
                img = ' img=<err>'
        # rels / image part name
        lines.append('  %s name=%r%s%s%s' % (typ, name, pos, txt, img))

open(r'D:\AI\pptdemo\last_slides.txt', 'w', encoding='utf-8').write('\n'.join(lines))
print('done', len(lines))
