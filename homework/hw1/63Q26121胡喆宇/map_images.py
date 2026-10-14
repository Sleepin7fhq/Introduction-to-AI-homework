# -*- coding: utf-8 -*-
from pptx import Presentation
from PIL import Image
import io

p = Presentation(r'D:\AI\pptdemo\东南大学PPT.pptx')
lines = []
n = len(p.slides)
seen = {}
for idx in range(max(0, n-8), n):
    s = p.slides[idx]
    for sh in s.shapes:
        if sh.shape_type == 13:  # PICTURE
            try:
                img = sh.image
                blob = img.blob
                sha1 = img.sha1
                im = Image.open(io.BytesIO(blob))
                dims = '%dx%d' % im.size
                fmt = im.format
                if sha1 not in seen:
                    seen[sha1] = True
                    lines.append('IMG sha1=%s  %s  %s  ext=%s' % (sha1[:10], dims, fmt, img.ext))
            except Exception as e:
                lines.append('err %s' % e)

open(r'D:\AI\pptdemo\map_images.txt', 'w', encoding='utf-8').write('\n'.join(lines))
print('done', len(seen))
