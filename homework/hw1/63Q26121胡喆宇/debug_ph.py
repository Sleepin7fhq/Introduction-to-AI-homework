# -*- coding: utf-8 -*-
from pptx import Presentation
p = Presentation(r'D:\AI\pptdemo\高考压轴题完整评分解答.pptx')
for i, s in enumerate(p.slides):
    phs = []
    for sh in s.placeholders:
        phs.append((sh.placeholder_format.idx, str(sh.placeholder_format.type), repr(sh.text_frame.text)))
    print('SLIDE %d placeholders:' % (i+1), phs)
    if i >= 2:
        break
