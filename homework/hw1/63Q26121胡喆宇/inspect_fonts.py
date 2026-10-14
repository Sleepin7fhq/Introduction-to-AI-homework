# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Pt

p = Presentation(r'D:\AI\pptdemo\东南大学PPT.pptx')
lines = []

def run_info(run):
    f = run.font
    color = None
    try:
        if f.color and f.color.type is not None:
            color = f.color.rgb
    except Exception:
        color = '?'
    return 'run=%r font=%s size=%s bold=%s color=%s' % (run.text[:30], f.name, f.size, f.bold, color)

def dump_shape(sh, indent='  '):
    lines.append(indent + 'SHAPE name=%r type=%s' % (getattr(sh, 'name', None), sh.shape_type))
    if sh.has_text_frame:
        for pi, para in enumerate(sh.text_frame.paragraphs):
            lines.append(indent + '  para%d align=%s' % (pi, para.alignment))
            for r in para.runs:
                lines.append(indent + '    ' + run_info(r))

# title slide layout placeholders
lines.append('==== LAYOUT 0 (标题幻灯片) placeholders ====')
for sh in p.slide_layouts[0].placeholders:
    lines.append('placeholder idx=%s type=%s pos=(%.2f,%.2f)' % (sh.placeholder_format.idx, sh.placeholder_format.type, sh.left/914400, sh.top/914400))
    dump_shape(sh)

lines.append('==== LAYOUT 8 (仅标题) ====')
for sh in p.slide_layouts[8].placeholders:
    lines.append('placeholder idx=%s type=%s' % (sh.placeholder_format.idx, sh.placeholder_format.type))
    dump_shape(sh)

lines.append('==== LAYOUT 15 (末尾) placeholders ====')
for sh in p.slide_layouts[15].placeholders:
    lines.append('placeholder idx=%s type=%s' % (sh.placeholder_format.idx, sh.placeholder_format.type))
    dump_shape(sh)

lines.append('==== ACTUAL SLIDE 0 title runs ====')
for sh in p.slides[0].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        dump_shape(sh)

# theme relationship of master 0
lines.append('==== MASTER rels ====')
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
for rel in p.slide_masters[0].part.rels.values():
    if 'theme' in rel.reltype:
        lines.append('master0 theme rel: %s -> %s' % (rel.reltype, rel.target_ref))

open(r'D:\AI\pptdemo\font_dump.txt', 'w', encoding='utf-8').write('\n'.join(lines))
print('done', len(lines))
