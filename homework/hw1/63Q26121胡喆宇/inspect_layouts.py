# -*- coding: utf-8 -*-
from pptx import Presentation

p = Presentation(r'D:\AI\pptdemo\东南大学PPT.pptx')
lines = []

def desc_shape(sh):
    info = []
    if sh.has_text_frame:
        txt = sh.text_frame.text.strip().replace('\n', ' / ')
        info.append('text=%r' % txt[:60])
    if sh.is_placeholder:
        info.append('ph_type=%s ph_idx=%s' % (sh.placeholder_format.type, sh.placeholder_format.idx))
    info.append('pos=(%.2f,%.2f) size=(%.2f,%.2f)' % (sh.left/914400, sh.top/914400, sh.width/914400, sh.height/914400))
    return '    ' + ' | '.join(info)

lines.append('========== SLIDE MASTERS ==========')
for i, m in enumerate(p.slide_masters):
    lines.append('MASTER %d (layouts: %d)' % (i, len(m.slide_layouts)))
    for sh in m.shapes:
        lines.append(desc_shape(sh))

lines.append('========== LAYOUTS (all) ==========')
for i, l in enumerate(p.slide_layouts):
    master_name = 'M0' if l.slide_master is p.slide_masters[0] else 'M1'
    lines.append('LAYOUT %d | %s (%s)' % (i, l.name, master_name))
    for sh in l.shapes:
        lines.append(desc_shape(sh))

lines.append('========== SAMPLE SLIDES ==========')
for idx in [0, 1, 2, 3, 4, 5]:
    if idx < len(p.slides):
        s = p.slides[idx]
        lines.append('SLIDE %d (layout=%r)' % (idx, s.slide_layout.name))
        for sh in s.shapes:
            lines.append(desc_shape(sh))

open(r'D:\AI\pptdemo\layout_dump.txt', 'w', encoding='utf-8').write('\n'.join(lines))
print('done', len(lines))
