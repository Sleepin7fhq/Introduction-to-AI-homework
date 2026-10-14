# -*- coding: utf-8 -*-
from pptx import Presentation
from lxml import etree

p = Presentation(r'D:\AI\pptdemo\东南大学PPT.pptx')
lines = []
ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}

def txtbody_info(name, shape):
    lines.append('--- %s ---' % name)
    if not shape.has_text_frame:
        return
    txBody = shape.text_frame._txBody
    # bodyPr
    bp = txBody.find('a:bodyPr', ns)
    lines.append('bodyPr anchor=%s' % (bp.get('anchor') if bp is not None else None))
    for p in txBody.findall('a:p', ns):
        ppr = p.find('a:pPr', ns)
        lvl = ppr.get('lvl') if ppr is not None else None
        lines.append('  p lvl=%s' % lvl)
        for r in p.findall('a:r', ns):
            rpr = r.find('a:rPr', ns)
            if rpr is not None:
                sz = rpr.get('sz'); b = rpr.get('b')
                latin = rpr.find('a:latin', ns)
                ea = rpr.find('a:ea', ns)
                sf = rpr.find('a:solidFill', ns)
                col = None
                if sf is not None:
                    c = sf.find('a:srgbClr', ns)
                    if c is not None:
                        col = c.get('val')
                lines.append('    run sz=%s bold=%s latin=%s ea=%s color=%s' % (sz, b, latin.get('typeface') if latin is not None else None, ea.get('typeface') if ea is not None else None, col))
            t = r.find('a:t', ns)
            lines.append('      text=%r' % (t.text[:40] if t is not None and t.text else ''))

# master title/body
for i, m in enumerate(p.slide_masters):
    for sh in m.placeholders:
        if sh.placeholder_format.type == 1:  # TITLE
            txtbody_info('MASTER %d TITLE' % i, sh)
        if sh.placeholder_format.type == 2:  # BODY
            txtbody_info('MASTER %d BODY' % i, sh)

# layout 0 title placeholder idx 10, 11
for sh in p.slide_layouts[0].placeholders:
    if sh.placeholder_format.idx in (10, 11, 12):
        txtbody_info('LAYOUT0 ph%d' % sh.placeholder_format.idx, sh)

# layout 8 section title idx 11
for sh in p.slide_layouts[8].placeholders:
    if sh.placeholder_format.idx == 11:
        txtbody_info('LAYOUT8 章节标题 ph11', sh)

# layout 7 main title idx 12
for sh in p.slide_layouts[7].placeholders:
    if sh.placeholder_format.idx in (12, 13):
        txtbody_info('LAYOUT7 ph%d' % sh.placeholder_format.idx, sh)

open(r'D:\AI\pptdemo\size_dump.txt', 'w', encoding='utf-8').write('\n'.join(lines))
print('done')
