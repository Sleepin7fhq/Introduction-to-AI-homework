# -*- coding: utf-8 -*-
from pptx import Presentation

p = Presentation(r'D:\AI\pptdemo\高考压轴题完整评分解答2.pptx')
lines = []
lines.append('slides: %d' % len(p.slides))
for i, s in enumerate(p.slides):
    pics = []
    for sh in s.shapes:
        if sh.shape_type == 13:
            try:
                im = sh.image
                pics.append('PIC %dx%d @(%.2f,%.2f,%.2f,%.2f)' % (im.size[0], im.size[1], sh.left/914400, sh.top/914400, sh.width/914400, sh.height/914400))
            except Exception as e:
                pics.append('PIC err %s' % e)
    if pics:
        lines.append('slide %d pictures:' % (i+1))
        lines.extend(['  '+p for p in pics])
    # verify title text present
    if i == 0:
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        lines.append('  title text present: %s' % any('压轴题' in t for t in texts))
    if i == 14:
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        lines.append('  last text present: %s' % any('命题逻辑' in t for t in texts))

open(r'D:\AI\pptdemo\verify_v2.txt', 'w', encoding='utf-8').write('\n'.join(lines))
print('done')
