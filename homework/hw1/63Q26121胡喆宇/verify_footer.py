# -*- coding: utf-8 -*-
from pptx import Presentation
p = Presentation(r'D:\AI\pptdemo\高考压轴题完整评分解答.pptx')
lines = []
for i, s in enumerate(p.slides):
    texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    has_motto = any('止于至善' in t for t in texts)
    has_pn = any(t.strip() == str(i+1) for t in texts)
    lines.append('slide %d: motto=%s pageno=%s' % (i+1, has_motto, has_pn))
open(r'D:\AI\pptdemo\verify_footer.txt', 'w', encoding='utf-8').write('\n'.join(lines))
print('done')
